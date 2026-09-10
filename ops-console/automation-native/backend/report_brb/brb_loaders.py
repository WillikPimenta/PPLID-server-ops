# -*- coding: utf-8 -*-
"""Leitura e preparação das abas do BRB_Report_atualizado.xlsx."""
from __future__ import annotations

from functools import lru_cache
from pathlib import Path

import pandas as pd

from report_brb.brb_filters import (
    add_case_columns,
    classificar_conforme,
    classificar_fn_fp,
    contestacao_batimento_key_from_row,
    contestacao_batimento_keys_from_df,
    dedupe_by_case_key,
    filter_period,
    filter_period_any,
    is_possivel_ataque,
    match_client,
    norm_protocolo,
    parse_excel_date,
    safe_str,
)
from report_brb.client_registry import get_client_config
from report_brb.config_brb import ATTACK_KEYWORDS, EXPECTED_SHEETS, EXCEL_PATH, MES_PT, SUPPLEMENT_SHEETS
from report_brb.na_falhas_normalize import (
    normalize_na_falhas_raw,
    read_na_falhas_excel,
    supplement_na_falhas_satisfied,
)

_AUDITADOS_USECOLS = (
    "Data",
    "Data análise",
    "Protocolo",
    "Etapa",
    "Cliente",
    "Workflow",
    "Tipo de análise",
    "Cenário",
    "STATUS",
    "IRREGULARIDADES_APONTADAS",
    "Novo cenário",
    "Cadastrado anteriormente",
    "Resultado Destino",
    "Resultado Origem",
    "Protocolo Destino",
)

_SUPPLEMENT_RAW_CACHE: dict[str, dict[str, pd.DataFrame]] = {}
_WORKBOOK_SHEETS_CACHE: dict[str, list[str]] = {}

# Códigos legados presentes na base BRB. O 240 representa a conclusão sem
# risco; os demais abaixo representam conclusões com risco em diferentes
# motores/cenários. Resultados textuais continuam sendo a fonte preferencial.
_BRB_RESULTADO_SEM_RISCO = {"240"}
_BRB_RESULTADO_COM_RISCO = {"100", "130", "180", "190", "220"}


def _risk_bucket(value) -> str:
    """Normaliza resultados textuais/legados para COM_RISCO ou SEM_RISCO."""
    raw = safe_str(value).upper()
    if not raw:
        return ""
    if "SEM RISCO" in raw:
        return "SEM_RISCO"
    if "COM RISCO" in raw:
        return "COM_RISCO"
    code = raw.removesuffix(".0")
    if code in _BRB_RESULTADO_SEM_RISCO:
        return "SEM_RISCO"
    if code in _BRB_RESULTADO_COM_RISCO:
        return "COM_RISCO"
    return ""


def _is_face_fraudadores(value) -> bool:
    text = safe_str(value).upper()
    return "FACE ENCONTRADA" in text and "FRAUDADOR" in text


def _risk_change_mask(df: pd.DataFrame) -> pd.Series:
    """Identifica mudança de classificação entre COM e SEM risco."""
    if df.empty:
        return pd.Series(dtype=bool, index=df.index)
    cliente = df["RESULTADO DO CLIENTE"].map(_risk_bucket)
    auditoria = df["RESULTADO DA AUDITORIA"].map(_risk_bucket)
    return (
        cliente.isin(("COM_RISCO", "SEM_RISCO"))
        & auditoria.isin(("COM_RISCO", "SEM_RISCO"))
        & cliente.ne(auditoria)
    )


def _notification_mask(df: pd.DataFrame) -> pd.Series:
    """Regra BRB: mudança COM↔SEM risco; face em fraudadores sempre entra."""
    if df.empty:
        return pd.Series(dtype=bool, index=df.index)
    face_fraudadores = df["MOTIVO DA FALHA"].map(_is_face_fraudadores)
    return _risk_change_mask(df) | face_fraudadores


class BRBDataBundle:
    def __init__(
        self,
        na_demandas: pd.DataFrame,
        na_falhas: pd.DataFrame,
        treinamentos: pd.DataFrame,
        falhas_gerais: pd.DataFrame,
        contestacao: pd.DataFrame,
        auditados: pd.DataFrame | None = None,
        treinamentos_horas: pd.DataFrame | None = None,
        inicio=None,
        fim=None,
        fg_linhas_brutas: int = 0,
        fg_duplicatas: int = 0,
        contestacao_merge_stats: dict[str, int] | None = None,
    ):
        self.na_demandas = na_demandas
        self.na_falhas = na_falhas
        self.treinamentos = treinamentos
        self.treinamentos_horas = (
            treinamentos_horas if treinamentos_horas is not None else pd.DataFrame()
        )
        self.falhas_gerais = falhas_gerais
        self.contestacao = contestacao
        self.auditados = auditados if auditados is not None else pd.DataFrame()
        self.inicio = inicio
        self.fim = fim
        self.fg_linhas_brutas = fg_linhas_brutas
        self.fg_duplicatas = fg_duplicatas
        self.contestacao_merge_stats = contestacao_merge_stats


def _missing_workbook_sheets(required: list[str], sheet_names: list[str]) -> list[str]:
    names = set(sheet_names)
    missing: list[str] = []
    for sheet in required:
        if sheet == "NA_Falhas":
            if not supplement_na_falhas_satisfied(names):
                missing.append(sheet)
        elif sheet not in names:
            missing.append(sheet)
    return missing


def validate_workbook(path: Path, *, mode: str = "auto") -> list[str]:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {path}")
    cache_key = f"{path.resolve()}::{path.stat().st_mtime_ns}"
    sheet_names = _WORKBOOK_SHEETS_CACHE.get(cache_key)
    if sheet_names is None:
        sheet_names = pd.ExcelFile(path).sheet_names
        _WORKBOOK_SHEETS_CACHE.clear()
        _WORKBOOK_SHEETS_CACHE[cache_key] = sheet_names
    resolved = mode
    if resolved == "auto":
        has_full = not _missing_workbook_sheets(EXPECTED_SHEETS, sheet_names)
        resolved = "full" if has_full else "supplement"
    required = EXPECTED_SHEETS if resolved == "full" else SUPPLEMENT_SHEETS
    missing = _missing_workbook_sheets(required, sheet_names)
    if missing:
        raise ValueError(
            f"Abas ausentes em {path.name} (modo {resolved}): {missing}. "
            f"Encontradas: {sheet_names}"
        )
    return sheet_names


def validate_portfolio_contestacao_workbook(path: Path) -> None:
    """Exige aba Contestacao para complementar contestados ausentes na EO."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {path}")
    xl = pd.ExcelFile(path)
    if "Contestacao" in xl.sheet_names:
        return
    raise ValueError(
        "A planilha precisa da aba Contestacao para complementar contestações ausentes na EO."
    )


def load_supplement_from_excel(
    path: Path,
    inicio=None,
    fim=None,
    client_slug: str = "brb",
) -> dict[str, pd.DataFrame]:
    """Carrega apenas NA + Treinamentos do Excel operacional."""
    path = Path(path)
    validate_workbook(path, mode="supplement")
    cache_key = f"{path.resolve()}::{path.stat().st_mtime_ns}"
    raw = _SUPPLEMENT_RAW_CACHE.get(cache_key)
    if raw is None:
        xl_names = set(_WORKBOOK_SHEETS_CACHE[cache_key])
        raw = {
            "na_demandas": pd.read_excel(path, "NA_Demandas"),
            "na_falhas": read_na_falhas_excel(path),
            "treinamentos": pd.read_excel(path, "Treinamentos"),
            "treinamentos_horas": pd.read_excel(path, "TreinamentosComHoras") if "TreinamentosComHoras" in xl_names else pd.DataFrame(),
        }
        _prepare_training_frames(raw)
        _SUPPLEMENT_RAW_CACHE.clear()
        _SUPPLEMENT_RAW_CACHE[cache_key] = raw
    treinamentos_horas = pd.DataFrame()
    if not raw["treinamentos_horas"].empty:
        treinamentos_horas = _prep_treinamentos_horas(
            raw["treinamentos_horas"], inicio, fim, client_slug
        )
    return {
        "na_demandas": _prep_na_demandas(
            raw["na_demandas"], inicio, fim, client_slug
        ),
        "na_falhas": _prep_na_falhas(
            raw["na_falhas"], inicio, fim, client_slug
        ),
        "treinamentos": _prep_treinamentos(
            raw["treinamentos"], inicio, fim, client_slug
        ),
        "treinamentos_horas": treinamentos_horas,
    }


def load_bundle_hybrid(
    path: Path | None = None,
    inicio=None,
    fim=None,
    client_slug: str = "brb",
    *,
    use_eo_db: bool = True,
    excel_contestacao: pd.DataFrame | None = None,
) -> BRBDataBundle:
    """EO (auditados/FG/contestação) + Excel suplemento (NA/treinamentos)."""
    client_slug = (client_slug or "brb").strip().lower()
    if not use_eo_db:
        return load_workbook(path, inicio=inicio, fim=fim, client_slug=client_slug)

    from report_brb.db_loaders import load_contestacao_batimento_keys, load_eo_core_frames

    eo = load_eo_core_frames(client_slug, inicio=inicio, fim=fim)
    falhas_gerais = eo["falhas_gerais"]
    contestacao = eo["contestacao"]
    contestacao_merge_stats: dict[str, int] | None = None

    contestacao_sheet = excel_contestacao
    if contestacao_sheet is None and path and Path(path).is_file():
        xl_names = set(pd.ExcelFile(path).sheet_names)
        if "Contestacao" in xl_names:
            contestacao_sheet = pd.read_excel(path, "Contestacao")
    if contestacao_sheet is not None and not contestacao_sheet.empty:
        from apps.brb_report.services.client_catalog import resolve_client_config

        cfg = resolve_client_config(client_slug)
        batimento_keys = load_contestacao_batimento_keys(
            int(eo["id_cliente"]),
            cliente_nome=cfg.get("nome", client_slug),
        )
        contestacao, contestacao_merge_stats = merge_contestacao_eo_with_excel(
            contestacao,
            contestacao_sheet,
            batimento_keys=batimento_keys,
            inicio=inicio,
            fim=fim,
            client_slug=client_slug,
        )

    sup = None
    # Quando o usuário envia um suplemento nesta geração, ele é a fonte da
    # visão de treinamento. O banco só funciona como fallback sem arquivo.
    if path and Path(path).exists():
        sup = load_supplement_from_excel(path, inicio, fim, client_slug)
    elif use_eo_db:
        try:
            from apps.brb_report.services.supplement_db_loaders import load_supplement_from_db

            sup = load_supplement_from_db(client_slug, inicio=inicio, fim=fim)
        except Exception:  # noqa: BLE001 — DB opcional
            sup = None

    if sup is None:
        sup = {
            "na_demandas": pd.DataFrame(),
            "na_falhas": pd.DataFrame(),
            "treinamentos": pd.DataFrame(),
            "treinamentos_horas": pd.DataFrame(),
        }

    na_falhas = _enrich_na_falhas(sup["na_falhas"], falhas_gerais)
    falhas_gerais = _join_conforme_fg(falhas_gerais, contestacao)
    falhas_gerais = _flag_fg_contestacao_externa(falhas_gerais, contestacao)

    return BRBDataBundle(
        na_demandas=sup["na_demandas"],
        na_falhas=na_falhas,
        treinamentos=sup["treinamentos"],
        treinamentos_horas=sup["treinamentos_horas"],
        falhas_gerais=falhas_gerais,
        contestacao=contestacao,
        auditados=eo["auditados"],
        inicio=inicio,
        fim=fim,
        fg_linhas_brutas=eo["fg_linhas_brutas"],
        fg_duplicatas=eo["fg_duplicatas"],
        contestacao_merge_stats=contestacao_merge_stats,
    )


def load_workbook(
    path: Path | None = None,
    inicio=None,
    fim=None,
    client_slug: str = "brb",
) -> BRBDataBundle:
    path = Path(path or EXCEL_PATH)
    client_slug = (client_slug or "brb").strip().lower()
    validate_workbook(path)
    client_cfg = get_client_config(client_slug)
    fg_sheet = client_cfg.get("fg_sheet", "Falhas_Gerais-BRB")

    na_demandas = _prep_na_demandas(
        pd.read_excel(path, "NA_Demandas"), inicio, fim, client_slug
    )
    na_falhas = _prep_na_falhas(read_na_falhas_excel(path), inicio, fim, client_slug)
    treinamentos = _prep_treinamentos(
        pd.read_excel(path, "Treinamentos"), inicio, fim, client_slug
    )
    xl_names = set(pd.ExcelFile(path).sheet_names)
    if "TreinamentosComHoras" in xl_names:
        treinamentos_horas = _prep_treinamentos_horas(
            pd.read_excel(path, "TreinamentosComHoras"), inicio, fim, client_slug
        )
    else:
        treinamentos_horas = pd.DataFrame()
    fg_raw = pd.read_excel(path, fg_sheet) if fg_sheet in xl_names else pd.DataFrame()
    if not fg_raw.empty and "Cliente" in fg_raw.columns:
        fg_raw = fg_raw[fg_raw["Cliente"].map(lambda v: match_client(v, client_slug))].copy()
    falhas_gerais, fg_linhas, fg_dup = _prep_falhas_gerais(fg_raw, inicio, fim)
    contestacao = _prep_contestacao(
        pd.read_excel(path, "Contestacao"), inicio, fim, client_slug
    )
    auditados = _prep_auditados(_read_auditados(path), inicio, fim, client_slug)

    na_falhas = _enrich_na_falhas(na_falhas, falhas_gerais)
    falhas_gerais = _join_conforme_fg(falhas_gerais, contestacao)
    falhas_gerais = _flag_fg_contestacao_externa(falhas_gerais, contestacao)

    return BRBDataBundle(
        na_demandas=na_demandas,
        na_falhas=na_falhas,
        treinamentos=treinamentos,
        treinamentos_horas=treinamentos_horas,
        falhas_gerais=falhas_gerais,
        contestacao=contestacao,
        auditados=auditados,
        inicio=inicio,
        fim=fim,
        fg_linhas_brutas=fg_linhas,
        fg_duplicatas=fg_dup,
    )


def _read_auditados(path: Path) -> pd.DataFrame:
    """Leitura enxuta da aba Auditados (só colunas necessárias)."""
    xl = pd.ExcelFile(path)
    if "Auditados" not in xl.sheet_names:
        return pd.DataFrame()
    available = set(xl.parse("Auditados", nrows=0).columns.astype(str))
    usecols = [c for c in _AUDITADOS_USECOLS if c in available]
    if not usecols:
        return pd.DataFrame()
    return pd.read_excel(path, "Auditados", usecols=usecols)


def _prep_auditados(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    if df is None or df.empty:
        return pd.DataFrame()
    out = df.copy()
    if "Cliente" in out.columns:
        out = out[out["Cliente"].map(lambda v: match_client(v, client_slug))].copy()
    if out.empty:
        return out
    if "Data" in out.columns:
        out["Data"] = parse_excel_date(out["Data"])
    if "Protocolo" in out.columns:
        out["_protocolo_norm"] = out["Protocolo"].map(norm_protocolo)
    else:
        out["_protocolo_norm"] = ""
    return filter_period(out, "Data", inicio, fim)


def _prep_na_demandas(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    if df.empty or "Cliente" not in df.columns:
        return pd.DataFrame()
    out = df[df["Cliente"].map(lambda v: match_client(v, client_slug))].copy()
    if out.empty:
        return out
    if "Data da Abertura" in out.columns:
        out["Data da Abertura"] = parse_excel_date(out["Data da Abertura"])
    if "Data do Retorno" in out.columns:
        out["Data do Retorno"] = parse_excel_date(out["Data do Retorno"])
    for c in ("Quantidade de Protolocos", "Falhas Manuais", "Falhas Processuais", "Falhas Automáticas"):
        if c in out.columns:
            out[c] = pd.to_numeric(out[c], errors="coerce").fillna(0).astype(int)
    return filter_demandas_period(out, inicio, fim)


def _norm_mes(value) -> str:
    s = safe_str(value).upper()
    for ch in ("Á", "À", "Â", "Ã"):
        s = s.replace(ch, "A")
    return s.replace("Ç", "C")


def filter_demandas_period(df: pd.DataFrame, inicio, fim) -> pd.DataFrame:
    """NA_Demandas: filtra por Data da Abertura e/ou coluna Mês dentro do intervalo."""
    if df.empty or (inicio is None and fim is None):
        return df
    start = pd.Timestamp(inicio) if inicio else pd.Timestamp("1900-01-01")
    end = pd.Timestamp(fim) if fim else pd.Timestamp.today().normalize() + pd.Timedelta(days=1) - pd.Timedelta(seconds=1)

    by_date = filter_period(df, "Data da Abertura", inicio, fim)

    def mes_no_periodo(mes_val) -> bool:
        mes_num = MES_PT.get(_norm_mes(mes_val))
        if not mes_num:
            return False
        for year in range(start.year, end.year + 1):
            month_start = pd.Timestamp(year=year, month=mes_num, day=1)
            month_end = month_start + pd.offsets.MonthEnd(0)
            if month_end >= start and month_start <= end:
                return True
        return False

    by_mes = df[df["Mês"].map(mes_no_periodo)] if "Mês" in df.columns else df.iloc[0:0]
    idx = by_date.index.union(by_mes.index)
    return df.loc[idx].copy()


def _prep_na_falhas(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    out = normalize_na_falhas_raw(df)
    if out.empty or "CLIENTE" not in out.columns:
        return pd.DataFrame()
    out = out[out["CLIENTE"].map(lambda v: match_client(v, client_slug))].copy()
    out["_protocolo_norm"] = out["PROTOCOLO"].map(norm_protocolo)
    for c in ("DATA DE CADASTRO", "DATA DE NOTIFICAÇÃO"):
        if c in out.columns:
            out[c] = parse_excel_date(out[c])
    out["possivel_ataque"] = (
        out["MOTIVO DA FALHA"].map(lambda x: is_possivel_ataque(x, ATTACK_KEYWORDS))
        | out["RESULTADO DA AUDITORIA"].map(lambda x: is_possivel_ataque(x, ATTACK_KEYWORDS))
    )
    out["notificacao_ativa"] = _notification_mask(out)
    out["divergencia_risco"] = _risk_change_mask(out)
    # Cadastro ∪ notificação: inclui falhas cadastradas antes, mas notificadas no período
    out = filter_period_any(
        out, ["DATA DE CADASTRO", "DATA DE NOTIFICAÇÃO"], inicio, fim
    )
    # A aba pode conter apontamentos que não mudam a classificação de risco.
    # Eles permanecem na fonte, mas não compõem os indicadores de notificação.
    return out.loc[out["notificacao_ativa"]].copy()


@lru_cache(maxsize=50_000)
def _training_title_client_slugs(title: str) -> tuple[str, ...]:
    from report_brb.client_registry import CLIENTS

    return tuple(slug for slug in CLIENTS if match_client(title, slug))


def _training_title_matches_client(value: object, client_slug: str) -> bool:
    explicit_clients = _training_title_client_slugs(str(value or ""))
    if client_slug in explicit_clients:
        return True
    return not explicit_clients


def _prepare_training_frames(raw: dict[str, pd.DataFrame]) -> None:
    """Normaliza treinamento uma vez por arquivo, antes do loop de clientes."""
    training = raw.get("treinamentos")
    if training is not None and not training.empty and not training.attrs.get("brb_prepared"):
        training = training.copy()
        col = "Event: EventTitle"
        training["scope_treinamento"] = training[col].map(training_scope)
        status_col = "Status" if "Status" in training.columns else None
        if status_col:
            training = training[~training[status_col].fillna("").astype(str).str.casefold().str.contains("cancel", na=False)].copy()
        training["tipo_treinamento"] = training[col].map(classificar_treinamento)
        for c in ("AssignmentDate", "SignatureDate", "Session: StartDate", "Session: FinalDate"):
            if c in training.columns:
                training[c] = parse_excel_date(training[c])
        training.attrs["brb_prepared"] = True
        raw["treinamentos"] = training

    hours = raw.get("treinamentos_horas")
    if hours is not None and not hours.empty and not hours.attrs.get("brb_prepared"):
        hours = hours.copy()
        col = "Event: EventTitle"
        hours["tipo_treinamento"] = hours[col].map(classificar_treinamento)
        hours["scope_treinamento"] = hours[col].map(training_scope)
        status_col = "SessionStatus" if "SessionStatus" in hours.columns else None
        if status_col:
            hours = hours[~hours[status_col].fillna("").astype(str).str.casefold().str.contains("cancel", na=False)].copy()
        for c in ("StartDate", "FinalDate", "RequestDate", "Event: EventDeadline"):
            if c in hours.columns:
                hours[c] = parse_excel_date(hours[c])
        dur_col = "Event: EstimatedDuration"
        hours["horas"] = hours[dur_col].map(parse_duration_hours) if dur_col in hours.columns else 0.0
        hours.loc[hours["tipo_treinamento"].eq("Newsletter"), "horas"] = 0.0
        hours.attrs["brb_prepared"] = True
        raw["treinamentos_horas"] = hours


def training_scope(value: object) -> str:
    title = str(value or "")
    return "cliente" if _training_title_client_slugs(title) else "corporativo"


def _prep_treinamentos(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    col = "Event: EventTitle"
    out = df[df[col].map(lambda value: _training_title_matches_client(value, client_slug))].copy()
    if "scope_treinamento" not in out.columns:
        out["scope_treinamento"] = out[col].map(training_scope)
    if "tipo_treinamento" not in out.columns:
        out["tipo_treinamento"] = out[col].map(classificar_treinamento)
    for c in ("AssignmentDate", "SignatureDate", "Session: StartDate", "Session: FinalDate"):
        if c in out.columns and not pd.api.types.is_datetime64_any_dtype(out[c]):
            out[c] = parse_excel_date(out[c])
    ref = "Session: StartDate" if "Session: StartDate" in out.columns else "AssignmentDate"
    return filter_period(out, ref, inicio, fim)


def parse_duration_hours(val) -> float:
    """Converte duração (timedelta / time / 'HH:MM:SS') em horas decimais."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return 0.0
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        # Excel raro: já em horas ou fração de dia
        v = float(val)
        if 0 < v < 1:
            return v * 24.0
        return v
    if hasattr(val, "total_seconds"):
        try:
            return max(0.0, float(val.total_seconds()) / 3600.0)
        except (TypeError, ValueError):
            return 0.0
    if hasattr(val, "hour") and hasattr(val, "minute"):
        try:
            return float(val.hour) + float(val.minute) / 60.0 + float(getattr(val, "second", 0) or 0) / 3600.0
        except (TypeError, ValueError):
            return 0.0
    s = str(val).strip()
    if not s or s.lower() in ("nan", "nat", "none"):
        return 0.0
    parts = s.split(":")
    try:
        if len(parts) == 3:
            return int(parts[0]) + int(parts[1]) / 60.0 + int(parts[2]) / 3600.0
        if len(parts) == 2:
            return int(parts[0]) + int(parts[1]) / 60.0
    except (TypeError, ValueError):
        return 0.0
    return 0.0


def classificar_treinamento(title: object) -> str:
    """Classifica o evento pelo título sem exigir que o cliente esteja no texto."""
    text = str(title or "").strip().casefold()
    if "reforço operacional" in text or "reforco operacional" in text:
        return "Reforço Operacional"
    for keyword, label in (
        ("follow-up", "Follow-up"),
        ("follow up", "Follow-up"),
        ("refresh", "Refresh"),
        ("workshop", "Workshop"),
        ("update", "Update"),
        ("newsletter", "Newsletter"),
        ("upgrade", "Upgrade"),
        ("dinâmica", "Dinâmica"),
        ("dinamica", "Dinâmica"),
    ):
        if keyword in text:
            return label
    return "Outros"


def _prep_treinamentos_horas(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    """Sessões com duração estimada (esforço) — filtro por EventTitle."""
    if df.empty or "Event: EventTitle" not in df.columns:
        return pd.DataFrame()
    out = df[df["Event: EventTitle"].map(lambda value: _training_title_matches_client(value, client_slug))].copy()
    status_col = "SessionStatus" if "SessionStatus" in out.columns else None
    if status_col:
        out = out[~out[status_col].fillna("").astype(str).str.casefold().str.contains("cancel", na=False)].copy()
    for c in ("StartDate", "FinalDate", "RequestDate", "Event: EventDeadline"):
        if c in out.columns:
            out[c] = parse_excel_date(out[c])
    dur_col = "Event: EstimatedDuration"
    if "horas" not in out.columns:
        out["horas"] = out[dur_col].map(parse_duration_hours) if dur_col in out.columns else 0.0
    if "tipo_treinamento" not in out.columns:
        out["tipo_treinamento"] = out["Event: EventTitle"].map(classificar_treinamento)
    if "scope_treinamento" not in out.columns:
        out["scope_treinamento"] = out["Event: EventTitle"].map(training_scope)
    out.loc[out["tipo_treinamento"].eq("Newsletter"), "horas"] = 0.0
    if "StartDate" in out.columns and "FinalDate" in out.columns:
        out["_periodo_ref"] = out["StartDate"].where(out["StartDate"].notna(), out["FinalDate"])
        out = filter_period(out, "_periodo_ref", inicio, fim)
    elif "StartDate" in out.columns:
        out = filter_period(out, "StartDate", inicio, fim)
    return out


def _prep_falhas_gerais(df: pd.DataFrame, inicio, fim) -> tuple[pd.DataFrame, int, int]:
    if df.empty:
        return add_case_columns(df, "Protocolo", "Matrícula Agente"), 0, 0
    out = add_case_columns(df, "Protocolo", "Matrícula Agente")
    for c in ("Data de Análise", "Data Auditoria"):
        if c in out.columns:
            out[c] = parse_excel_date(out[c])
    cen = out.get("Novo cenário", out.get("Cenário", pd.Series(dtype=str)))
    out["fn_fp"] = cen.map(classificar_fn_fp)
    filtered = filter_period(out, "Data de Análise", inicio, fim)
    linhas_brutas = len(filtered)
    deduped = dedupe_by_case_key(filtered, "Data de Análise")
    duplicatas = max(0, linhas_brutas - len(deduped))
    return deduped, linhas_brutas, duplicatas


def _prep_contestacao(df: pd.DataFrame, inicio, fim, client_slug: str = "brb") -> pd.DataFrame:
    out = df[df["Cliente"].map(lambda v: match_client(v, client_slug))].copy()
    out = add_case_columns(out, "Protocolo", "Matrícula")
    out["classificacao_conforme"] = out["CONFORME"].map(classificar_conforme)
    for c in ("Data de Análise", "Data de Cadastro", "Data de Conclusão", "Data"):
        if c in out.columns:
            out[c] = parse_excel_date(out[c])
    out = dedupe_by_case_key(out, "Data de Análise")
    # O período da Contestação é definido pela data de recebimento (coluna Data).
    # Data de Análise identifica quando ocorreu a análise que foi contestada.
    return filter_period(out, "Data", inicio, fim)


def merge_contestacao_eo_with_excel(
    eo_cont: pd.DataFrame,
    excel_raw: pd.DataFrame | None,
    *,
    batimento_keys: set[str] | None = None,
    inicio=None,
    fim=None,
    client_slug: str = "brb",
) -> tuple[pd.DataFrame, dict[str, int]]:
    """Acrescenta linhas da planilha Contestação que não existem na EO/TSV (batimento)."""
    stats = {
        "eo_rows": len(eo_cont) if eo_cont is not None else 0,
        "excel_rows_raw": 0,
        "excel_appended": 0,
        "excel_skipped_batimento": 0,
    }
    if excel_raw is None or excel_raw.empty:
        return (eo_cont if eo_cont is not None else pd.DataFrame()), stats

    excel_prep = _prep_contestacao(excel_raw, inicio, fim, client_slug)
    stats["excel_rows_raw"] = len(excel_prep)
    if excel_prep.empty:
        return (eo_cont if eo_cont is not None else pd.DataFrame()), stats

    existing = batimento_keys if batimento_keys is not None else contestacao_batimento_keys_from_df(eo_cont)
    append_rows: list[pd.Series] = []
    for _, row in excel_prep.iterrows():
        key = contestacao_batimento_key_from_row(row)
        if not key.replace("|", ""):
            stats["excel_skipped_batimento"] += 1
            continue
        if key in existing:
            stats["excel_skipped_batimento"] += 1
            continue
        existing.add(key)
        append_rows.append(row)

    stats["excel_appended"] = len(append_rows)
    base = eo_cont.copy() if eo_cont is not None and not eo_cont.empty else pd.DataFrame()
    if not base.empty and "_contestacao_source" not in base.columns:
        base = base.copy()
        base["_contestacao_source"] = "eo"
    if not append_rows:
        return base, stats

    extra = pd.DataFrame(append_rows)
    extra["_contestacao_source"] = "excel_supplement"
    merged = pd.concat([base, extra], ignore_index=True, sort=False)
    merged = dedupe_by_case_key(merged, "Data de Análise")
    return merged, stats


def _find_column(df: pd.DataFrame, *tokens: str) -> str | None:
    from report_brb.brb_normalize import strip_accents

    for col in df.columns:
        norm = strip_accents(safe_str(col)).lower()
        if all(t.lower() in norm for t in tokens):
            return col
    return None


def _is_contestacao_externa(val) -> bool:
    from report_brb.brb_normalize import strip_accents

    s = strip_accents(safe_str(val)).lower()
    return "contest" in s and "extern" in s


def _flag_fg_contestacao_externa(fg: pd.DataFrame, cont: pd.DataFrame) -> pd.DataFrame:
    """Marca FG Contestação Externa cujo protocolo já existe na aba Contestacao."""
    if fg.empty:
        return fg
    out = fg.copy()
    tipo_an_col = _find_column(out, "tipo", "anal")
    cont_protos = (
        set(cont["_protocolo_norm"]) - {""} if not cont.empty and "_protocolo_norm" in cont.columns else set()
    )
    if tipo_an_col:
        out["contestacao_externa_fg"] = out[tipo_an_col].map(_is_contestacao_externa)
    else:
        out["contestacao_externa_fg"] = False
    out["duplicado_contestacao"] = (
        out["contestacao_externa_fg"].fillna(False).astype(bool) & out["_protocolo_norm"].isin(cont_protos)
    )
    return out


def _enrich_na_falhas(na: pd.DataFrame, fg: pd.DataFrame) -> pd.DataFrame:
    if na.empty:
        return na
    if fg.empty:
        out = na.copy()
        out["matricula_fg"] = ""
        out["sem_matricula_fg"] = True
        return out
    # 1 protocolo → 1 matrícula FG (evita duplicar linhas da NA no merge 1:N)
    lookup = (
        fg[["_protocolo_norm", "_matricula_norm", "Nome Agente"]]
        .dropna(subset=["_protocolo_norm"])
        .drop_duplicates(subset=["_protocolo_norm"], keep="first")
    )
    merged = na.merge(
        lookup,
        left_on="_protocolo_norm",
        right_on="_protocolo_norm",
        how="left",
        suffixes=("", "_fg"),
    )
    merged["matricula_fg"] = merged["_matricula_norm"]
    merged["sem_matricula_fg"] = merged["matricula_fg"].isna() | (merged["matricula_fg"] == "")
    return merged


def _join_conforme_fg(fg: pd.DataFrame, cont: pd.DataFrame) -> pd.DataFrame:
    if fg.empty:
        return fg
    if cont.empty:
        fg = fg.copy()
        fg["CONFORME"] = ""
        fg["classificacao_conforme"] = "sem_avaliacao"
        return fg
    sub = cont[
        ["chave_caso", "CONFORME", "classificacao_conforme", "Cenário", "Origem"]
    ].drop_duplicates(subset=["chave_caso"], keep="last")
    sub = sub.rename(columns={"Cenário": "Cenario_Contestacao"})
    out = fg.merge(sub, on="chave_caso", how="left")
    out["classificacao_conforme"] = out["classificacao_conforme"].fillna("sem_avaliacao")
    return out


_CLIENT_COLUMN_CANDIDATES = (
    ("NA_Demandas", "Cliente"),
    ("NA_Falhas", "CLIENTE"),
    ("Falhas", "Cliente"),
    ("Contestacao", "Cliente"),
    ("Treinamentos", "Event: EventTitle"),
    ("TreinamentosComHoras", "Event: EventTitle"),
)

# Auditados omitido no preview — 150k+ linhas; detecção BRB vem das demais abas.
_PREVIEW_CLIENT_ROW_CAP = 15_000


def _sheet_row_counts(path: Path, sheet_names: list[str]) -> dict[str, int]:
    """Contagem rápida de linhas (sem carregar células via pandas)."""
    from openpyxl import load_workbook

    counts: dict[str, int] = {}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for name in sheet_names:
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            max_row = ws.max_row or 0
            counts[name] = max(0, max_row - 1)
    finally:
        wb.close()
    return counts


def _read_client_column_preview(path: Path, sheet: str, col: str) -> pd.Series:
    """Amostra colunas de cliente para preview (evita ler Auditados inteiro)."""
    return pd.read_excel(path, sheet, usecols=[col], nrows=_PREVIEW_CLIENT_ROW_CAP)[col]


def preview_workbook(path: Path) -> dict:
    """Valida planilha e retorna contagens por aba + clientes detectados."""
    from report_brb.client_registry import CLIENTS, list_clients

    path = Path(path)
    validate_workbook(path)
    xl = pd.ExcelFile(path)
    count_sheets = [s for s in EXPECTED_SHEETS if s in xl.sheet_names]
    if "Falhas" in xl.sheet_names and "NA_Falhas" not in xl.sheet_names:
        count_sheets.append("Falhas")
    counts_raw = _sheet_row_counts(path, count_sheets)
    sheets = {s: counts_raw[s] for s in EXPECTED_SHEETS if s in counts_raw}
    if "NA_Falhas" not in sheets and "Falhas" in counts_raw:
        sheets["NA_Falhas"] = counts_raw["Falhas"]

    detected: dict[str, int] = {}
    for sheet, col in _CLIENT_COLUMN_CANDIDATES:
        if sheet not in xl.sheet_names:
            continue
        try:
            series = _read_client_column_preview(path, sheet, col)
        except (KeyError, ValueError):
            continue
        for val, count in series.value_counts(dropna=True).items():
            label = safe_str(val)
            if not label:
                continue
            detected[label] = detected.get(label, 0) + int(count)

    registry_hits = []
    for slug, cfg in CLIENTS.items():
        hits = sum(
            count for label, count in detected.items() if match_client(label, slug)
        )
        if hits:
            registry_hits.append(
                {
                    "slug": slug,
                    "nome": cfg.get("nome_curto", slug),
                    "enabled": bool(cfg.get("enabled", True)),
                    "rows_matched": hits,
                }
            )

    return {
        "filename": path.name,
        "sheets": sheets,
        "extra_sheets": [s for s in xl.sheet_names if s not in EXPECTED_SHEETS],
        "clients_registry": list_clients(enabled_only=False),
        "clients_detected": sorted(
            registry_hits, key=lambda x: (-x["rows_matched"], x["nome"])
        ),
        "client_values_sample": sorted(
            detected.items(), key=lambda x: -x[1]
        )[:25],
    }
