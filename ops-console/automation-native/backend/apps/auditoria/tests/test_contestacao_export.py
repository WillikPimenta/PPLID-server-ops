from __future__ import annotations

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.auditoria.models import AuditoriaAtividadeProtocolo, AuditoriaAtividadeProtocoloEtapa
from apps.auditoria.services.atividade_import import confirm_import, create_import_staging, preview_import
from apps.auditoria.services.contestacao_export import (
    RETORNO_TEMPLATE_PATH,
    build_retorno_filename,
    derive_retorno_values,
)
from apps.auditoria.tests.test_atividade_import import build_sample_workbook

User = get_user_model()


@override_settings(ACCESS_ENFORCEMENT=False)
class ContestacaoRetornoExportTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(
            username="retorno_user",
            email="retorno@test.local",
            password="test12345",
        )
        self.client.force_authenticate(user=self.user)
        file_bytes = build_sample_workbook()
        preview = preview_import(file_bytes, "QI-8279 - PICPAY - 01072026.xlsx")
        staging = create_import_staging(
            user=self.user,
            file_bytes=file_bytes,
            filename="QI-8279 - PICPAY - 01072026.xlsx",
            preview=preview,
        )
        self.atividade = confirm_import(
            staging=staging,
            user=self.user,
            nome="QI-8279 - PICPAY - 01072026",
            data_recepcao=date(2026, 7, 1),
        )
        self.protocolo = self.atividade.protocolos.order_by("excel_row").first()

    def test_template_exists_in_repo(self):
        self.assertTrue(RETORNO_TEMPLATE_PATH.exists())

    def test_build_retorno_filename(self):
        self.assertEqual(
            build_retorno_filename("QI-8279 - PICPAY - 01072026.xlsx"),
            "QI-8279 - PICPAY - 01072026_Retorno_IDF.xlsx",
        )
        self.assertEqual(
            build_retorno_filename("QI-8279 - PICPAY - 01072026"),
            "QI-8279 - PICPAY - 01072026_Retorno_IDF.xlsx",
        )
        self.assertEqual(
            build_retorno_filename("QI-8279 - PICPAY - 01072026_Retorno_IDF"),
            "QI-8279 - PICPAY - 01072026_Retorno_IDF.xlsx",
        )

    def test_export_uses_atividade_nome_in_filename(self):
        self.atividade.nome = "QI-8279 - PICPAY - 01072026"
        self.atividade.nome_arquivo_original = "arquivo_importado_diferente.xlsx"
        self.atividade.save(update_fields=["nome", "nome_arquivo_original"])

        response = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/retorno.xlsx"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "QI-8279 - PICPAY - 01072026_Retorno_IDF.xlsx",
            response["Content-Disposition"],
        )

    def test_export_uses_blank_template_and_fills_rows(self):
        save = self.client.patch(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/protocolos/{self.protocolo.id}/analise/",
            {
                "consideracoes_finais": "Documento revisado na intranet.",
                "tipo_conclusao": "Manual",
                "finalizar": True,
                "etapas": [
                    {
                        "resultado_correto": "",
                        "tipo_falha": "Automático",
                        "etapa_falha": "DIGITALIZAÇÃO",
                        "situacao": "improcedente",
                    }
                ],
            },
            format="json",
        )
        self.assertEqual(save.status_code, 200)

        response = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/retorno.xlsx"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "QI-8279 - PICPAY - 01072026_Retorno_IDF.xlsx",
            response["Content-Disposition"],
        )

        workbook = load_workbook(filename=BytesIO(response.content), data_only=True)
        sheet = workbook["Contestação de Resultado"]
        self.assertEqual(sheet["B9"].value, "PROTOCOLO")
        self.assertEqual(sheet["G9"].value, "RESULTADO PÓS AUDITORIA")
        self.assertEqual(str(sheet["B10"].value), "123456")
        self.assertEqual(sheet["C10"].value, "Risk Manager - PICPAY")
        self.assertEqual(sheet["F10"].value, "SEM RISCO APARENTE")
        self.assertEqual(sheet["G10"].value, "SEM RISCO APARENTE")
        self.assertEqual(sheet["H10"].value, "MANUAL")
        self.assertEqual(sheet["I10"].value, "SEM FALHA")
        self.assertEqual(sheet["K10"].value, "Documento revisado na intranet.")
        self.assertEqual(sheet["L10"].value, "IMPROCEDENTE")
        # Segundo protocolo do import também entra no template.
        self.assertEqual(str(sheet["B11"].value), "123457")
        # Bloco CONSIDERAÇÕES permanece no padrão do template.
        self.assertEqual(sheet["B15"].value, "CONSIDERAÇÕES")

    def test_export_expands_table_formatting_without_merging_data_rows(self):
        for index in range(7):
            AuditoriaAtividadeProtocolo.objects.create(
                atividade=self.atividade,
                protocolo=f"90000{index}",
                workflow="Risk Manager",
                nivel_hierarquico="Integrador",
                resultado_contestado="SEM RISCO APARENTE",
                conclusao_contestacao="IMPROCEDENTE",
                excel_row=12 + index,
            )

        response = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/retorno.xlsx"
        )
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(filename=BytesIO(response.content))
        sheet = workbook["Contestação de Resultado"]
        merged_ranges = list(sheet.merged_cells.ranges)

        self.assertIn("B20:L20", {str(cell_range) for cell_range in merged_ranges})
        self.assertFalse(
            any(
                cell_range.min_row <= 18 and cell_range.max_row >= 10
                for cell_range in merged_ranges
            )
        )
        self.assertIn(
            "L10:L18",
            {str(conditional.sqref) for conditional in sheet.conditional_formatting},
        )

        for row_number in range(9, 19):
            self.assertIsNone(sheet.row_dimensions[row_number].height)
            for col_number in range(2, 13):
                cell = sheet.cell(row=row_number, column=col_number)
                self.assertEqual(cell.alignment.horizontal, "left")
                self.assertTrue(cell.alignment.wrap_text)

    def test_export_preserves_import_order_and_fills_matricula(self):
        protocolos = list(self.atividade.protocolos.order_by("excel_row", "id"))
        protocolos[0].protocolo = "999999"
        protocolos[0].save(update_fields=["protocolo", "updated_at"])
        protocolos[1].protocolo = "111111"
        protocolos[1].save(update_fields=["protocolo", "updated_at"])
        AuditoriaAtividadeProtocoloEtapa.objects.create(
            protocolo=protocolos[0],
            ordem=0,
            agente="c12345a",
            tipo_falha="Colaborador",
            situacao="procedente",
        )

        response = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/retorno.xlsx"
        )
        self.assertEqual(response.status_code, 200)

        workbook = load_workbook(filename=BytesIO(response.content), data_only=True)
        sheet = workbook["Contestação de Resultado"]
        self.assertEqual(sheet["M9"].value, "MATRÍCULA")
        self.assertEqual(str(sheet["B10"].value), "999999")
        self.assertEqual(sheet["M10"].value, "c12345a")
        self.assertEqual(str(sheet["B11"].value), "111111")

    def test_derive_procedente_values(self):
        response = self.client.patch(
            f"/api/v1/qualidade/auditoria/atividades/{self.atividade.id}/protocolos/{self.protocolo.id}/analise/",
            {
                "consideracoes_finais": "CPF divergente.",
                "tipo_conclusao": "Manual",
                "etapas": [
                    {
                        "resultado_correto": "COM RISCO - CPF DIVERGENTE",
                        "tipo_falha": "Processual",
                        "etapa_falha": "AUDITORIA",
                        "agente": "c12345a",
                        "situacao": "procedente",
                        "motivo_falha": "NÃO SINALIZADO - CPF DIVERGENTE",
                    }
                ],
            },
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        self.protocolo.refresh_from_db()
        values = derive_retorno_values(self.protocolo)
        self.assertEqual(values["resultado_pos_auditoria"], "COM RISCO - CPF DIVERGENTE")
        self.assertEqual(values["tipo_conclusao"], "MANUAL")
        self.assertEqual(values["tipo_falha"], "PROCESSUAL")
        self.assertEqual(values["cenario"], "NÃO SINALIZADO - CPF DIVERGENTE")
        self.assertEqual(values["detalhamento"], "CPF divergente.")
        self.assertEqual(values["conclusao_contestacao"], "PROCEDENTE")
