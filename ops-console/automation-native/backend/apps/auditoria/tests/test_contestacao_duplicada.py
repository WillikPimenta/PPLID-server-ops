from __future__ import annotations

from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.auditoria.models import AuditoriaAtividade, AuditoriaAtividadeProtocolo, AuditoriaFalhaCadastro
from apps.auditoria.services.atividade_import import confirm_import, create_import_staging, preview_import
from apps.auditoria.services.contestacao_export import build_row_values
from apps.auditoria.services.contestacao_duplicada import (
    contestacao_row_identity,
    lookup_prior_tratados_contestacao,
)
from apps.auditoria.services.contestacao_import import ParsedProtocolRow
from apps.auditoria.tests.test_atividade_import import build_sample_workbook

User = get_user_model()


def build_workbook_with_rows(rows: list[dict[str, str]]) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contestação de Resultado"
    sheet["B9"] = "PROTOCOLO"
    sheet["C9"] = "WORKFLOW"
    sheet["D9"] = "NÍVEL HIERÁRQUICO"
    sheet["F9"] = "RESULTADO CONTESTADO"
    for index, row in enumerate(rows, start=10):
        sheet[f"B{index}"] = row.get("protocolo", "")
        sheet[f"C{index}"] = row.get("workflow", "Risk Manager - PICPAY")
        sheet[f"D{index}"] = row.get("nivel_hierarquico", "Risk Manager - PICPAY - Digitalizador")
        sheet[f"F{index}"] = row.get("resultado_contestado", "")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _import_and_finalize_protocol(
    client: APIClient,
    user,
    *,
    protocolo: str = "123456",
    link_demanda: str = "https://example.com/demanda/original",
) -> tuple[AuditoriaAtividade, AuditoriaAtividadeProtocolo]:
    workbook = build_sample_workbook()
    preview = preview_import(workbook, "QI-ORIG - PICPAY - 01072026.xlsx")
    staging = create_import_staging(
        user=user,
        file_bytes=workbook,
        filename="QI-ORIG - PICPAY - 01072026.xlsx",
        preview=preview,
    )
    atividade = confirm_import(
        staging=staging,
        user=user,
        nome="Atividade original",
        link_demanda=link_demanda,
        data_recepcao=date(2026, 7, 1),
    )
    protocolo_obj = atividade.protocolos.filter(protocolo=protocolo).first()
    assert protocolo_obj is not None

    save = client.patch(
        f"/api/v1/qualidade/auditoria/atividades/{atividade.id}/protocolos/{protocolo_obj.id}/analise/",
        {
            "consideracoes_finais": "Análise original concluída.",
            "tipo_conclusao": "Manual",
            "finalizar": True,
            "etapas": [
                {
                    "resultado_correto": "SEM RISCO APARENTE",
                    "tipo_falha": "Automático",
                    "etapa_falha": "DIGITALIZAÇÃO",
                    "situacao": "improcedente",
                }
            ],
        },
        format="json",
    )
    assert save.status_code == 200, save.data
    protocolo_obj.refresh_from_db()
    return atividade, protocolo_obj


@override_settings(ACCESS_ENFORCEMENT=False)
class ContestacaoDuplicadaTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_superuser(
            username="duplicado_user",
            email="duplicado@test.local",
            password="test12345",
        )
        self.client.force_authenticate(user=self.user)
        self.atividade_origem, self.protocolo_origem = _import_and_finalize_protocol(
            self.client,
            self.user,
            link_demanda="https://example.com/demanda/original",
        )
        self.tratado = AuditoriaFalhaCadastro.objects.get(pk=self.protocolo_origem.tratado_id)

    def test_preview_flags_protocolo_ja_tratado(self):
        preview = preview_import(build_sample_workbook(), "QI-8279 - PICPAY - 01072026.xlsx")
        self.assertEqual(len(preview.protocolos_ja_tratados), 1)
        self.assertEqual(preview.protocolos_ja_tratados[0]["protocolo"], "123456")
        self.assertEqual(preview.protocolos_ja_tratados[0]["tratado_id"], self.tratado.id)
        self.assertEqual(
            preview.protocolos_ja_tratados[0]["link_demanda"],
            "https://example.com/demanda/original",
        )

    def test_confirm_import_espelha_protocolo_duplicado(self):
        file_bytes = build_sample_workbook()
        preview = preview_import(file_bytes, "QI-8279 - PICPAY - 01072026.xlsx")
        staging = create_import_staging(
            user=self.user,
            file_bytes=file_bytes,
            filename="QI-8279 - PICPAY - 01072026.xlsx",
            preview=preview,
        )
        atividade = confirm_import(
            staging=staging,
            user=self.user,
            nome="Atividade duplicada",
            link_demanda="https://example.com/demanda/nova",
            data_recepcao=date(2026, 7, 2),
        )

        espelhado = atividade.protocolos.get(protocolo="123456")
        novo = atividade.protocolos.get(protocolo="123457")
        self.assertEqual(espelhado.tratado_referencia_id, self.tratado.id)
        self.assertEqual(espelhado.status, AuditoriaAtividadeProtocolo.STATUS_CONCLUIDO)
        self.assertEqual(novo.status, AuditoriaAtividadeProtocolo.STATUS_PENDENTE)
        self.assertEqual(espelhado.etapas.count(), 1)
        self.assertEqual(espelhado.consideracoes_finais, "Análise original concluída.")
        self.assertIsNone(espelhado.tratado_id)

        detail = self.client.get(f"/api/v1/qualidade/auditoria/atividades/{atividade.id}/")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual(detail.data["protocolos_tratamento_pendentes"], 1)
        self.assertEqual(detail.data["protocolos_espelhados"], 1)

        list_response = self.client.get("/api/v1/qualidade/auditoria/atividades/")
        self.assertEqual(list_response.status_code, 200)
        row = next(item for item in list_response.data["results"] if item["id"] == atividade.id)
        self.assertEqual(row["protocolos_espelhados"], 1)

        protocolos = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{atividade.id}/protocolos/"
        )
        self.assertEqual(protocolos.status_code, 200)
        espelhado_payload = next(item for item in protocolos.data["results"] if item["protocolo"] == "123456")
        self.assertTrue(espelhado_payload["espelhado"])
        self.assertEqual(espelhado_payload["tratado_referencia_id"], self.tratado.id)

    def test_export_includes_duplicado_note_in_detalhamento(self):
        file_bytes = build_sample_workbook()
        preview = preview_import(file_bytes, "QI-8279 - PICPAY - 01072026.xlsx")
        staging = create_import_staging(
            user=self.user,
            file_bytes=file_bytes,
            filename="QI-8279 - PICPAY - 01072026.xlsx",
            preview=preview,
        )
        atividade = confirm_import(
            staging=staging,
            user=self.user,
            nome="Atividade duplicada export",
            data_recepcao=date(2026, 7, 2),
        )
        espelhado = (
            AuditoriaAtividadeProtocolo.objects.select_related("tratado_referencia", "tratado_referencia__atividade")
            .prefetch_related("etapas")
            .get(atividade=atividade, protocolo="123456")
        )
        row = build_row_values(espelhado)
        self.assertIn("Protocolo duplicado", row["detalhamento"])
        self.assertIn("https://example.com/demanda/original", row["detalhamento"])
        self.assertEqual(row["conclusao_contestacao"], "IMPROCEDENTE")

        response = self.client.get(
            f"/api/v1/qualidade/auditoria/atividades/{atividade.id}/retorno.xlsx"
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(filename=BytesIO(response.content), data_only=True)
        sheet = workbook["Contestação de Resultado"]
        detalhamento = str(sheet["K10"].value or "")
        self.assertIn("Protocolo duplicado", detalhamento)
        self.assertIn("https://example.com/demanda/original", detalhamento)

    def test_patch_analise_blocked_for_espelhado(self):
        file_bytes = build_sample_workbook()
        preview = preview_import(file_bytes, "QI-8279 - PICPAY - 01072026.xlsx")
        staging = create_import_staging(
            user=self.user,
            file_bytes=file_bytes,
            filename="QI-8279 - PICPAY - 01072026.xlsx",
            preview=preview,
        )
        atividade = confirm_import(
            staging=staging,
            user=self.user,
            nome="Atividade bloqueio espelho",
            data_recepcao=date(2026, 7, 2),
        )
        espelhado = atividade.protocolos.get(protocolo="123456")
        response = self.client.patch(
            f"/api/v1/qualidade/auditoria/atividades/{atividade.id}/protocolos/{espelhado.id}/analise/",
            {"finalizar": True, "etapas": []},
            format="json",
        )
        self.assertEqual(response.status_code, 409)
        self.assertIn("duplicado", response.data["detail"].lower())

    def test_does_not_espelhar_when_resultado_contestado_differs(self):
        workbook = build_workbook_with_rows(
            [
                {
                    "protocolo": "123456",
                    "resultado_contestado": "DOCUMENTO AUSENTE",
                },
                {
                    "protocolo": "123457",
                    "resultado_contestado": "SEM RISCO APARENTE",
                },
            ]
        )
        preview = preview_import(workbook, "QI-diff - PICPAY - 01072026.xlsx")
        self.assertEqual(preview.protocolos_ja_tratados, [])

        staging = create_import_staging(
            user=self.user,
            file_bytes=workbook,
            filename="QI-diff - PICPAY - 01072026.xlsx",
            preview=preview,
        )
        atividade = confirm_import(
            staging=staging,
            user=self.user,
            nome="Atividade resultado diferente",
            data_recepcao=date(2026, 7, 3),
        )
        protocolo = atividade.protocolos.get(protocolo="123456")
        self.assertIsNone(protocolo.tratado_referencia_id)
        self.assertEqual(protocolo.status, AuditoriaAtividadeProtocolo.STATUS_PENDENTE)

    def test_matches_case_insensitive_workflow_and_resultado(self):
        rows = [
            ParsedProtocolRow(
                protocolo="123456",
                workflow="risk manager - picpay",
                nivel_hierarquico="Risk Manager - PICPAY - Digitalizador",
                resultado_contestado="sem risco aparente",
                excel_row=10,
            )
        ]
        prior_map = lookup_prior_tratados_contestacao(rows)
        identity = contestacao_row_identity("123456", "risk manager - picpay", "sem risco aparente")
        self.assertIn(identity, prior_map)
        self.assertEqual(prior_map[identity].id, self.tratado.id)
