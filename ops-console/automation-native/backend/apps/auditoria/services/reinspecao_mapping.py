from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

MAPPING_SHEET = "Matriz completa"
MAPPING_MAX_BYTES = 5 * 1024 * 1024
MAPPING_HEADERS = (
    "IC/CO",
    "Código",
    "Irregularidades de Inspeção",
    "Ilha de atendimento",
    "Etapa de atendimento",
    "Tipo de Documento",
    "Categoria da falha. Crítica / Procedimento",
    "O que pode ser o motivo da falha",
)

STATUS_MATCHED = "matched"
STATUS_UNMATCHED = "unmatched"
STATUS_AMBIGUOUS = "ambiguous"
STATUS_PREEXISTING = "preexisting"
STATUS_CONFLICT = "conflict"

_ZERO_WIDTH = "\u200b\u200c\u200d\ufeff"
_DASHES = r"\-\u2010\u2011\u2012\u2013\u2014\u2212"
_IRREGULARITY_PATTERN = re.compile(
    rf"^\s*(?P<classification>IC|CO|Corre[cç][aã]o)\s*[{_DASHES}]\s*"
    rf"(?P<code>[0-9]+|Corre[cç][aã]o)\s*[{_DASHES}]\s*"
    r"(?P<description>.+?)\s*$",
    re.IGNORECASE,
)


def normalize_display_text(value: Any) -> str:
    """Limpa controles invisíveis sem destruir o texto canônico de exibição."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = text.replace("\u00a0", " ")
    text = text.translate({ord(char): None for char in _ZERO_WIDTH})
    return re.sub(r"\s+", " ", text).strip()


def normalize_comparison_text(value: Any) -> str:
    return normalize_display_text(value).casefold()


def normalize_code(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = str(int(value))
    elif isinstance(value, int):
        value = str(value)
    return normalize_comparison_text(value)


@dataclass(frozen=True)
class ExtractedIrregularity:
    classification: str = ""
    code: str = ""
    normalized_code: str = ""
    description: str = ""
    normalized_description: str = ""
    parsed: bool = False


def extract_irregularity_key(value: Any) -> ExtractedIrregularity:
    text = normalize_display_text(value)
    match = _IRREGULARITY_PATTERN.match(text)
    if match is None:
        return ExtractedIrregularity(description=text)
    code = normalize_display_text(match.group("code"))
    description = normalize_display_text(match.group("description"))
    return ExtractedIrregularity(
        classification=match.group("classification").upper(),
        code=code,
        normalized_code=normalize_code(code),
        description=description,
        normalized_description=normalize_comparison_text(description),
        parsed=True,
    )


@dataclass(frozen=True)
class MappingVariant:
    source_row: int
    classification: str
    code_original: str
    code_normalized: str
    description_original: str
    description_normalized: str
    island: str
    stage: str
    document_type: str
    category: str
    scenario: str
    mapping_version: str
    source_hash: str

    @classmethod
    def from_model(cls, row) -> "MappingVariant":
        return cls(
            source_row=row.source_row,
            classification=row.classificacao,
            code_original=row.codigo_original,
            code_normalized=row.codigo_normalizado,
            description_original=row.descricao_original,
            description_normalized=row.descricao_normalizada,
            island=row.ilha,
            stage=row.etapa,
            document_type=row.tipo_documento,
            category=row.categoria,
            scenario=row.cenario,
            mapping_version=row.mapping_version,
            source_hash=row.source_hash,
        )


@dataclass(frozen=True)
class MappingResolution:
    extracted: ExtractedIrregularity
    scenario: str = ""
    stage: str = ""
    scenario_status: str = STATUS_UNMATCHED
    stage_status: str = STATUS_UNMATCHED
    mapping_version: str = ""
    source_hash: str = ""
    candidates: tuple[MappingVariant, ...] = ()

    def diagnostic_candidates(self) -> list[dict[str, Any]]:
        return [
            {
                "source_row": row.source_row,
                "classification": row.classification,
                "code": row.code_original,
                "description": row.description_original,
                "scenario": row.scenario,
                "stage": row.stage,
                "document_type": row.document_type,
            }
            for row in self.candidates
        ]


class ReinspecaoMappingResolver:
    """Resolvedor em memória; a matriz é carregada uma vez por lote/request."""

    def __init__(self, variants: Iterable[MappingVariant]):
        self.variants = tuple(variants)
        self.by_code: dict[str, list[MappingVariant]] = {}
        for row in self.variants:
            self.by_code.setdefault(row.code_normalized, []).append(row)
        versions = {row.mapping_version for row in self.variants if row.mapping_version}
        hashes = {row.source_hash for row in self.variants if row.source_hash}
        if len(versions) > 1 or len(hashes) > 1:
            raise ValueError("Há mais de uma versão/hash ativa da matriz de Reinspeção.")
        self.mapping_version = next(iter(versions), "")
        self.source_hash = next(iter(hashes), "")

    @classmethod
    def from_active_mapping(cls) -> "ReinspecaoMappingResolver":
        from apps.auditoria.models import ReinspecaoIrregularidadeMapping

        rows = ReinspecaoIrregularidadeMapping.objects.filter(
            active=True,
            review_status=ReinspecaoIrregularidadeMapping.REVIEW_APPROVED,
        ).order_by("source_row")
        return cls(MappingVariant.from_model(row) for row in rows)

    def resolve(self, value: Any) -> MappingResolution:
        extracted = extract_irregularity_key(value)
        if not extracted.parsed:
            return MappingResolution(extracted=extracted)

        candidates = list(self.by_code.get(extracted.normalized_code, ()))
        if extracted.normalized_code == normalize_code("Correção"):
            candidates = [
                row
                for row in candidates
                if row.description_normalized == extracted.normalized_description
                and (
                    row.classification.upper() not in {"IC", "CO"}
                    or row.classification.upper() == extracted.classification
                )
            ]
        if not candidates:
            return MappingResolution(
                extracted=extracted,
                mapping_version=self.mapping_version,
                source_hash=self.source_hash,
            )

        scenario_values = sorted({row.scenario for row in candidates if row.scenario})
        stage_values = sorted({row.stage for row in candidates if row.stage})
        scenario_status = (
            STATUS_MATCHED
            if len(scenario_values) == 1
            else STATUS_AMBIGUOUS
            if len(scenario_values) > 1
            else STATUS_UNMATCHED
        )
        stage_status = (
            STATUS_MATCHED
            if len(stage_values) == 1
            else STATUS_AMBIGUOUS
            if len(stage_values) > 1
            else STATUS_UNMATCHED
        )
        return MappingResolution(
            extracted=extracted,
            scenario=scenario_values[0] if scenario_status == STATUS_MATCHED else "",
            stage=stage_values[0] if stage_status == STATUS_MATCHED else "",
            scenario_status=scenario_status,
            stage_status=stage_status,
            mapping_version=self.mapping_version,
            source_hash=self.source_hash,
            candidates=tuple(candidates),
        )


def status_for_existing(existing: Any, resolved: str, matrix_status: str) -> str:
    existing_text = normalize_display_text(existing)
    if not existing_text:
        return matrix_status
    if not resolved:
        return STATUS_PREEXISTING
    if existing_text == normalize_display_text(resolved):
        return STATUS_PREEXISTING
    return STATUS_CONFLICT


@dataclass(frozen=True)
class MappingDataset:
    mapping_version: str
    source_hash: str
    sheet_name: str
    source_range: str
    headers: tuple[str, ...]
    variants: tuple[MappingVariant, ...]
    preview_token: str
    source_label: str = ""
    warnings: tuple[str, ...] = ()

    def stats(self) -> dict[str, Any]:
        by_code: dict[str, list[MappingVariant]] = {}
        for row in self.variants:
            by_code.setdefault(row.code_normalized, []).append(row)
        duplicate_codes: list[dict[str, Any]] = []
        for code, rows in sorted(by_code.items()):
            if len(rows) < 2:
                continue
            scenarios = sorted({row.scenario for row in rows if row.scenario})
            stages = sorted({row.stage for row in rows if row.stage})
            duplicate_codes.append(
                {
                    "code": rows[0].code_original,
                    "rows": [row.source_row for row in rows],
                    "scenario_status": STATUS_AMBIGUOUS if len(scenarios) > 1 else STATUS_MATCHED,
                    "stage_status": STATUS_AMBIGUOUS if len(stages) > 1 else STATUS_MATCHED,
                    "scenario_values": scenarios,
                    "stage_values": stages,
                }
            )
        numeric = [row for row in self.variants if row.code_original.isdigit()]
        textual = [row for row in self.variants if row.code_original and not row.code_original.isdigit()]
        return {
            "rows": len(self.variants),
            "numeric_rows": len(numeric),
            "numeric_unique": len({row.code_normalized for row in numeric}),
            "textual_rows": len(textual),
            "missing_code_rows": sum(not row.code_original for row in self.variants),
            "duplicate_codes": duplicate_codes,
            "scenario_values": len({row.scenario for row in self.variants if row.scenario}),
            "stage_values": len({row.stage for row in self.variants if row.stage}),
        }


def _dataset_token(*, version: str, source_hash: str, variants: Iterable[MappingVariant]) -> str:
    payload = {
        "version": version,
        "source_hash": source_hash.lower(),
        "rows": [asdict(row) for row in variants],
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def load_mapping_xlsx(path: str | Path, *, mapping_version: str) -> MappingDataset:
    source = Path(path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("A matriz deve ser um arquivo .xlsx sem macros.")
    size = source.stat().st_size
    if size <= 0 or size > MAPPING_MAX_BYTES:
        raise ValueError(f"Tamanho inválido da matriz: {size} bytes.")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=False, data_only=False, keep_vba=False)
    try:
        if MAPPING_SHEET not in workbook.sheetnames:
            raise ValueError(f"A aba obrigatória {MAPPING_SHEET!r} não foi encontrada.")
        worksheet = workbook[MAPPING_SHEET]
        headers = tuple(normalize_display_text(worksheet.cell(1, col).value) for col in range(1, 9))
        if tuple(map(normalize_comparison_text, headers)) != tuple(
            map(normalize_comparison_text, MAPPING_HEADERS)
        ):
            raise ValueError(f"Cabeçalhos inválidos na matriz: {headers!r}.")

        variants: list[MappingVariant] = []
        warnings: list[str] = []
        last_row = 1
        for row_number in range(2, worksheet.max_row + 1):
            cells = [worksheet.cell(row_number, col) for col in range(1, 9)]
            if all(cell.value in (None, "") for cell in cells):
                continue
            if any(cell.data_type == "f" or cell.hyperlink is not None for cell in cells):
                raise ValueError(f"Fórmula ou hyperlink não permitido na linha {row_number}.")
            values = [cell.value for cell in cells]
            code_original = normalize_display_text(values[1])
            if isinstance(values[1], float) and values[1].is_integer():
                code_original = str(int(values[1]))
            elif isinstance(values[1], int):
                code_original = str(values[1])
            if not code_original:
                warnings.append(f"linha_{row_number}_sem_codigo")
            description = normalize_display_text(values[2])
            variants.append(
                MappingVariant(
                    source_row=row_number,
                    classification=normalize_display_text(values[0]).upper(),
                    code_original=code_original,
                    code_normalized=normalize_code(code_original),
                    description_original=description,
                    description_normalized=normalize_comparison_text(description),
                    island=normalize_display_text(values[3]),
                    stage=normalize_display_text(values[4]),
                    document_type=normalize_display_text(values[5]),
                    category=normalize_display_text(values[6]),
                    scenario=normalize_display_text(values[7]),
                    mapping_version=mapping_version,
                    source_hash=source_hash,
                )
            )
            last_row = row_number
    finally:
        workbook.close()

    token = _dataset_token(version=mapping_version, source_hash=source_hash, variants=variants)
    return MappingDataset(
        mapping_version=mapping_version,
        source_hash=source_hash,
        sheet_name=MAPPING_SHEET,
        source_range=f"A1:H{last_row}",
        headers=headers,
        variants=tuple(variants),
        preview_token=token,
        source_label=source.name,
        warnings=tuple(warnings),
    )


def load_mapping_json(path: str | Path) -> MappingDataset:
    source = Path(path)
    payload = json.loads(source.read_text(encoding="utf-8"))
    version = normalize_display_text(payload.get("mapping_version"))
    source_hash = normalize_display_text(payload.get("source_hash")).lower()
    if not version or not re.fullmatch(r"[0-9a-f]{64}", source_hash):
        raise ValueError("Fonte normalizada sem versão ou SHA-256 válido.")
    variants = tuple(MappingVariant(**row) for row in payload.get("rows", []))
    if not variants:
        raise ValueError("Fonte normalizada sem linhas de mapeamento.")
    expected = _dataset_token(version=version, source_hash=source_hash, variants=variants)
    declared = normalize_display_text(payload.get("preview_token"))
    if declared and declared != expected:
        raise ValueError("O conteúdo normalizado não corresponde ao preview_token declarado.")
    return MappingDataset(
        mapping_version=version,
        source_hash=source_hash,
        sheet_name=normalize_display_text(payload.get("sheet_name")) or MAPPING_SHEET,
        source_range=normalize_display_text(payload.get("source_range")),
        headers=tuple(payload.get("headers") or MAPPING_HEADERS),
        variants=variants,
        preview_token=expected,
        source_label=source.name,
        warnings=tuple(payload.get("warnings") or ()),
    )


def load_mapping_source(path: str | Path, *, mapping_version: str = "") -> MappingDataset:
    source = Path(path)
    if source.suffix.lower() == ".json":
        dataset = load_mapping_json(source)
        if mapping_version and dataset.mapping_version != mapping_version:
            raise ValueError("A versão informada difere da versão da fonte normalizada.")
        return dataset
    if not mapping_version:
        raise ValueError("--mapping-version é obrigatório para importar uma matriz XLSX.")
    return load_mapping_xlsx(source, mapping_version=mapping_version)


def export_mapping_json(dataset: MappingDataset, path: str | Path) -> None:
    payload = {
        "schema": 1,
        "mapping_version": dataset.mapping_version,
        "source_hash": dataset.source_hash,
        "sheet_name": dataset.sheet_name,
        "source_range": dataset.source_range,
        "headers": list(dataset.headers),
        "preview_token": dataset.preview_token,
        "warnings": list(dataset.warnings),
        "rows": [asdict(row) for row in dataset.variants],
    }
    Path(path).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
