"""Importação Compliance Auditoria (relatório de conferência) com amostragem A×B.

Formato esperado (CSV/XLSX):
  Protocolo; Tipo/Status conferencia; Matricula Inspetor; Nome Inspetor; Data/Hora da Conferência

- Nome Inspetor é ignorado na chave (usa-se só Matricula Inspetor).
- A = protocolos por agente (matrícula)
- B = quantidade de auditores (matrículas distintas amostradas)
- Total importado = A × B
"""

from __future__ import annotations

import random
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import load_workbook

from apps.auditoria.services.import_tabular import decode_csv_rows
from apps.auditoria.services.reinspecao_import import (
    ParsedReinspecaoRow,
    cell_to_str,
    normalize_header,
    parse_datetime_cell,
    resolve_matricula_inspetor,
)

HEADER_TO_FIELD = {
    "PROTOCOLO": "protocolo",
    "TIPO/STATUS CONFERENCIA": "tipo_status",
    "TIPO STATUS CONFERENCIA": "tipo_status",
    "MATRICULA INSPETOR": "matricula_inspetor",
    "MATRICULA DO INSPETOR": "matricula_inspetor",
    "NOME INSPETOR": "nome_inspetor",  # só exibição
    "DATA/HORA DA CONFERENCIA": "data_conferencia",
    "DATA HORA DA CONFERENCIA": "data_conferencia",
}

REQUIRED_HEADER_FIELDS = ("protocolo", "matricula_inspetor")


@dataclass
class ParsedComplianceRow:
    protocolo: str
    matricula_inspetor: str
    nome_inspetor: str
    tipo_status: str
    data_conferencia: datetime | None = None
    excel_row: int = 0

    def to_pendente_row(self) -> ParsedReinspecaoRow:
        return ParsedReinspecaoRow(
            protocolo=self.protocolo,
            matricula_inspetor=self.matricula_inspetor,
            descricao_irregularidades=self.tipo_status or "Conferência",
            data_analise=self.data_conferencia,
            excel_row=self.excel_row,
        )


@dataclass
class AgentSamplePreview:
    matricula: str
    nome_agente: str
    quantidade_protocolos: int
    protocolos: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "matricula": self.matricula,
            "nome_agente": self.nome_agente,
            "quantidade_protocolos": self.quantidade_protocolos,
            "protocolos": self.protocolos,
        }


@dataclass
class ComplianceImportPreview:
    sheet_name: str
    header_row: int
    total_lidas: int
    matriculas_disponiveis: int
    protocolos_por_agente: int
    total_selecionado: int
    agentes: list[AgentSamplePreview] = field(default_factory=list)
    rows: list[ParsedComplianceRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    amostra: list[dict[str, Any]] = field(default_factory=list)
    protocolos_duplicados_arquivo: list[str] = field(default_factory=list)
    protocolos_ignorados: list[dict[str, Any]] = field(default_factory=list)
    total_ignorados_duplicata: int = 0
    registros_importacao: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "kind": "auditoria_compliance_amostra",
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "total_lidas": self.total_lidas,
            "total_falhas": self.total_selecionado,
            "linhas_filtradas": max(0, self.total_lidas - self.total_selecionado),
            "matriculas_disponiveis": self.matriculas_disponiveis,
            "protocolos_por_agente": self.protocolos_por_agente,
            "quantidade_auditores": self.matriculas_disponiveis,
            "total_selecionado": self.total_selecionado,
            "agentes": [agent.to_dict() for agent in self.agentes],
            "errors": self.errors,
            "warnings": self.warnings,
            "protocolos_duplicados_arquivo": self.protocolos_duplicados_arquivo,
            "protocolos_ignorados": self.protocolos_ignorados,
            "total_ignorados_duplicata": self.total_ignorados_duplicata,
            "registros_importacao": self.registros_importacao,
        }


def _cell_at(row: tuple[Any, ...], column_map: dict[str, int], field: str) -> Any:
    idx = column_map.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def find_compliance_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows, start=1):
        column_map: dict[str, int] = {}
        for col_index, cell in enumerate(row):
            key = normalize_header(cell)
            # Normaliza barras/hífens residual
            key = re.sub(r"[/_\-]+", " ", key)
            key = re.sub(r"\s+", " ", key).strip()
            field = HEADER_TO_FIELD.get(key)
            if not field:
                # Tentativas extras após remoção de acentos já feita em normalize_header
                compact = key.replace(" ", "")
                aliases = {
                    "TIPOSTATUSCONFERENCIA": "tipo_status",
                    "MATRICULAINSPETOR": "matricula_inspetor",
                    "MATRICULADOINSPETOR": "matricula_inspetor",
                    "NOMEINSPETOR": "nome_inspetor",
                    "DATAHORADACONFERENCIA": "data_conferencia",
                }
                field = aliases.get(compact)
            if field and field not in column_map:
                column_map[field] = col_index
        if all(field in column_map for field in REQUIRED_HEADER_FIELDS):
            return row_index, column_map
    return None


def parse_compliance_rows(
    rows: list[tuple[Any, ...]],
    *,
    sheet_name: str,
) -> tuple[list[ParsedComplianceRow], int, list[str]]:
    header_info = find_compliance_header_row(rows)
    if not header_info:
        return (
            [],
            0,
            [
                "Cabeçalho não encontrado. É necessário ter as colunas "
                "Protocolo e Matricula Inspetor "
                "(Nome Inspetor é ignorado na chave; use Matricula Inspetor)."
            ],
        )

    header_row, column_map = header_info
    parsed: list[ParsedComplianceRow] = []
    total_lidas = 0

    for excel_row, row in enumerate(rows[header_row:], start=header_row + 1):
        if row is None or all(cell is None or cell_to_str(cell) == "" for cell in row):
            continue

        protocolo = cell_to_str(_cell_at(row, column_map, "protocolo"))
        if not protocolo:
            continue

        total_lidas += 1

        matricula_raw = _cell_at(row, column_map, "matricula_inspetor")
        matricula = resolve_matricula_inspetor(matricula_raw)
        # Nome só para prévia — não entra na chave.
        nome = cell_to_str(_cell_at(row, column_map, "nome_inspetor"))
        tipo = cell_to_str(_cell_at(row, column_map, "tipo_status"))
        data_conf = parse_datetime_cell(_cell_at(row, column_map, "data_conferencia"))

        parsed.append(
            ParsedComplianceRow(
                protocolo=protocolo,
                matricula_inspetor=matricula,
                nome_inspetor=nome,
                tipo_status=tipo,
                data_conferencia=data_conf,
                excel_row=excel_row,
            )
        )

    return parsed, total_lidas, []


def parse_compliance_file(file_bytes: bytes, filename: str) -> tuple[list[ParsedComplianceRow], str, int, list[str]]:
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        rows = decode_csv_rows(file_bytes)
        parsed, total_lidas, errors = parse_compliance_rows(rows, sheet_name="CSV")
        return parsed, "CSV", 1 if not errors else 0, errors

    workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = [tuple(row) for row in sheet.iter_rows(values_only=True)]
        parsed, total_lidas, errors = parse_compliance_rows(rows, sheet_name=sheet.title or "Planilha1")
        header_row = 1 if not errors else 0
        return parsed, sheet.title or "Planilha1", header_row, errors
    finally:
        workbook.close()


def _display_name_for_matricula(rows: list[ParsedComplianceRow]) -> str:
    counts: dict[str, int] = defaultdict(int)
    for row in rows:
        name = (row.nome_inspetor or "").strip()
        if name:
            counts[name] += 1
    if not counts:
        return rows[0].matricula_inspetor if rows else ""
    return max(counts.items(), key=lambda item: item[1])[0]


def sample_compliance_import(
    *,
    file_bytes: bytes,
    filename: str,
    protocolos_por_agente: int,
    rng: random.Random | None = None,
) -> ComplianceImportPreview:
    """Para cada matrícula distinta do arquivo, amostra A protocolos (total A × distinct)."""
    if protocolos_por_agente < 1:
        return ComplianceImportPreview(
            sheet_name="",
            header_row=0,
            total_lidas=0,
            matriculas_disponiveis=0,
            protocolos_por_agente=protocolos_por_agente,
            total_selecionado=0,
            errors=["Informe a quantidade de protocolos por agente (A) maior que zero."],
        )

    parsed, sheet_name, header_row, errors = parse_compliance_file(file_bytes, filename)
    if errors:
        return ComplianceImportPreview(
            sheet_name=sheet_name,
            header_row=header_row,
            total_lidas=0,
            matriculas_disponiveis=0,
            protocolos_por_agente=protocolos_por_agente,
            total_selecionado=0,
            errors=errors,
        )

    by_matricula: dict[str, list[ParsedComplianceRow]] = defaultdict(list)
    for row in parsed:
        # Não amostrar matrículas inválidas agrupadas como "sistema".
        if row.matricula_inspetor == "sistema":
            continue
        by_matricula[row.matricula_inspetor].append(row)

    if not by_matricula:
        return ComplianceImportPreview(
            sheet_name=sheet_name,
            header_row=header_row,
            total_lidas=len(parsed),
            matriculas_disponiveis=0,
            protocolos_por_agente=protocolos_por_agente,
            total_selecionado=0,
            errors=["Nenhuma matrícula de inspetor válida encontrada no arquivo."],
        )

    rng = rng or random.Random()
    selected_rows: list[ParsedComplianceRow] = []
    shortfall_names: list[str] = []

    for matricula in sorted(by_matricula.keys()):
        pool = by_matricula[matricula]
        take = min(protocolos_por_agente, len(pool))
        chosen = rng.sample(pool, take) if take < len(pool) else list(pool)
        chosen_sorted = sorted(chosen, key=lambda item: item.protocolo)
        selected_rows.extend(chosen_sorted)
        nome = _display_name_for_matricula(pool)
        if len(pool) < protocolos_por_agente:
            shortfall_names.append(nome or matricula)

    from apps.auditoria.services.reinspecao_import_dedupe import (
        apply_compliance_import_dedupe,
        build_registros_importacao,
        compliance_row_registro,
        ignorados_to_preview,
    )

    import_rows, dup_arquivo, ignorados_portal = apply_compliance_import_dedupe(
        selected_rows,
        contexto="auditoria_compliance",
    )
    dedupe_preview = ignorados_to_preview(dup_arquivo, ignorados_portal)
    registros_importacao = build_registros_importacao(
        kept_rows=import_rows,
        duplicados_arquivo=dup_arquivo,
        ignorados_portal=ignorados_portal,
        row_to_registro=lambda row: compliance_row_registro(
            row,
            contexto="auditoria_compliance",
        ),
    )

    by_matricula_import: dict[str, list[ParsedComplianceRow]] = defaultdict(list)
    for row in import_rows:
        by_matricula_import[row.matricula_inspetor].append(row)

    agentes = []
    for matricula in sorted(by_matricula_import.keys()):
        pool = by_matricula_import[matricula]
        nome = _display_name_for_matricula(pool)
        agentes.append(
            AgentSamplePreview(
                matricula=matricula,
                nome_agente=nome,
                quantidade_protocolos=len(pool),
                protocolos=[row.protocolo for row in pool],
            )
        )

    warnings: list[str] = []
    if shortfall_names:
        warnings.append(
            f"Existem agentes com quantidade de protocolos abaixo do mínimo definido ({protocolos_por_agente})."
        )
    if dedupe_preview["total_ignorados_duplicata"]:
        warnings.append(
            f"{dedupe_preview['total_ignorados_duplicata']} registro(s) ignorado(s) por duplicidade "
            "(no arquivo ou já existentes no portal)."
        )

    amostra = []

    return ComplianceImportPreview(
        sheet_name=sheet_name,
        header_row=header_row,
        total_lidas=len(parsed),
        matriculas_disponiveis=len(by_matricula),
        protocolos_por_agente=protocolos_por_agente,
        total_selecionado=len(import_rows),
        agentes=agentes,
        rows=import_rows,
        amostra=amostra,
        warnings=warnings,
        protocolos_duplicados_arquivo=dedupe_preview["protocolos_duplicados_arquivo"],
        protocolos_ignorados=dedupe_preview["protocolos_ignorados"],
        total_ignorados_duplicata=dedupe_preview["total_ignorados_duplicata"],
        registros_importacao=registros_importacao,
    )


def create_compliance_import_staging(
    *,
    user,
    file_bytes: bytes,
    filename: str,
    preview: ComplianceImportPreview,
) -> Any:
    from datetime import timedelta

    from django.core.files.base import ContentFile

    from apps.auditoria.models import AuditoriaComplianceImportStaging

    expires_at = timezone.now() + timedelta(hours=2)
    staging = AuditoriaComplianceImportStaging(
        nome_arquivo=filename,
        preview={
            **preview.to_dict(),
            "fila_contexto": "auditoria_compliance",
            "rows": [
                {
                    "protocolo": row.protocolo,
                    "matricula_inspetor": row.matricula_inspetor,
                    "descricao_irregularidades": row.tipo_status or "Conferência",
                    "data_analise": (
                        row.data_conferencia.isoformat() if row.data_conferencia else None
                    ),
                    "excel_row": row.excel_row,
                    "nome_inspetor": row.nome_inspetor,
                }
                for row in preview.rows
            ],
        },
        created_by=user,
        expires_at=expires_at,
    )
    staging.arquivo.save(filename, ContentFile(file_bytes), save=False)
    staging.save()
    return staging
