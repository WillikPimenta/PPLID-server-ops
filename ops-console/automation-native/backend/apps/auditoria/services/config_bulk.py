from __future__ import annotations

from io import BytesIO
import unicodedata

from django.db import transaction
from openpyxl import load_workbook

from apps.auditoria.models import AuditoriaCatalogItem, AuditoriaMotivoFalha
from apps.auditoria.services.catalog_items import normalize_catalog_value
from apps.auditoria.services.config_change_log import build_changes, log_config_change, snapshot_fields


class ConfigBulkError(Exception):
    pass


def _key(value) -> str:
    text = unicodedata.normalize("NFKD", str("" if value is None else value))
    return "_".join(text.encode("ascii", "ignore").decode("ascii").lower().strip().split())


def _boolean(value, *, default=True) -> bool:
    if value is None or str(value).strip() == "":
        return default
    normalized = _key(value)
    if normalized in {"sim", "true", "1", "ativo"}:
        return True
    if normalized in {"nao", "false", "0", "inativo"}:
        return False
    raise ConfigBulkError(f"Valor invalido para Ativo: {value}.")


def read_xlsx_rows(upload) -> list[dict]:
    if not upload or not (upload.name or "").lower().endswith(".xlsx"):
        raise ConfigBulkError("Envie um arquivo no formato .xlsx.")
    content = upload.read()
    if len(content) > 5 * 1024 * 1024:
        raise ConfigBulkError("A planilha deve ter no maximo 5 MB.")
    try:
        sheet = load_workbook(BytesIO(content), read_only=True, data_only=True).active
    except Exception as exc:
        raise ConfigBulkError("Nao foi possivel ler a planilha Excel.") from exc
    rows = sheet.iter_rows(values_only=True)
    headers = [_key(value) for value in next(rows, ())]
    if not any(headers):
        raise ConfigBulkError("A planilha nao possui cabecalho.")
    result = []
    for excel_row, values in enumerate(rows, start=2):
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        if any(value is not None and str(value).strip() for value in values):
            row["_excel_row"] = excel_row
            result.append(row)
    if not result:
        raise ConfigBulkError("A planilha nao possui dados para importar.")
    if len(result) > 2000:
        raise ConfigBulkError("A planilha excede o limite de 2.000 linhas.")
    return result


def _order(value, fallback: int) -> int:
    if value is None or str(value).strip() == "":
        return fallback
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise ConfigBulkError(f"Ordem invalida: {value}.") from exc
    if parsed < 0:
        raise ConfigBulkError("A ordem nao pode ser negativa.")
    return parsed


SCENARIO_FIELDS = ("motivo", "criticidade", "segmentos", "subsegmento", "sort_order", "active")
SCENARIO_COLUMNS = [
    {"key": "ordem", "label": "Ordem"}, {"key": "cenario", "label": "Cenario"},
    {"key": "criticidade", "label": "Criticidade"}, {"key": "segmentos", "label": "Segmentos"},
    {"key": "subsegmento", "label": "Subsegmento"}, {"key": "ativo", "label": "Ativo"},
]


def _preview_result(rows: list[dict], columns: list[dict]) -> dict:
    summary = {
        "total": len(rows),
        "new": sum(row["status"] == "novo" for row in rows),
        "updated": sum(row["status"] == "atualizado" for row in rows),
        "reactivated": sum(row["status"] == "reativado" for row in rows),
        "deactivated": sum(row["status"] == "inativado" for row in rows),
        "unchanged": sum(row["status"] == "sem_alteracao" for row in rows),
        "duplicates": sum(row["status"] == "duplicado" for row in rows),
        "errors": sum(row["status"] == "erro" for row in rows),
    }
    summary["actions"] = summary["new"] + summary["updated"] + summary["reactivated"] + summary["deactivated"]
    return {"columns": columns, "rows": rows, "summary": summary, "can_confirm": summary["errors"] == 0}


def preview_catalog_xlsx(*, catalog: str, upload) -> dict:
    source_rows = read_xlsx_rows(upload)
    existing = set(AuditoriaCatalogItem.objects.filter(catalog=catalog).values_list("value", flat=True))
    seen = set(existing)
    next_order = (AuditoriaCatalogItem.objects.filter(catalog=catalog).order_by("-sort_order").values_list("sort_order", flat=True).first() or 0) + 1
    preview = []
    for row in source_rows:
        try:
            value = normalize_catalog_value(catalog, str(row.get("valor") or row.get("value") or "").strip())
            if not value:
                raise ConfigBulkError("Informe o Valor.")
            data = {
                "ordem": _order(row.get("ordem") if "ordem" in row else row.get("sort_order"), next_order),
                "valor": value,
                "ativo": _boolean(row.get("ativo") if "ativo" in row else row.get("active")),
            }
            duplicate = value in seen
            preview.append({"linha": row["_excel_row"], "status": "duplicado" if duplicate else "novo", "mensagem": "Registro ja existente; sera ignorado." if duplicate else "", "dados": data})
            seen.add(value)
            next_order = max(next_order + 1, data["ordem"] + 1)
        except ConfigBulkError as exc:
            preview.append({"linha": row["_excel_row"], "status": "erro", "mensagem": str(exc), "dados": {}})
    return _preview_result(preview, [{"key": "ordem", "label": "Ordem"}, {"key": "valor", "label": "Valor"}, {"key": "ativo", "label": "Ativo"}])


def _scenario_row(row: dict, fallback_order: int, *, force_active: bool = False) -> dict:
    data = {
        "ordem": _order(row.get("ordem") if "ordem" in row else row.get("sort_order"), fallback_order),
        "cenario": str(row.get("cenario") or row.get("motivo") or row.get("alerta") or "").strip(),
        "criticidade": str(row.get("criticidade") or "").strip(),
        "segmentos": str(row.get("segmentos") or row.get("segmento") or "").strip(),
        "subsegmento": str(row.get("subsegmento") or "").strip(),
        "ativo": True if force_active else _boolean(row.get("ativo") if "ativo" in row else row.get("active")),
    }
    if not all((data["cenario"], data["criticidade"], data["segmentos"], data["subsegmento"])):
        raise ConfigBulkError("Preencha Cenario, Criticidade, Segmentos e Subsegmento.")
    return data


def preview_scenarios_xlsx(*, upload, mode: str = "append") -> dict:
    if mode not in {"append", "replace"}:
        raise ConfigBulkError("Modo de importacao invalido.")
    source_rows = read_xlsx_rows(upload)
    existing_items = {item.motivo: item for item in AuditoriaMotivoFalha.objects.all()}
    seen = set(existing_items) if mode == "append" else set()
    next_order = (AuditoriaMotivoFalha.objects.order_by("-sort_order").values_list("sort_order", flat=True).first() or 0) + 1
    preview = []
    imported_names = set()
    for index, row in enumerate(source_rows):
        try:
            fallback_order = index if mode == "replace" else next_order
            data = _scenario_row(row, fallback_order, force_active=mode == "replace")
            if data["cenario"] in imported_names:
                raise ConfigBulkError("Cenario repetido dentro da planilha.")
            imported_names.add(data["cenario"])
            current = existing_items.get(data["cenario"])
            if mode == "append":
                duplicate = data["cenario"] in seen
                status_value = "duplicado" if duplicate else "novo"
                message = "Cenario ja existente; sera ignorado." if duplicate else ""
            elif current is None:
                status_value, message = "novo", "Novo cenario sera criado e ativado."
            else:
                changed = any((
                    current.criticidade != data["criticidade"], current.segmentos != data["segmentos"],
                    current.subsegmento != data["subsegmento"], current.sort_order != data["ordem"],
                ))
                if not current.active:
                    status_value, message = "reativado", "Cenario existente sera atualizado e reativado."
                elif changed:
                    status_value, message = "atualizado", "Dados do cenario serao atualizados."
                else:
                    status_value, message = "sem_alteracao", "Cenario ja esta atualizado."
            preview.append({"linha": row["_excel_row"], "status": status_value, "mensagem": message, "dados": data})
            seen.add(data["cenario"])
            next_order = max(next_order + 1, data["ordem"] + 1)
        except ConfigBulkError as exc:
            preview.append({"linha": row["_excel_row"], "status": "erro", "mensagem": str(exc), "dados": {}})
    if mode == "replace":
        for item in existing_items.values():
            if item.active and item.motivo not in imported_names:
                preview.append({
                    "linha": None,
                    "status": "inativado",
                    "mensagem": "Ausente da planilha; sera inativado.",
                    "dados": {"ordem": item.sort_order, "cenario": item.motivo, "criticidade": item.criticidade,
                              "segmentos": item.segmentos, "subsegmento": item.subsegmento, "ativo": False},
                })
    return _preview_result(preview, SCENARIO_COLUMNS)


@transaction.atomic
def import_catalog_xlsx(*, catalog: str, upload, user) -> dict:
    rows = read_xlsx_rows(upload)
    existing = set(AuditoriaCatalogItem.objects.filter(catalog=catalog).values_list("value", flat=True))
    next_order = (AuditoriaCatalogItem.objects.filter(catalog=catalog).order_by("-sort_order").values_list("sort_order", flat=True).first() or 0) + 1
    created = skipped = 0
    for row in rows:
        value = normalize_catalog_value(catalog, str(row.get("valor") or row.get("value") or "").strip())
        if not value:
            raise ConfigBulkError(f"Linha {row['_excel_row']}: informe o Valor.")
        if value in existing:
            skipped += 1
            continue
        item = AuditoriaCatalogItem.objects.create(
            catalog=catalog,
            value=value,
            label=str(row.get("nome_exibido") or row.get("label") or value).strip(),
            sort_order=_order(row.get("ordem") if "ordem" in row else row.get("sort_order"), next_order),
            active=_boolean(row.get("ativo") if "ativo" in row else row.get("active")),
        )
        fields = ("catalog", "value", "label", "sort_order", "active")
        log_config_change(instance=item, changes=build_changes({field: None for field in fields}, snapshot_fields(item, fields)), user=user)
        existing.add(value)
        next_order = max(next_order + 1, item.sort_order + 1)
        created += 1
    return {"created": created, "skipped": skipped}


@transaction.atomic
def import_scenarios_xlsx(*, upload, user, mode: str = "append") -> dict:
    if mode not in {"append", "replace"}:
        raise ConfigBulkError("Modo de importacao invalido.")
    rows = read_xlsx_rows(upload)
    existing_items = {item.motivo: item for item in AuditoriaMotivoFalha.objects.select_for_update()}
    existing = set(existing_items)
    next_order = (AuditoriaMotivoFalha.objects.order_by("-sort_order").values_list("sort_order", flat=True).first() or 0) + 1
    created = skipped = updated = reactivated = deactivated = 0
    imported_names = set()
    for index, row in enumerate(rows):
        try:
            data = _scenario_row(row, index if mode == "replace" else next_order, force_active=mode == "replace")
        except ConfigBulkError as exc:
            raise ConfigBulkError(f"Linha {row['_excel_row']}: {exc}") from exc
        motivo = data["cenario"]
        if motivo in imported_names:
            raise ConfigBulkError(f"Linha {row['_excel_row']}: cenario repetido dentro da planilha.")
        imported_names.add(motivo)
        item = existing_items.get(motivo)
        if item is not None and mode == "append":
            skipped += 1
            continue
        if item is None:
            item = AuditoriaMotivoFalha.objects.create(
                motivo=motivo, criticidade=data["criticidade"], segmentos=data["segmentos"],
                subsegmento=data["subsegmento"], sort_order=data["ordem"], active=data["ativo"],
            )
            log_config_change(instance=item, changes=build_changes({field: None for field in SCENARIO_FIELDS}, snapshot_fields(item, SCENARIO_FIELDS)), user=user)
            created += 1
        else:
            before = snapshot_fields(item, SCENARIO_FIELDS)
            was_inactive = not item.active
            item.criticidade = data["criticidade"]
            item.segmentos = data["segmentos"]
            item.subsegmento = data["subsegmento"]
            item.sort_order = data["ordem"]
            item.active = True
            changes = build_changes(before, snapshot_fields(item, SCENARIO_FIELDS))
            if changes:
                item.save()
                log_config_change(instance=item, changes=changes, user=user)
                if was_inactive:
                    reactivated += 1
                else:
                    updated += 1
        existing.add(motivo)
        next_order = max(next_order + 1, item.sort_order + 1)
    if mode == "replace":
        for item in existing_items.values():
            if item.active and item.motivo not in imported_names:
                before = snapshot_fields(item, SCENARIO_FIELDS)
                item.active = False
                item.save(update_fields=["active", "updated_at"])
                log_config_change(instance=item, changes=build_changes(before, snapshot_fields(item, SCENARIO_FIELDS)), user=user)
                deactivated += 1
    return {"created": created, "updated": updated, "reactivated": reactivated, "deactivated": deactivated, "skipped": skipped}


@transaction.atomic
def bulk_update_catalog(*, catalog: str, rows: list[dict], user) -> int:
    if not isinstance(rows, list) or not rows or len(rows) > 2000:
        raise ConfigBulkError("Informe entre 1 e 2.000 registros.")
    ids = [row.get("id") for row in rows]
    targets = {item.pk: item for item in AuditoriaCatalogItem.objects.select_for_update().filter(catalog=catalog, pk__in=ids)}
    if len(targets) != len(set(ids)):
        raise ConfigBulkError("Um ou mais registros nao pertencem a este catalogo.")
    planned = {}
    for row in rows:
        item = targets[row["id"]]
        value = normalize_catalog_value(catalog, str(row.get("value") or "").strip())
        if not value:
            raise ConfigBulkError("Todos os registros devem possuir Valor.")
        if value in planned and planned[value] != item.pk:
            raise ConfigBulkError(f"Valor duplicado na alteracao em massa: {value}.")
        planned[value] = item.pk
    conflict = AuditoriaCatalogItem.objects.filter(catalog=catalog, value__in=planned).exclude(pk__in=ids).values_list("value", flat=True).first()
    if conflict:
        raise ConfigBulkError(f"O valor {conflict} ja existe no catalogo.")
    changed_count = 0
    fields = ("catalog", "value", "label", "sort_order", "active")
    for row in rows:
        item = targets[row["id"]]
        before = snapshot_fields(item, fields)
        item.value = normalize_catalog_value(catalog, str(row.get("value") or "").strip())
        item.label = str(row.get("label") or item.value).strip()
        item.sort_order = _order(row.get("sort_order"), item.sort_order)
        item.active = _boolean(row.get("active"), default=item.active)
        changes = build_changes(before, snapshot_fields(item, fields))
        if changes:
            item.save()
            log_config_change(instance=item, changes=changes, user=user)
            changed_count += 1
    return changed_count


@transaction.atomic
def bulk_update_scenarios(*, rows: list[dict], user) -> int:
    if not isinstance(rows, list) or not rows or len(rows) > 2000:
        raise ConfigBulkError("Informe entre 1 e 2.000 registros.")
    ids = [row.get("id") for row in rows]
    targets = {item.pk: item for item in AuditoriaMotivoFalha.objects.select_for_update().filter(pk__in=ids)}
    if len(targets) != len(set(ids)):
        raise ConfigBulkError("Um ou mais cenarios nao foram encontrados.")
    planned = {}
    for row in rows:
        motivo = str(row.get("motivo") or "").strip()
        if not all((motivo, str(row.get("criticidade") or "").strip(), str(row.get("segmentos") or "").strip(), str(row.get("subsegmento") or "").strip())):
            raise ConfigBulkError("Preencha todos os campos obrigatorios dos cenarios.")
        if motivo in planned and planned[motivo] != row["id"]:
            raise ConfigBulkError(f"Cenario duplicado na alteracao em massa: {motivo}.")
        planned[motivo] = row["id"]
    conflict = AuditoriaMotivoFalha.objects.filter(motivo__in=planned).exclude(pk__in=ids).values_list("motivo", flat=True).first()
    if conflict:
        raise ConfigBulkError(f"O cenario {conflict} ja existe.")
    changed_count = 0
    fields = ("motivo", "criticidade", "segmentos", "subsegmento", "sort_order", "active")
    for row in rows:
        item = targets[row["id"]]
        before = snapshot_fields(item, fields)
        for field in ("motivo", "criticidade", "segmentos", "subsegmento"):
            setattr(item, field, str(row.get(field) or "").strip())
        item.sort_order = _order(row.get("sort_order"), item.sort_order)
        item.active = _boolean(row.get("active"), default=item.active)
        changes = build_changes(before, snapshot_fields(item, fields))
        if changes:
            item.save()
            log_config_change(instance=item, changes=changes, user=user)
            changed_count += 1
    return changed_count
