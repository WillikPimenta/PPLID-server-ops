from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from typing import Any

from django.utils import timezone
from openpyxl import load_workbook

MAX_PREVIEW_ROWS = 10
TIPO_FALHA_REINSPECAO = "reinspecao"
SISTEMA_USUARIO = "sistema"
# Padrão operacional LAN: c#####a ou c#####q (ex.: c19131q).
OPERACAO_LAN_PATTERN = re.compile(r"^c\d{5}[aq]$", re.IGNORECASE)

IRREGULARIDADES_DATA_RESPOSTA_ESPELHO = (
    "IC - 668 - Plano no GED diverge do Plano no Contrato de Habilitação",
    "CO - 49 - Comprovante de Endereço Ausente",
)

HEADER_TO_FIELD = {
    "PROTOCOLO": "protocolo",
    "DATA DA CONTESTACAO": "data_contestacao",
    "DATA CONTESTACAO": "data_contestacao",
    "MATRICULA DO INSPETOR": "matricula_inspetor",
    "DATA DE RESPOSTA": "data_resposta",
    "DATA DA RESPOSTA": "data_resposta",
    "DESCRICAO DAS IRREGULARIDADES": "descricao_irregularidades",
}

REQUIRED_HEADER_FIELDS = (
    "protocolo",
    "data_contestacao",
    "data_resposta",
    "matricula_inspetor",
    "descricao_irregularidades",
)


@dataclass
class ParsedReinspecaoRow:
    protocolo: str
    matricula_inspetor: str
    descricao_irregularidades: str
    data_contestacao: datetime | None = None
    data_analise: datetime | None = None
    excel_row: int = 0
    nome_inspetor: str = ""
    codigo_irregularidade: str = ""
    classificacao_irregularidade: str = ""
    cenario_mapeado: str = ""
    etapa_mapeada: str = ""
    mapping_scenario_status: str = ""
    mapping_stage_status: str = ""
    mapping_version: str = ""
    mapping_source_hash: str = ""
    mapping_candidates: list[dict[str, Any]] = field(default_factory=list)


@dataclass(frozen=True)
class ReinspecaoGedObservation:
    row: ParsedReinspecaoRow
    ged_elegivel: bool
    motivo: str


@dataclass
class ReinspecaoImportPreview:
    sheet_name: str
    header_row: int
    total_falhas: int
    total_lidas: int = 0
    linhas_filtradas: int = 0
    amostra: list[dict[str, Any]] = field(default_factory=list)
    rows: list[ParsedReinspecaoRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    protocolos_duplicados_arquivo: list[str] = field(default_factory=list)
    protocolos_ignorados: list[dict[str, Any]] = field(default_factory=list)
    total_ignorados_duplicata: int = 0
    registros_importacao: list[dict[str, Any]] = field(default_factory=list)
    resumo_planilha: dict[str, Any] = field(default_factory=dict)
    mapping_summary: dict[str, Any] = field(default_factory=dict)
    mapping_issues: list[dict[str, Any]] = field(default_factory=list)
    mapping_version: str = ""
    mapping_source_hash: str = ""
    observations: list[ReinspecaoGedObservation] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "total_falhas": self.total_falhas,
            "total_lidas": self.total_lidas,
            "linhas_filtradas": self.linhas_filtradas,
            "errors": self.errors,
            "protocolos_duplicados_arquivo": self.protocolos_duplicados_arquivo,
            "protocolos_ignorados": self.protocolos_ignorados,
            "total_ignorados_duplicata": self.total_ignorados_duplicata,
            "registros_importacao": self.registros_importacao,
            "resumo_planilha": self.resumo_planilha,
            "mapping_summary": self.mapping_summary,
            "mapping_issues": self.mapping_issues,
            "mapping_version": self.mapping_version,
            "mapping_source_hash": self.mapping_source_hash,
        }


def apply_mapping_to_preview(
    preview: ReinspecaoImportPreview,
    *,
    resolver=None,
) -> ReinspecaoImportPreview:
    """Enriquece a prévia em memória com um único carregamento da matriz."""
    from apps.auditoria.services.reinspecao_mapping import (
        STATUS_AMBIGUOUS,
        STATUS_MATCHED,
        STATUS_UNMATCHED,
        ReinspecaoMappingResolver,
    )

    resolver = resolver or ReinspecaoMappingResolver.from_active_mapping()
    counters = {
        "parsed": 0,
        "matched": 0,
        "unmatched": 0,
        "ambiguous_scenario": 0,
        "ambiguous_stage": 0,
    }
    issues: dict[tuple[str, str, str], dict[str, Any]] = {}
    rows_by_excel_row: dict[int, ParsedReinspecaoRow] = {}
    for row in preview.rows:
        resolution = resolver.resolve(row.descricao_irregularidades)
        row.codigo_irregularidade = resolution.extracted.code
        row.classificacao_irregularidade = resolution.extracted.classification
        row.cenario_mapeado = resolution.scenario
        row.etapa_mapeada = resolution.stage
        row.mapping_scenario_status = resolution.scenario_status
        row.mapping_stage_status = resolution.stage_status
        row.mapping_version = resolution.mapping_version
        row.mapping_source_hash = resolution.source_hash
        row.mapping_candidates = resolution.diagnostic_candidates()
        rows_by_excel_row[row.excel_row] = row

        counters["parsed"] += int(resolution.extracted.parsed)
        counters["matched"] += int(
            resolution.scenario_status == STATUS_MATCHED
            or resolution.stage_status == STATUS_MATCHED
        )
        counters["unmatched"] += int(
            resolution.scenario_status == STATUS_UNMATCHED
            and resolution.stage_status == STATUS_UNMATCHED
        )
        counters["ambiguous_scenario"] += int(
            resolution.scenario_status == STATUS_AMBIGUOUS
        )
        counters["ambiguous_stage"] += int(
            resolution.stage_status == STATUS_AMBIGUOUS
        )
        if (
            resolution.scenario_status != STATUS_MATCHED
            or resolution.stage_status != STATUS_MATCHED
        ):
            key = (
                resolution.extracted.code or "<sem_codigo>",
                resolution.scenario_status,
                resolution.stage_status,
            )
            issue = issues.setdefault(
                key,
                {
                    "code": resolution.extracted.code,
                    "classification": resolution.extracted.classification,
                    "scenario_status": resolution.scenario_status,
                    "stage_status": resolution.stage_status,
                    "rows": [],
                    "candidates": resolution.diagnostic_candidates(),
                },
            )
            issue["rows"].append(row.excel_row)

    for registro in preview.registros_importacao:
        row = rows_by_excel_row.get(int(registro.get("excel_row") or 0))
        if row is None:
            continue
        registro.update(
            {
                "codigo_irregularidade": row.codigo_irregularidade,
                "cenario": row.cenario_mapeado,
                "etapa": row.etapa_mapeada,
                "mapping_scenario_status": row.mapping_scenario_status,
                "mapping_stage_status": row.mapping_stage_status,
            }
        )

    preview.mapping_summary = {"total": len(preview.rows), **counters}
    preview.mapping_issues = sorted(issues.values(), key=lambda item: item["code"] or "")
    preview.mapping_version = resolver.mapping_version
    preview.mapping_source_hash = resolver.source_hash
    if preview.rows and not resolver.mapping_version:
        preview.errors.append(
            "Nenhuma matriz de Reinspeção aprovada e ativa está disponível. "
            "Importe a matriz antes de confirmar novos registros."
        )
    return preview


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def is_empty_data_resposta(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, datetime):
        return False
    text = cell_to_str(value)
    if not text:
        return True
    return text in {"-", "—", "–", "N/A", "NA", "NULL"}


def parse_datetime_cell(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        dt = value
        if timezone.is_naive(dt):
            return timezone.make_aware(dt, timezone.get_current_timezone())
        return dt
    text = cell_to_str(value)
    if not text or text in {"-", "—", "–"}:
        return None
    try:
        dt = datetime.fromisoformat(text.replace("Z", "+00:00"))
        if timezone.is_naive(dt):
            return timezone.make_aware(dt, timezone.get_current_timezone())
        return dt
    except ValueError:
        pass
    for fmt in (
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
        "%d/%m/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            dt = datetime.strptime(text, fmt)
            return timezone.make_aware(dt, timezone.get_current_timezone())
        except ValueError:
            continue
    return None


def normalize_irregularidade(value: Any) -> str:
    text = normalize_header(value)
    text = re.sub(r"\s*[-–—]\s*", " - ", text)
    return re.sub(r"\s+", " ", text).strip()


IRREGULARIDADES_DATA_RESPOSTA_ESPELHO_NORMALIZADAS = frozenset(
    normalize_irregularidade(value) for value in IRREGULARIDADES_DATA_RESPOSTA_ESPELHO
)

ELEGIBILIDADE_DETALHE_LABELS = {
    "irregularidade_vazia": "Irregularidade vazia",
    "respondida": "Contestação já respondida",
    "data_invalida": "Data de contestação ou resposta inválida",
    "data_inconsistente": "Data de resposta anterior à contestação",
    "data_resposta_vazia": "Data de resposta vazia",
    "data_resposta_espelho": "Datas de contestação e resposta iguais",
}


@dataclass(frozen=True)
class LinhaResumoPlanilha:
    protocolo: str
    excel_row: int
    excecao: bool
    elegivel: bool
    motivo: str
    matricula_inspetor: str
    descricao: str
    data: str | None


def empty_resumo_planilha() -> dict[str, Any]:
    return {
        "total_protocolos": 0,
        "total_respondidos": 0,
        "total_irregularidades_vazias": 0,
        "total_excecoes": 0,
        "total_agentes": 0,
    }


def is_irregularidade_ic668_co49(descricao: Any) -> bool:
    text = cell_to_str(descricao)
    if not text:
        return False
    return normalize_irregularidade(text) in IRREGULARIDADES_DATA_RESPOSTA_ESPELHO_NORMALIZADAS


def is_protocolo_excecao(linha: LinhaResumoPlanilha) -> bool:
    """Exceção operacional: IC 668/CO 49 com datas de contestação e resposta iguais."""
    return linha.excecao and linha.elegivel and linha.motivo == "data_resposta_espelho"


def elegibilidade_detalhe_label(motivo: str) -> str:
    return ELEGIBILIDADE_DETALHE_LABELS.get(motivo, motivo.replace("_", " "))


def format_observacao_registro(
    linha: LinhaResumoPlanilha,
    *,
    dedupe_motivo: str | None,
) -> str:
    from apps.auditoria.services.reinspecao_import_dedupe import MOTIVO_LABELS

    if dedupe_motivo:
        motivo = MOTIVO_LABELS.get(dedupe_motivo, dedupe_motivo)
        return f"{motivo} — não será importado"
    if not linha.elegivel:
        return f"{elegibilidade_detalhe_label(linha.motivo)} — não será importado"
    if is_protocolo_excecao(linha):
        return "Exceção IC 668/CO 49"
    return ""


def build_registros_planilha_preview(
    linhas: list[LinhaResumoPlanilha],
    *,
    import_rows: list[ParsedReinspecaoRow],
    duplicados_arquivo: list[Any],
    ignorados_portal: list[Any],
) -> list[dict[str, Any]]:
    from apps.auditoria.services.reinspecao_import_dedupe import MOTIVO_LABELS

    del import_rows  # elegibilidade já refletida em linhas
    dedupe_por_linha = {
        (item.protocolo, item.excel_row): item
        for item in (*duplicados_arquivo, *ignorados_portal)
    }
    registros: list[dict[str, Any]] = []
    for linha in linhas:
        ignorado = dedupe_por_linha.get((linha.protocolo, linha.excel_row))
        dedupe_motivo = ignorado.motivo if ignorado else None
        if dedupe_motivo:
            status = "duplicado"
        elif not linha.elegivel:
            status = "filtrado"
        else:
            status = "ok"
        registro: dict[str, Any] = {
            "status_importacao": status,
            "protocolo": linha.protocolo,
            "matricula_inspetor": linha.matricula_inspetor,
            "descricao": linha.descricao,
            "data": linha.data,
            "excel_row": linha.excel_row,
            "observacao": format_observacao_registro(linha, dedupe_motivo=dedupe_motivo),
            "motivo": dedupe_motivo if status == "duplicado" else (linha.motivo if status == "filtrado" else None),
            "motivo_label": (
                MOTIVO_LABELS.get(dedupe_motivo, dedupe_motivo)
                if dedupe_motivo
                else (elegibilidade_detalhe_label(linha.motivo) if status == "filtrado" else None)
            ),
        }
        if ignorado:
            registro["dedupe_key"] = ignorado.dedupe_key
        registros.append(registro)
    registros.sort(
        key=lambda item: (
            int(item.get("excel_row") or 0),
            str(item.get("protocolo") or ""),
        )
    )
    return registros


def avaliar_elegibilidade_reinspecao(
    *,
    descricao_irregularidades: Any,
    data_contestacao: Any,
    data_resposta: Any,
) -> tuple[bool, str]:
    """Decide se a linha representa uma contestação ainda não respondida."""
    descricao = cell_to_str(descricao_irregularidades)
    if not descricao:
        return False, "irregularidade_vazia"

    contestacao_dt = parse_datetime_cell(data_contestacao)
    if contestacao_dt is None:
        return False, "data_invalida"

    if is_empty_data_resposta(data_resposta):
        return True, "data_resposta_vazia"

    irregularidade_normalizada = normalize_irregularidade(descricao)
    if irregularidade_normalizada not in IRREGULARIDADES_DATA_RESPOSTA_ESPELHO_NORMALIZADAS:
        return False, "respondida"

    resposta_dt = parse_datetime_cell(data_resposta)
    if resposta_dt is None:
        return False, "data_invalida"

    if contestacao_dt == resposta_dt:
        return True, "data_resposta_espelho"
    if resposta_dt > contestacao_dt:
        return False, "respondida"
    return False, "data_inconsistente"


def normalize_matricula(value: Any) -> str:
    text = cell_to_str(value)
    if not text:
        return ""
    return text.strip().lower()


def is_operacao_lan(value: str) -> bool:
    """True se matrícula segue o padrão operacional c#####a / c#####q."""
    return bool(OPERACAO_LAN_PATTERN.match((value or "").strip()))


def resolve_matricula_inspetor(value: Any) -> str:
    """
    Normaliza matrícula do inspetor para persistência.
    Padrão c#####a/c#####q é mantido; qualquer outro valor vira 'sistema'.
    """
    matricula = normalize_matricula(value)
    if is_operacao_lan(matricula):
        return matricula
    return SISTEMA_USUARIO


def find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows, start=1):
        column_map: dict[str, int] = {}
        for col_index, cell in enumerate(row):
            key = normalize_header(cell)
            field = HEADER_TO_FIELD.get(key)
            if field and field not in column_map:
                column_map[field] = col_index
        if all(field in column_map for field in REQUIRED_HEADER_FIELDS):
            return row_index, column_map
    return None


def _cell_at(row: tuple[Any, ...], column_map: dict[str, int], field: str) -> Any:
    idx = column_map.get(field)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _parse_reinspecao_rows(
    rows: list[tuple[Any, ...]],
    *,
    sheet_name: str,
    allow_no_eligible: bool = False,
) -> ReinspecaoImportPreview:
    header_info = find_header_row(rows)
    if not header_info:
        return ReinspecaoImportPreview(
            sheet_name=sheet_name,
            header_row=0,
            total_falhas=0,
            errors=[
                "Cabeçalho não encontrado. É necessário ter as colunas "
                "Protocolo, Data da contestação, Data de resposta, "
                "Matricula do Inspetor e Descrição das Irregularidades."
            ],
        )

    header_row, column_map = header_info
    parsed_rows: list[ParsedReinspecaoRow] = []
    total_lidas = 0
    linhas_filtradas = 0
    resumo = empty_resumo_planilha()
    matriculas_lidas: set[str] = set()
    linhas_resumo: list[LinhaResumoPlanilha] = []
    observations: list[ReinspecaoGedObservation] = []

    for excel_row, row in enumerate(rows[header_row:], start=header_row + 1):
        if row is None or all(cell is None or cell_to_str(cell) == "" for cell in row):
            continue

        protocolo = cell_to_str(_cell_at(row, column_map, "protocolo"))
        if not protocolo:
            continue

        total_lidas += 1
        resumo["total_protocolos"] += 1
        matriculas_lidas.add(
            resolve_matricula_inspetor(_cell_at(row, column_map, "matricula_inspetor"))
        )
        descricao_raw = _cell_at(row, column_map, "descricao_irregularidades")
        descricao = cell_to_str(descricao_raw)
        data_contestacao_raw = _cell_at(row, column_map, "data_contestacao")
        data_resposta_raw = _cell_at(row, column_map, "data_resposta")
        matricula_inspetor = resolve_matricula_inspetor(
            _cell_at(row, column_map, "matricula_inspetor")
        )
        data_contestacao = parse_datetime_cell(data_contestacao_raw)

        excecao = is_irregularidade_ic668_co49(descricao_raw)

        elegivel, motivo = avaliar_elegibilidade_reinspecao(
            descricao_irregularidades=descricao_raw,
            data_contestacao=data_contestacao_raw,
            data_resposta=data_resposta_raw,
        )
        if descricao and data_contestacao is not None:
            observations.append(
                ReinspecaoGedObservation(
                    row=ParsedReinspecaoRow(
                        protocolo=protocolo,
                        matricula_inspetor=matricula_inspetor,
                        descricao_irregularidades=descricao,
                        data_contestacao=data_contestacao,
                        excel_row=excel_row,
                    ),
                    ged_elegivel=elegivel,
                    motivo=motivo,
                )
            )
        linhas_resumo.append(
            LinhaResumoPlanilha(
                protocolo=protocolo,
                excel_row=excel_row,
                excecao=excecao,
                elegivel=elegivel,
                motivo=motivo,
                matricula_inspetor=matricula_inspetor,
                descricao=descricao,
                data=data_contestacao.isoformat() if data_contestacao else None,
            )
        )
        if excecao and elegivel and motivo == "data_resposta_espelho":
            resumo["total_excecoes"] += 1
        if not elegivel:
            linhas_filtradas += 1
            if motivo == "respondida":
                resumo["total_respondidos"] += 1
            elif motivo == "irregularidade_vazia":
                resumo["total_irregularidades_vazias"] += 1
            continue

        parsed_rows.append(
            ParsedReinspecaoRow(
                protocolo=protocolo,
                matricula_inspetor=matricula_inspetor,
                descricao_irregularidades=descricao,
                data_contestacao=data_contestacao,
                excel_row=excel_row,
            )
        )

    resumo["total_agentes"] = len(matriculas_lidas)

    from apps.auditoria.services.reinspecao_import_dedupe import (
        apply_reinspecao_import_dedupe,
        ignorados_to_preview,
    )

    import_rows: list[ParsedReinspecaoRow] = []
    dup_arquivo: list[Any] = []
    ignorados_portal: list[Any] = []
    if parsed_rows:
        import_rows, dup_arquivo, ignorados_portal = apply_reinspecao_import_dedupe(parsed_rows)

    registros_importacao = build_registros_planilha_preview(
        linhas_resumo,
        import_rows=import_rows,
        duplicados_arquivo=dup_arquivo,
        ignorados_portal=ignorados_portal,
    )

    if not parsed_rows:
        return ReinspecaoImportPreview(
            sheet_name=sheet_name,
            header_row=header_row,
            total_falhas=0,
            total_lidas=total_lidas,
            linhas_filtradas=linhas_filtradas,
            resumo_planilha=resumo,
            registros_importacao=registros_importacao,
            observations=observations,
            errors=(
                []
                if allow_no_eligible and observations
                else [
                    "Nenhuma linha válida após os filtros "
                    "(irregularidade vazia ou registro considerado respondido)."
                ]
            ),
        )

    dedupe_preview = ignorados_to_preview(dup_arquivo, ignorados_portal)

    return ReinspecaoImportPreview(
        sheet_name=sheet_name,
        header_row=header_row,
        total_falhas=len(import_rows),
        total_lidas=total_lidas,
        linhas_filtradas=linhas_filtradas,
        rows=import_rows,
        protocolos_duplicados_arquivo=dedupe_preview["protocolos_duplicados_arquivo"],
        protocolos_ignorados=dedupe_preview["protocolos_ignorados"],
        total_ignorados_duplicata=dedupe_preview["total_ignorados_duplicata"],
        registros_importacao=registros_importacao,
        resumo_planilha=resumo,
        observations=observations,
    )


def parse_reinspecao_csv(file_bytes: bytes) -> ReinspecaoImportPreview:
    from apps.auditoria.services.import_tabular import decode_csv_rows

    try:
        rows = decode_csv_rows(file_bytes)
    except Exception as exc:
        return ReinspecaoImportPreview(
            sheet_name="",
            header_row=0,
            total_falhas=0,
            errors=[f"Não foi possível ler o arquivo CSV: {exc}"],
        )
    return _parse_reinspecao_rows(rows, sheet_name="CSV")


def parse_reinspecao_workbook(file_bytes: bytes) -> ReinspecaoImportPreview:
    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as exc:
        return ReinspecaoImportPreview(
            sheet_name="",
            header_row=0,
            total_falhas=0,
            errors=[f"Não foi possível ler o arquivo Excel: {exc}"],
        )

    sheet_name = workbook.sheetnames[0] if workbook.sheetnames else ""
    if not sheet_name:
        return ReinspecaoImportPreview(
            sheet_name="",
            header_row=0,
            total_falhas=0,
            errors=["O arquivo não possui abas."],
        )

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()
    return _parse_reinspecao_rows(rows, sheet_name=sheet_name)


def parse_reinspecao_file(file_bytes: bytes, filename: str = "") -> ReinspecaoImportPreview:
    from apps.auditoria.services.import_tabular import is_csv_filename

    if is_csv_filename(filename):
        return parse_reinspecao_csv(file_bytes)
    return parse_reinspecao_workbook(file_bytes)


def rows_from_reinspecao_preview_payload(preview: dict[str, Any] | None) -> list[ParsedReinspecaoRow]:
    if not preview:
        return []
    rows_payload = preview.get("rows") or []
    result: list[ParsedReinspecaoRow] = []
    for item in rows_payload:
        if not isinstance(item, dict):
            continue
        protocolo = cell_to_str(item.get("protocolo"))
        matricula = resolve_matricula_inspetor(item.get("matricula_inspetor"))
        descricao = cell_to_str(item.get("descricao_irregularidades"))
        if not protocolo or not descricao:
            continue
        data_raw = item.get("data_contestacao")
        data_contestacao = None
        if isinstance(data_raw, datetime):
            data_contestacao = data_raw
        elif data_raw:
            data_contestacao = parse_datetime_cell(data_raw)
        analise_raw = item.get("data_analise")
        data_analise = None
        if isinstance(analise_raw, datetime):
            data_analise = analise_raw
        elif analise_raw:
            data_analise = parse_datetime_cell(analise_raw)
        result.append(
            ParsedReinspecaoRow(
                protocolo=protocolo,
                matricula_inspetor=matricula,
                descricao_irregularidades=descricao,
                data_contestacao=data_contestacao,
                data_analise=data_analise,
                excel_row=int(item.get("excel_row") or 0),
                nome_inspetor=cell_to_str(
                    item.get("nome_inspetor") or item.get("nome_agente") or ""
                ),
                codigo_irregularidade=cell_to_str(item.get("codigo_irregularidade")),
                classificacao_irregularidade=cell_to_str(
                    item.get("classificacao_irregularidade")
                ),
                cenario_mapeado=cell_to_str(item.get("cenario_mapeado")),
                etapa_mapeada=cell_to_str(item.get("etapa_mapeada")),
                mapping_scenario_status=cell_to_str(
                    item.get("mapping_scenario_status")
                ),
                mapping_stage_status=cell_to_str(item.get("mapping_stage_status")),
                mapping_version=cell_to_str(item.get("mapping_version")),
                mapping_source_hash=cell_to_str(item.get("mapping_source_hash")),
                mapping_candidates=(
                    item.get("mapping_candidates")
                    if isinstance(item.get("mapping_candidates"), list)
                    else []
                ),
            )
        )
    return result
