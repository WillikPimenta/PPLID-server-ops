from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

SHEET_TARGET = "Contestação de Resultado"
MAX_PREVIEW_ROWS = 10

HEADER_TO_FIELD = {
    "PROTOCOLO": "protocolo",
    "WORKFLOW": "workflow",
    "NIVEL HIERARQUICO": "nivel_hierarquico",
    "MATRICULA": "matricula",
    "N DO CONTRATO": "numero_contrato",
    "NUMERO DO CONTRATO": "numero_contrato",
    "RESULTADO CONTESTADO": "resultado_contestado",
    "RESULTADO POS AUDITORIA": "resultado_pos_auditoria",
    "TIPO DE CONCLUSAO": "tipo_conclusao",
    "TIPO DE FALHA": "tipo_falha",
    "CENARIO": "cenario",
    "DETALHAMENTO": "detalhamento",
    "CONCLUSAO DA CONTESTACAO": "conclusao_contestacao",
    "DATA DE ANALISE": "data_analise",
    "DATA ANALISE": "data_analise",
}

REQUIRED_FIELDS = ("protocolo", "workflow", "nivel_hierarquico", "resultado_contestado")
IGNORED_PROTOCOL_LABELS = frozenset({"CONSIDERACOES"})


@dataclass
class ParsedProtocolRow:
    protocolo: str
    workflow: str
    nivel_hierarquico: str
    resultado_contestado: str
    numero_contrato: str = ""
    resultado_pos_auditoria: str = ""
    tipo_conclusao: str = ""
    tipo_falha: str = ""
    cenario: str = ""
    detalhamento: str = ""
    conclusao_contestacao: str = ""
    data_analise: str = ""
    excel_row: int = 0


@dataclass
class ContestacaoImportPreview:
    sheet_name: str
    header_row: int
    total_protocolos: int
    protocolos_duplicados: list[str] = field(default_factory=list)
    protocolos_ja_tratados: list[dict[str, Any]] = field(default_factory=list)
    linhas_ignoradas: int = 0
    workflow: str = ""
    nivel_hierarquico: str = ""
    cliente: str = ""
    link_demanda: str = ""
    amostra: list[dict[str, Any]] = field(default_factory=list)
    rows: list[ParsedProtocolRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheet_name": self.sheet_name,
            "header_row": self.header_row,
            "total_protocolos": self.total_protocolos,
            "protocolos_duplicados": self.protocolos_duplicados,
            "protocolos_ja_tratados": self.protocolos_ja_tratados,
            "total_protocolos_ja_tratados": len(self.protocolos_ja_tratados),
            "cliente": self.cliente,
            "link_demanda": self.link_demanda,
            "amostra": self.amostra,
            "errors": self.errors,
        }


def parse_import_filename_metadata(filename: str) -> dict[str, str]:
    base = (filename or "").rsplit(".", 1)[0].strip()
    parts = [part.strip() for part in base.split(" - ") if part.strip()]
    cliente = ""
    if len(parts) >= 3:
        cliente = parts[1]
    elif len(parts) == 2:
        cliente = parts[1]
    return {"cliente": cliente, "link_demanda": ""}


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.upper()
    text = text.replace("º", "O").replace("°", "O").replace("Nº", "N")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def is_ignored_protocol_row(protocolo: str) -> bool:
    return normalize_text(protocolo) in IGNORED_PROTOCOL_LABELS


def find_target_sheet(sheet_names: list[str]) -> str | None:
    for name in sheet_names:
        normalized = normalize_text(name)
        if normalized == normalize_text(SHEET_TARGET):
            return name
    for name in sheet_names:
        normalized = normalize_text(name)
        if "CONTESTACAO" in normalized and "RESULTADO" in normalized:
            return name
    return None


def find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int]] | None:
    for row_index, row in enumerate(rows, start=1):
        column_map: dict[str, int] = {}
        for col_index, cell in enumerate(row):
            header_key = normalize_text(cell)
            field_name = HEADER_TO_FIELD.get(header_key)
            if field_name:
                column_map[field_name] = col_index
        if "protocolo" in column_map:
            return row_index, column_map
    return None


def _parse_contestacao_rows(
    rows: list[tuple[Any, ...]],
    *,
    sheet_name: str,
    empty_error: str,
) -> ContestacaoImportPreview:
    preview = ContestacaoImportPreview(sheet_name=sheet_name, header_row=0, total_protocolos=0)

    header_info = find_header_row(rows)
    if not header_info:
        preview.errors.append('Não foi possível localizar a linha de cabeçalho com a coluna "PROTOCOLO".')
        return preview

    header_row, column_map = header_info
    preview.header_row = header_row

    missing_columns = [field for field in REQUIRED_FIELDS if field not in column_map]
    if missing_columns:
        labels = ", ".join(field.replace("_", " ").upper() for field in missing_columns)
        preview.errors.append(f"Colunas obrigatórias ausentes no cabeçalho: {labels}.")
        return preview

    parsed_rows: list[ParsedProtocolRow] = []
    ignored_rows = 0

    for row_index, row in enumerate(rows[header_row:], start=header_row + 1):
        protocolo = cell_to_str(row[column_map["protocolo"]] if column_map["protocolo"] < len(row) else "")
        if not protocolo or is_ignored_protocol_row(protocolo):
            ignored_rows += 1
            continue

        def get_field(field_name: str) -> str:
            col = column_map.get(field_name)
            if col is None or col >= len(row):
                return ""
            return cell_to_str(row[col])

        parsed_rows.append(
            ParsedProtocolRow(
                protocolo=protocolo,
                workflow=get_field("workflow"),
                nivel_hierarquico=get_field("nivel_hierarquico"),
                resultado_contestado=get_field("resultado_contestado"),
                numero_contrato=get_field("numero_contrato"),
                resultado_pos_auditoria=get_field("resultado_pos_auditoria"),
                tipo_conclusao=get_field("tipo_conclusao"),
                tipo_falha=get_field("tipo_falha"),
                cenario=get_field("cenario"),
                detalhamento=get_field("detalhamento"),
                conclusao_contestacao=get_field("conclusao_contestacao"),
                data_analise=get_field("data_analise"),
                excel_row=row_index,
            )
        )

    preview.linhas_ignoradas = ignored_rows
    preview.rows = parsed_rows
    preview.total_protocolos = len(parsed_rows)

    counts: dict[str, int] = {}
    for item in parsed_rows:
        counts[item.protocolo] = counts.get(item.protocolo, 0) + 1
    preview.protocolos_duplicados = sorted([protocol for protocol, count in counts.items() if count > 1])

    if parsed_rows:
        preview.workflow = parsed_rows[0].workflow
        preview.nivel_hierarquico = parsed_rows[0].nivel_hierarquico

    preview.amostra = [
        {
            "protocolo": item.protocolo,
            "workflow": item.workflow,
            "nivel_hierarquico": item.nivel_hierarquico,
            "resultado_contestado": item.resultado_contestado,
            "excel_row": item.excel_row,
        }
        for item in parsed_rows[:MAX_PREVIEW_ROWS]
    ]

    if not parsed_rows:
        preview.errors.append(empty_error)

    return preview


def parse_contestacao_csv(file_bytes: bytes) -> ContestacaoImportPreview:
    from apps.auditoria.services.import_tabular import decode_csv_rows

    try:
        rows = decode_csv_rows(file_bytes)
    except Exception:
        preview = ContestacaoImportPreview(sheet_name="", header_row=0, total_protocolos=0)
        preview.errors.append("Não foi possível ler o arquivo CSV.")
        return preview

    return _parse_contestacao_rows(
        rows,
        sheet_name="CSV",
        empty_error="Nenhum protocolo preenchido foi encontrado no arquivo CSV.",
    )


def parse_contestacao_workbook(file_bytes: bytes) -> ContestacaoImportPreview:
    preview = ContestacaoImportPreview(sheet_name="", header_row=0, total_protocolos=0)

    try:
        workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        preview.errors.append("Não foi possível ler o arquivo Excel.")
        return preview

    sheet_name = find_target_sheet(workbook.sheetnames)
    if not sheet_name:
        preview.errors.append('A aba "Contestação de Resultado" não foi encontrada no arquivo.')
        workbook.close()
        return preview

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    return _parse_contestacao_rows(
        rows,
        sheet_name=sheet_name,
        empty_error="Nenhum protocolo preenchido foi encontrado na aba de contestação.",
    )


def parse_contestacao_file(file_bytes: bytes, filename: str = "") -> ContestacaoImportPreview:
    from apps.auditoria.services.import_tabular import is_csv_filename

    if is_csv_filename(filename):
        return parse_contestacao_csv(file_bytes)
    return parse_contestacao_workbook(file_bytes)


def rows_from_preview_payload(payload: dict[str, Any]) -> list[ParsedProtocolRow]:
    rows = payload.get("rows") or []
    parsed: list[ParsedProtocolRow] = []
    for item in rows:
        if not isinstance(item, dict):
            continue
        parsed.append(
            ParsedProtocolRow(
                protocolo=str(item.get("protocolo") or "").strip(),
                workflow=str(item.get("workflow") or "").strip(),
                nivel_hierarquico=str(item.get("nivel_hierarquico") or "").strip(),
                resultado_contestado=str(item.get("resultado_contestado") or "").strip(),
                numero_contrato=str(item.get("numero_contrato") or "").strip(),
                resultado_pos_auditoria=str(item.get("resultado_pos_auditoria") or "").strip(),
                tipo_conclusao=str(item.get("tipo_conclusao") or "").strip(),
                tipo_falha=str(item.get("tipo_falha") or "").strip(),
                cenario=str(item.get("cenario") or "").strip(),
                detalhamento=str(item.get("detalhamento") or "").strip(),
                conclusao_contestacao=str(item.get("conclusao_contestacao") or "").strip(),
                data_analise=str(item.get("data_analise") or "").strip(),
                excel_row=int(item.get("excel_row") or 0),
            )
        )
    return [row for row in parsed if row.protocolo and not is_ignored_protocol_row(row.protocolo)]
