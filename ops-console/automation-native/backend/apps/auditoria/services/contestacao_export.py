from __future__ import annotations

import re
from copy import copy
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.cell_range import CellRange

from apps.auditoria.models import (
    AuditoriaAtividade,
    AuditoriaAtividadeProtocolo,
    AuditoriaAtividadeProtocoloEtapa,
    AuditoriaFalhaCadastro,
)
from apps.auditoria.services.contestacao_import import (
    find_header_row,
    find_target_sheet,
    normalize_text,
)
from apps.auditoria.services.contestacao_duplicada import (
    append_duplicado_note_to_detalhamento,
    format_duplicado_detalhamento_note,
)
from apps.auditoria.services.qualidade_promocao import tratados_qs
from apps.auditoria.services.text_format import is_tipo_falha_automatico

RETORNO_TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "retorno_contestacao.xlsx"

ROW_FIELDS = (
    "protocolo",
    "workflow",
    "nivel_hierarquico",
    "matricula",
    "numero_contrato",
    "resultado_contestado",
    "resultado_pos_auditoria",
    "tipo_conclusao",
    "tipo_falha",
    "cenario",
    "detalhamento",
    "conclusao_contestacao",
)


def build_retorno_filename(original_name: str) -> str:
    base = (original_name or "retorno").rsplit(".", 1)[0].strip() or "retorno"
    if re.search(r"_Retorno_IDF$", base, flags=re.IGNORECASE):
        return f"{base}.xlsx"
    return f"{base}_Retorno_IDF.xlsx"


def _export_upper(value: str) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    return text.upper()


def _export_tipo_conclusao(value: str, *, fallback_tipo_falha: str = "") -> str:
    raw = (value or "").strip()
    if not raw and fallback_tipo_falha:
        if is_tipo_falha_automatico(fallback_tipo_falha):
            return "AUTOMÁTICO"
        return "MANUAL"
    if not raw:
        return ""
    key = normalize_text(raw)
    if "AUTOMAT" in key:
        return "AUTOMÁTICO"
    if "MANUAL" in key:
        return "MANUAL"
    if "PROCESSUAL" in key:
        return "PROCESSUAL"
    return _export_upper(raw)


def _primary_etapa(
    etapas: list[AuditoriaAtividadeProtocoloEtapa],
) -> AuditoriaAtividadeProtocoloEtapa | None:
    procedente = next(
        (item for item in etapas if item.situacao == AuditoriaAtividadeProtocoloEtapa.SITUACAO_PROCEDENTE),
        None,
    )
    if procedente is not None:
        return procedente
    return next((item for item in etapas if (item.situacao or "").strip()), None)


def derive_retorno_values(protocolo: AuditoriaAtividadeProtocolo) -> dict[str, str]:
    """Monta valores de retorno a partir da análise preenchida na intranet."""
    etapas = list(protocolo.etapas.all().order_by("ordem"))
    primary = _primary_etapa(etapas)

    situacao = (protocolo.situacao or "").strip().lower()
    if not situacao and primary is not None:
        situacao = (primary.situacao or "").strip().lower()

    conclusao = ""
    if situacao == AuditoriaAtividadeProtocolo.SITUACAO_PROCEDENTE:
        conclusao = "PROCEDENTE"
    elif situacao == AuditoriaAtividadeProtocolo.SITUACAO_IMPROCEDENTE:
        conclusao = "IMPROCEDENTE"

    resultado_pos = ""
    tipo_falha_etapa = ""
    cenario = ""
    matricula = ""
    if primary is not None:
        resultado_pos = (primary.resultado_correto or "").strip()
        tipo_falha_etapa = (primary.tipo_falha or "").strip()
        cenario = (primary.motivo_falha or "").strip()
        matricula = (primary.agente or "").strip()

    if not matricula:
        matricula = next(
            (
                (item.agente or "").strip()
                for item in etapas
                if (item.agente or "").strip()
            ),
            "",
        )

    if not resultado_pos:
        resultado_pos = (protocolo.resultado_contestado or "").strip()
    if not resultado_pos:
        resultado_pos = (protocolo.resultado_pos_auditoria or "").strip()

    if situacao == AuditoriaAtividadeProtocolo.SITUACAO_IMPROCEDENTE:
        tipo_falha_export = "SEM FALHA"
    elif tipo_falha_etapa:
        tipo_falha_export = _export_upper(tipo_falha_etapa)
    else:
        tipo_falha_export = _export_upper(protocolo.tipo_falha)

    if not cenario:
        cenario = (protocolo.cenario or "").strip()

    detalhamento = (protocolo.consideracoes_finais or "").strip()
    if not detalhamento:
        detalhamento = (protocolo.detalhamento or "").strip()

    tipo_conclusao = _export_tipo_conclusao(
        protocolo.tipo_conclusao,
        fallback_tipo_falha=tipo_falha_etapa,
    )

    return {
        "numero_contrato": (protocolo.numero_contrato or "").strip(),
        "matricula": matricula,
        "resultado_pos_auditoria": resultado_pos,
        "tipo_conclusao": tipo_conclusao,
        "tipo_falha": tipo_falha_export,
        "cenario": cenario,
        "detalhamento": detalhamento,
        "conclusao_contestacao": conclusao or (protocolo.conclusao_contestacao or "").strip(),
    }


def build_row_values(protocolo: AuditoriaAtividadeProtocolo) -> dict[str, str]:
    if protocolo.tratado_referencia_id:
        tratado = protocolo.tratado_referencia
        if tratado is None:
            tratado = (
                tratados_qs(origem=AuditoriaFalhaCadastro.ORIGEM_CONTESTACAO)
                .filter(pk=protocolo.tratado_referencia_id)
                .select_related("atividade")
                .first()
            )
        if tratado is not None:
            row = build_row_values_from_tratado(tratado)
            note = format_duplicado_detalhamento_note(tratado)
            row["detalhamento"] = append_duplicado_note_to_detalhamento(row.get("detalhamento", ""), note)
            row["protocolo"] = (protocolo.protocolo or "").strip()
            row["workflow"] = (protocolo.workflow or row.get("workflow") or "").strip()
            row["nivel_hierarquico"] = (protocolo.nivel_hierarquico or row.get("nivel_hierarquico") or "").strip()
            row["resultado_contestado"] = (protocolo.resultado_contestado or row.get("resultado_contestado") or "").strip()
            return row

    retorno = derive_retorno_values(protocolo)
    return {
        "protocolo": (protocolo.protocolo or "").strip(),
        "workflow": (protocolo.workflow or "").strip(),
        "nivel_hierarquico": (protocolo.nivel_hierarquico or "").strip(),
        "resultado_contestado": (protocolo.resultado_contestado or "").strip(),
        **retorno,
    }


def build_row_values_from_tratado(falha: AuditoriaFalhaCadastro) -> dict[str, str]:
    """Monta linha de retorno a partir da etapa tratada canônica."""
    parsed = falha.brflow_parsed if isinstance(falha.brflow_parsed, dict) else {}
    situacao = (falha.status or parsed.get("situacao") or "").strip().lower()

    conclusao = ""
    if situacao == "procedente":
        conclusao = "PROCEDENTE"
    elif situacao == "improcedente":
        conclusao = "IMPROCEDENTE"

    resultado_pos = (falha.novo_resultado or "").strip()
    tipo_falha_etapa = (falha.tipo_falha or "").strip()
    cenario = (falha.motivo_falha or "").strip()
    if not resultado_pos:
        resultado_pos = (falha.resultado_cliente or "").strip()

    if situacao == "improcedente":
        tipo_falha_export = "SEM FALHA"
    elif tipo_falha_etapa:
        tipo_falha_export = _export_upper(tipo_falha_etapa)
    else:
        tipo_falha_export = _export_upper(falha.tipo_falha)

    detalhamento = (falha.observacao or "").strip()
    if not detalhamento:
        detalhamento = (parsed.get("consideracoes_finais") or falha.descricao_irregularidades or "").strip()

    tipo_conclusao = _export_tipo_conclusao(
        str(parsed.get("tipo_conclusao") or ""),
        fallback_tipo_falha=tipo_falha_etapa,
    )

    return {
        "protocolo": (falha.protocolo or "").strip(),
        "workflow": str(parsed.get("workflow") or "").strip(),
        "nivel_hierarquico": str(parsed.get("nivel_hierarquico") or "").strip(),
        "matricula": (falha.usuario or "").strip(),
        "resultado_contestado": (falha.resultado_cliente or "").strip(),
        "numero_contrato": str(parsed.get("numero_contrato") or "").strip(),
        "resultado_pos_auditoria": resultado_pos,
        "tipo_conclusao": tipo_conclusao,
        "tipo_falha": tipo_falha_export,
        "cenario": cenario,
        "detalhamento": detalhamento,
        "conclusao_contestacao": conclusao,
    }


def _find_consideracoes_row(worksheet, *, start_row: int) -> int | None:
    for row_index in range(start_row, worksheet.max_row + 1):
        value = worksheet.cell(row=row_index, column=2).value
        if normalize_text(value) == "CONSIDERACOES":
            return row_index
    return None


def _copy_row_style(worksheet, *, source_row: int, target_row: int, min_col: int, max_col: int) -> None:
    for col in range(min_col, max_col + 1):
        source = worksheet.cell(row=source_row, column=col)
        target = worksheet.cell(row=target_row, column=col)
        if source.has_style:
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)


def _insert_rows_preserving_merged_ranges(worksheet, *, row: int, amount: int) -> None:
    shifted_ranges: list[CellRange] = []
    for merged_range in list(worksheet.merged_cells.ranges):
        if merged_range.min_row < row:
            continue
        shifted = CellRange(str(merged_range))
        shifted.shift(row_shift=amount)
        shifted_ranges.append(shifted)
        worksheet.unmerge_cells(str(merged_range))

    worksheet.insert_rows(row, amount=amount)

    for shifted in shifted_ranges:
        worksheet.merge_cells(str(shifted))


def _extend_column_conditional_formatting(
    worksheet,
    *,
    column: int,
    first_data_row: int,
    last_data_row: int,
) -> None:
    if last_data_row < first_data_row:
        return

    for conditional, existing_rules in list(worksheet.conditional_formatting._cf_rules.items()):
        ranges = [CellRange(str(cell_range)) for cell_range in conditional.sqref.ranges]
        changed = False
        for cell_range in ranges:
            if (
                cell_range.min_col <= column <= cell_range.max_col
                and cell_range.min_row <= first_data_row <= cell_range.max_row
                and cell_range.max_row < last_data_row
            ):
                cell_range.max_row = last_data_row
                changed = True
        if not changed:
            continue

        rules = list(existing_rules)
        del worksheet.conditional_formatting._cf_rules[conditional]
        sqref = " ".join(str(cell_range) for cell_range in ranges)
        for rule in rules:
            worksheet.conditional_formatting.add(sqref, rule)


def _format_data_table(
    worksheet,
    *,
    header_row: int,
    last_data_row: int,
    min_col: int,
    max_col: int,
) -> None:
    for row_number in range(header_row, last_data_row + 1):
        worksheet.row_dimensions[row_number].height = None
        for col_number in range(min_col, max_col + 1):
            cell = worksheet.cell(row=row_number, column=col_number)
            alignment = copy(cell.alignment)
            alignment.horizontal = "left"
            alignment.wrap_text = True
            cell.alignment = alignment


def _ensure_data_capacity(worksheet, *, header_row: int, protocol_count: int, column_map: dict[str, int]) -> None:
    first_data_row = header_row + 1
    consideracoes_row = _find_consideracoes_row(worksheet, start_row=first_data_row)
    if consideracoes_row is None:
        return

    # Mantém uma linha em branco antes de CONSIDERAÇÕES.
    available = max(0, consideracoes_row - first_data_row - 1)
    if protocol_count <= available:
        return

    extra = protocol_count - available
    _insert_rows_preserving_merged_ranges(
        worksheet,
        row=consideracoes_row,
        amount=extra,
    )

    style_source = first_data_row
    min_col = min(column_map.values()) + 1
    max_col = max(column_map.values()) + 1
    for offset in range(available, protocol_count):
        _copy_row_style(
            worksheet,
            source_row=style_source,
            target_row=first_data_row + offset,
            min_col=min_col,
            max_col=max_col,
        )

    tipo_col = column_map.get("tipo_conclusao")
    if tipo_col is not None:
        last_data_row = first_data_row + protocol_count - 1
        col_letter = get_column_letter(tipo_col + 1)
        sqref = f"{col_letter}{first_data_row}:{col_letter}{last_data_row}"
        worksheet.data_validations.dataValidation = [
            dv
            for dv in worksheet.data_validations.dataValidation
            if col_letter not in str(dv.sqref)
        ]
        validation = DataValidation(
            type="list",
            formula1='"AUTOMÁTICO,MANUAL"',
            allow_blank=True,
            showDropDown=False,
            showInputMessage=True,
            showErrorMessage=True,
        )
        validation.add(sqref)
        worksheet.add_data_validation(validation)


def _write_protocol_row(
    worksheet,
    *,
    row_number: int,
    column_map: dict[str, int],
    values: dict[str, str],
) -> None:
    for field in ROW_FIELDS:
        col_index = column_map.get(field)
        if col_index is None:
            continue
        value = (values.get(field) or "").strip()
        worksheet.cell(row=row_number, column=col_index + 1).value = value or None


def _ensure_matricula_column(worksheet, *, header_row: int, column_map: dict[str, int]) -> None:
    """Garante a coluna de matricula em templates antigos sem alterar as colunas existentes."""
    if "matricula" in column_map:
        return

    source_column = max(column_map.values()) + 1
    target_column = source_column + 1
    source_letter = get_column_letter(source_column)
    target_letter = get_column_letter(target_column)

    for row_number in range(header_row, worksheet.max_row + 1):
        source = worksheet.cell(row=row_number, column=source_column)
        target = worksheet.cell(row=row_number, column=target_column)
        if source.has_style:
            target.font = copy(source.font)
            target.border = copy(source.border)
            target.fill = copy(source.fill)
            target.number_format = source.number_format
            target.protection = copy(source.protection)
            target.alignment = copy(source.alignment)

    worksheet.cell(row=header_row, column=target_column).value = "MATRÍCULA"
    worksheet.column_dimensions[target_letter].width = (
        worksheet.column_dimensions[source_letter].width or 18
    )
    column_map["matricula"] = target_column - 1


def load_retorno_template() -> Any:
    if not RETORNO_TEMPLATE_PATH.exists():
        raise ValueError("Template de retorno da contestação não encontrado no repositório.")
    try:
        return load_workbook(filename=RETORNO_TEMPLATE_PATH)
    except Exception as exc:
        raise ValueError("Não foi possível ler o template de retorno da contestação.") from exc


def export_atividade_retorno(atividade: AuditoriaAtividade) -> tuple[bytes, str]:
    """
    Gera a planilha de retorno sempre a partir do template oficial em branco,
    preenchendo as linhas com os dados da atividade/análise.
    """
    protocolos = list(
        AuditoriaAtividadeProtocolo.objects.filter(atividade=atividade)
        .select_related("tratado_referencia", "tratado_referencia__atividade")
        .prefetch_related("etapas")
        .order_by("excel_row", "id")
    )
    tratados = list(
        tratados_qs(origem=AuditoriaFalhaCadastro.ORIGEM_CONTESTACAO)
        .filter(atividade=atividade)
        .order_by("created_at", "id")
    )
    row_values = [build_row_values(item) for item in protocolos]
    protocolo_ids = {item.id for item in protocolos}
    # Tratados novos já possuem seu protocolo histórico. Acrescenta somente
    # registros legados cujo intermediário não existe mais.
    legados_por_protocolo: dict[str, list[AuditoriaFalhaCadastro]] = {}
    for item in tratados:
        parsed = item.brflow_parsed if isinstance(item.brflow_parsed, dict) else {}
        vinculado_id = item.protocolo_origem_id or parsed.get("protocolo_origem_id")
        if vinculado_id in protocolo_ids:
            continue
        legados_por_protocolo.setdefault((item.protocolo or "").strip(), []).append(item)
    for grupo in legados_por_protocolo.values():
        canonical = next(
            (item for item in grupo if (item.status or "").strip().lower() == "procedente"),
            grupo[0],
        )
        row_values.append(build_row_values_from_tratado(canonical))
    filename = build_retorno_filename(atividade.nome or atividade.nome_arquivo_original or "retorno")

    workbook = load_retorno_template()
    sheet_name = find_target_sheet(workbook.sheetnames)
    if not sheet_name:
        raise ValueError('A aba "Contestação de Resultado" não foi encontrada no template.')

    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    header_info = find_header_row(rows)
    if not header_info:
        raise ValueError('Não foi possível localizar o cabeçalho com a coluna "PROTOCOLO".')

    header_row, column_map = header_info
    _ensure_matricula_column(
        worksheet,
        header_row=header_row,
        column_map=column_map,
    )
    _ensure_data_capacity(
        worksheet,
        header_row=header_row,
        protocol_count=len(row_values),
        column_map=column_map,
    )

    for offset, values in enumerate(row_values):
        _write_protocol_row(
            worksheet,
            row_number=header_row + 1 + offset,
            column_map=column_map,
            values=values,
        )

    first_data_row = header_row + 1
    last_data_row = header_row + len(row_values)
    min_col = min(column_map.values()) + 1
    max_col = max(column_map.values()) + 1
    conclusao_col = column_map.get("conclusao_contestacao")
    if conclusao_col is not None:
        _extend_column_conditional_formatting(
            worksheet,
            column=conclusao_col + 1,
            first_data_row=first_data_row,
            last_data_row=last_data_row,
        )
    _format_data_table(
        worksheet,
        header_row=header_row,
        last_data_row=max(header_row, last_data_row),
        min_col=min_col,
        max_col=max_col,
    )

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), filename
