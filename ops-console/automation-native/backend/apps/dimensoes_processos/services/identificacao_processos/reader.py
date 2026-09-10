"""Leitura do workbook Identificação dos Processos."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

DEFAULT_XLSX = (
    Path(__file__).resolve().parents[4]
    / "data"
    / "dimensoes_processos"
    / "identificacao_processos.xlsx"
)


def _sheet_names(wb) -> dict[str, str]:
    out: dict[str, str] = {}
    for name in wb.sheetnames:
        key = (
            name.lower()
            .replace("ç", "c")
            .replace("ã", "a")
            .replace("á", "a")
            .replace("é", "e")
            .replace("í", "i")
            .replace("ó", "o")
            .replace("ú", "u")
            .replace(" ", "_")
        )
        out[key] = name
    return out


def find_sheet(names: dict[str, str], *candidates: str, required: bool = True) -> str | None:
    for candidate in candidates:
        key = candidate.lower().replace(" ", "_")
        if key in names:
            return names[key]
    if required:
        raise KeyError(f"Aba não encontrada. Candidatos={candidates} disponíveis={list(names.values())}")
    return None


def _cell(row: tuple, idx: int | None):
    if idx is None or idx < 0 or idx >= len(row):
        return None
    return row[idx]


def as_str(val: Any, default: str = "") -> str:
    if val is None:
        return default
    return str(val).strip()


def as_int(val: Any) -> int | None:
    if val is None or val == "":
        return None
    try:
        return int(float(val))
    except (TypeError, ValueError):
        return None


def as_float(val: Any) -> float | None:
    if val is None or val == "":
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def as_decimal(val: Any) -> Decimal | None:
    parsed = as_float(val)
    if parsed is None:
        return None
    try:
        return Decimal(str(parsed))
    except (InvalidOperation, ValueError):
        return None


def as_bool01(val: Any, default: bool = False) -> bool:
    if val is None or val == "":
        return default
    if isinstance(val, bool):
        return val
    try:
        return bool(int(float(val)))
    except (TypeError, ValueError):
        text = str(val).strip().lower()
        if text in {"1", "true", "sim", "yes", "y"}:
            return True
        if text in {"0", "false", "nao", "não", "no", "n"}:
            return False
        return default


def as_date(val: Any) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    return None


def as_time(val: Any) -> time | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.time()
    if isinstance(val, time):
        return val
    return None


def norm_key(text: str) -> str:
    return (
        str(text)
        .strip()
        .lower()
        .replace("ç", "c")
        .replace("ã", "a")
        .replace("á", "a")
        .replace("à", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
        .replace("ü", "u")
    )


def header_index(ws) -> dict[str, int]:
    headers: dict[str, int] = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))):
        if cell is None:
            continue
        headers[norm_key(cell)] = idx
    return headers


def col(headers: dict[str, int], *names: str) -> int:
    for name in names:
        key = norm_key(name)
        if key in headers:
            return headers[key]
    raise KeyError(f"Coluna não encontrada: {names}")


def optional_col(headers: dict[str, int], *names: str) -> int | None:
    for name in names:
        key = norm_key(name)
        if key in headers:
            return headers[key]
    return None


@dataclass
class WorkbookData:
    sheet_names: list[str] = field(default_factory=list)
    produtos: list[dict[str, Any]] = field(default_factory=list)
    grupos_servico: list[dict[str, Any]] = field(default_factory=list)
    servicos: list[dict[str, Any]] = field(default_factory=list)
    clientes: list[dict[str, Any]] = field(default_factory=list)
    niveis_hierarquicos: list[dict[str, Any]] = field(default_factory=list)
    workflows: list[dict[str, Any]] = field(default_factory=list)
    etapas: list[dict[str, Any]] = field(default_factory=list)
    metas_etapa: list[dict[str, Any]] = field(default_factory=list)
    projecao_sla: list[dict[str, Any]] = field(default_factory=list)
    projecao_equipes: list[dict[str, Any]] = field(default_factory=list)


def _read_produtos(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "produto")
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id, c_tipo = col(headers, "id_produto"), col(headers, "tipo_produto")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        pid = as_int(_cell(row, c_id))
        if pid is None:
            continue
        rows.append({"id_produto": pid, "tipo_produto": as_str(_cell(row, c_tipo)) or str(pid)})
    return rows


def _read_grupos_servico(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "grupo_servico", "grupo_serviço", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id = col(headers, "id_servico", "id_grupo", "id_grupo_servico")
    c_nome = col(headers, "servico", "nome", "grupo")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        gid = as_int(_cell(row, c_id))
        if gid is None:
            continue
        rows.append({"id_grupo": gid, "nome": as_str(_cell(row, c_nome))})
    return rows


def _read_servicos(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "servico", "serviço", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id = col(headers, "id_servico")
    c_nome = col(headers, "servico")
    c_meta = col(headers, "meta dia", "meta_dia")
    c_grp = col(headers, "id_grupo_servico")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        sid = as_int(_cell(row, c_id))
        if sid is None:
            continue
        rows.append(
            {
                "id_servico": sid,
                "nome": as_str(_cell(row, c_nome)),
                "meta_dia": as_decimal(_cell(row, c_meta)),
                "id_grupo_servico": as_int(_cell(row, c_grp)),
            }
        )
    return rows


def _read_clientes(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "clientes", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id, c_nome = col(headers, "id_cliente"), col(headers, "nome_cliente")
    c_ops = col(headers, "operations")
    c_cls = col(headers, "id_classificacao")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = as_int(_cell(row, c_id))
        if cid is None:
            continue
        rows.append(
            {
                "id_cliente": cid,
                "nome": as_str(_cell(row, c_nome)),
                "operations": as_bool01(_cell(row, c_ops), True),
                "id_classificacao": as_int(_cell(row, c_cls)) or 0,
            }
        )
    return rows


def _read_niveis_hierarquicos(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "nivel_hierarquico", "nivel hierarquico", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id = col(headers, "id_nh")
    c_nome = col(headers, "nivel hierarquico", "nivel_hierarquico")
    c_ind = col(headers, "ind_considerar")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nid = as_int(_cell(row, c_id))
        if nid is None:
            continue
        rows.append(
            {
                "id_nh": nid,
                "nome": as_str(_cell(row, c_nome)),
                "ind_considerar": as_bool01(_cell(row, c_ind), True),
            }
        )
    return rows


def _read_workflows(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "workflow", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id = col(headers, "id_workflow")
    c_nome = col(headers, "workflow")
    c_ind = col(headers, "ind_considerar")
    c_prod = col(headers, "id_produto")
    c_tipo = col(headers, "tipo_atendimento")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        wid = as_int(_cell(row, c_id))
        if wid is None:
            continue
        rows.append(
            {
                "id_workflow": wid,
                "nome": as_str(_cell(row, c_nome)),
                "ind_considerar": as_bool01(_cell(row, c_ind), True),
                "id_produto": as_int(_cell(row, c_prod)),
                "tipo_atendimento": as_str(_cell(row, c_tipo)),
            }
        )
    return rows


def _read_etapas(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "etapas", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_id, c_nome = col(headers, "id_etapa"), col(headers, "nome_etapa")
    c_manual = optional_col(headers, "manual", "ind_manual", "tipo_analise")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        eid = as_int(_cell(row, c_id))
        if eid is None:
            continue
        if c_manual is None:
            manual = True
        else:
            manual_raw = _cell(row, c_manual)
            if isinstance(manual_raw, str):
                manual = manual_raw.strip().lower() in {"1", "true", "sim", "manual", "s", "yes"}
            elif manual_raw is None or manual_raw == "":
                manual = True
            else:
                manual = bool(manual_raw)
        rows.append({"id_etapa": eid, "nome": as_str(_cell(row, c_nome)), "manual": manual})
    return rows


def _read_metas_etapa(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "metas_etapa", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_ini = col(headers, "data_início", "data_inicio")
    c_fim = col(headers, "data_fim")
    c_et = col(headers, "id_etapa")
    c_meta = col(headers, "meta_dia")
    c_srv = col(headers, "id_serviço", "id_servico")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        eid = as_int(_cell(row, c_et))
        di = as_date(_cell(row, c_ini))
        meta = as_decimal(_cell(row, c_meta))
        if eid is None or di is None or meta is None:
            continue
        rows.append(
            {
                "data_inicio": di,
                "data_fim": as_date(_cell(row, c_fim)),
                "id_etapa": eid,
                "meta_dia": meta,
                "id_servico": as_int(_cell(row, c_srv)),
            }
        )
    return rows


def _read_projecao_sla(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "projecao_sla", "projeção_sla", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)
    c_cli = col(headers, "id_cliente")
    c_ini = col(headers, "data_início", "data_inicio")
    c_fim = col(headers, "data_fim")
    c_wf = col(headers, "id_workflow")
    c_nh = col(headers, "id_nivel_hierarquico")
    c_dias = col(headers, "dias_semana")
    c_hi = col(headers, "hora_início", "hora_inicio")
    c_hf = col(headers, "hora_fim")
    c_dur = col(headers, "duração_atendimento", "duracao_atendimento")
    c_sla = col(headers, "sla_segundos")
    c_flag = col(headers, "flag_ajuste_sla")
    c_aj = col(headers, "sla_ajuste")
    c_vol = col(headers, "volume")
    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid = as_int(_cell(row, c_cli))
        wid = as_int(_cell(row, c_wf))
        nid = as_int(_cell(row, c_nh))
        di = as_date(_cell(row, c_ini))
        if not cid or not wid or not nid or di is None:
            continue
        rows.append(
            {
                "id_cliente": cid,
                "data_inicio": di,
                "data_fim": as_date(_cell(row, c_fim)),
                "id_workflow": wid,
                "id_nivel_hierarquico": nid,
                "dias_semana": as_str(_cell(row, c_dias)),
                "hora_inicio": as_time(_cell(row, c_hi)),
                "hora_fim": as_time(_cell(row, c_hf)),
                "duracao_atendimento": as_float(_cell(row, c_dur)),
                "sla_segundos": as_int(_cell(row, c_sla)),
                "flag_ajuste_sla": as_float(_cell(row, c_flag)),
                "sla_ajuste": as_int(_cell(row, c_aj)),
                "volume": as_float(_cell(row, c_vol)),
            }
        )
    return rows


def _read_projecao_equipes(wb, names: dict[str, str]) -> list[dict[str, Any]]:
    sheet = find_sheet(names, "projecao_equipes", "projeção_equipes", required=False)
    if not sheet:
        return []
    ws = wb[sheet]
    headers = header_index(ws)

    def hc(*opts: str) -> int | None:
        for opt in opts:
            key = norm_key(opt)
            if key in headers:
                return headers[key]
        return None

    c_di = hc("data inicial", "data_inicial")
    c_df = hc("data final", "data_final")
    c_mat = hc("matrícula agente", "matricula agente", "matricula_agente")
    if c_di is None or c_mat is None:
        raise KeyError("Projeção_Equipes: colunas Data Inicial / Matrícula Agente obrigatórias")

    cols = {
        "nome_agente": hc("nome agente"),
        "email_agente": hc("e-mail agente", "email agente"),
        "atividade": hc("atividade"),
        "matricula_lider": hc("matrícula líder", "matricula lider"),
        "nome_lider": hc("nome líder", "nome lider"),
        "email_lider": hc("e-mail líder", "email lider"),
        "equipe": hc("equipe"),
        "matricula_facilitador": hc("matrícula facilitador", "matricula facilitador"),
        "id_operations": hc("id_operations"),
        "horario": hc("horário", "horario"),
        "turno": hc("turno"),
        "localidade": hc("localidade"),
        "data_admissao": hc("data de admissão", "data de admissao"),
        "funcao": hc("função", "funcao"),
        "setor": hc("setor"),
        "pcd": hc("pcd"),
        "desconto_meta": hc("desconto meta"),
        "matricula_oracle": hc("matricula oracle"),
        "matricula_ponto": hc("matrícula ponto", "matricula ponto"),
        "observacao": hc("observação", "observacao"),
        "matricula_supervisor": hc("matrícula supervisor", "matricula supervisor"),
        "supervisor": hc("supervisor"),
        "status": hc("status"),
        "banda": hc("banda"),
        "powerapps_id": hc("__powerappsid__"),
    }

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        di = as_date(_cell(row, c_di))
        mat = as_str(_cell(row, c_mat))
        if di is None or not mat:
            continue
        item: dict[str, Any] = {
            "data_inicial": di,
            "data_final": as_date(_cell(row, c_df)) if c_df is not None else None,
            "matricula_agente": mat,
        }
        for key, idx in cols.items():
            if idx is None:
                item[key] = "" if key not in {"id_operations", "data_admissao"} else None
                continue
            if key == "id_operations":
                item[key] = as_int(_cell(row, idx))
            elif key == "data_admissao":
                item[key] = as_date(_cell(row, idx))
            else:
                item[key] = as_str(_cell(row, idx))
        rows.append(item)
    return rows


def load_workbook_data(path: Path | str | None = None) -> WorkbookData:
    path = Path(path) if path else DEFAULT_XLSX
    if not path.is_file():
        raise FileNotFoundError(
            f"Workbook não encontrado: {path}. "
            "Coloque Identificação dos Processos.xlsx em "
            "backend/data/dimensoes_processos/identificacao_processos.xlsx"
        )

    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = _sheet_names(wb)
        return WorkbookData(
            sheet_names=list(wb.sheetnames),
            produtos=_read_produtos(wb, names),
            grupos_servico=_read_grupos_servico(wb, names),
            servicos=_read_servicos(wb, names),
            clientes=_read_clientes(wb, names),
            niveis_hierarquicos=_read_niveis_hierarquicos(wb, names),
            workflows=_read_workflows(wb, names),
            etapas=_read_etapas(wb, names),
            metas_etapa=_read_metas_etapa(wb, names),
            projecao_sla=_read_projecao_sla(wb, names),
            projecao_equipes=_read_projecao_equipes(wb, names),
        )
    finally:
        wb.close()
