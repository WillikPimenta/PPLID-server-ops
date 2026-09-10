# -*- coding: utf-8 -*-
"""Testes API import Identificação dos Processos."""

from __future__ import annotations

import tempfile
from datetime import date
from io import BytesIO
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TransactionTestCase, override_settings
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.access.constants import ROLE_OP_AGENTE, ROLE_PLAN_ANALISTA, role_group_name

User = get_user_model()
BASE = "/api/v1/dimensoes-processos/identificacao-processos"


def _build_workbook_bytes() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Produto")
    ws.append(["id_produto", "tipo_produto"])
    ws.append([1, "Documentoscopia"])

    ws = wb.create_sheet("Clientes")
    ws.append(["id_cliente", "nome_cliente", "operations", "id_classificacao"])
    ws.append([100, "Cliente API", 1, 0])

    ws = wb.create_sheet("Nivel_Hierarquico")
    ws.append(["id_nh", "nivel hierarquico", "ind_considerar"])
    ws.append([10, "NH API", 1])

    ws = wb.create_sheet("Workflow")
    ws.append(["id_workflow", "workflow", "ind_considerar", "id_produto", "tipo_atendimento"])
    ws.append([700, "WF API", 1, 1, "Manual"])

    ws = wb.create_sheet("Projeção_SLA")
    ws.append(
        [
            "id_cliente",
            "data_início",
            "data_fim",
            "id_workflow",
            "id_nivel_hierarquico",
            "dias_semana",
            "hora_início",
            "hora_fim",
            "duração_atendimento",
            "sla_segundos",
            "flag_ajuste_sla",
            "sla_ajuste",
            "volume",
        ]
    )
    ws.append([100, date(2026, 1, 1), None, 700, 10, "{0..4}", None, None, None, 3600, None, None, 42.0])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@override_settings(ACCESS_ENFORCEMENT=True, MEDIA_ROOT=tempfile.mkdtemp())
class IdentificacaoProcessosApiTests(TransactionTestCase):
    def setUp(self):
        self.analyst = User.objects.create_user(
            username="megazord_ident",
            password="x",
            email="megazord_ident@example.com",
        )
        Group.objects.get_or_create(name=role_group_name(ROLE_PLAN_ANALISTA))
        self.analyst.groups.add(Group.objects.get(name=role_group_name(ROLE_PLAN_ANALISTA)))

        self.denied = User.objects.create_user(
            username="op_agente_ident",
            password="x",
            email="op_agente_ident@example.com",
        )
        Group.objects.get_or_create(name=role_group_name(ROLE_OP_AGENTE))
        self.denied.groups.add(Group.objects.get(name=role_group_name(ROLE_OP_AGENTE)))

        self.workbook_bytes = _build_workbook_bytes()
        self.client = APIClient()

    def _upload(self, user=None):
        client = APIClient()
        client.force_authenticate(user=user or self.analyst)
        upload = SimpleUploadedFile(
            "identificacao.xlsx",
            self.workbook_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        return client.post(f"{BASE}/upload/", {"file": upload}, format="multipart")

    def test_upload_rejects_non_xlsx(self):
        self.client.force_authenticate(user=self.analyst)
        bad = SimpleUploadedFile("bad.txt", b"hello", content_type="text/plain")
        res = self.client.post(f"{BASE}/upload/", {"file": bad}, format="multipart")
        self.assertEqual(res.status_code, 400)

    def test_upload_preflight_import_flow(self):
        upload_res = self._upload()
        self.assertEqual(upload_res.status_code, 201)
        token = upload_res.data["upload"]["token"]

        self.client.force_authenticate(user=self.analyst)
        preflight = self.client.post(f"{BASE}/preflight/", {"token": token}, format="json")
        self.assertEqual(preflight.status_code, 200)
        self.assertTrue(preflight.data["ok"])
        self.assertTrue(preflight.data["preflight"]["ok"])

        import_res = self.client.post(
            f"{BASE}/import/",
            {"token": token, "confirm": True, "mode": "sync"},
            format="json",
        )
        self.assertIn(import_res.status_code, (200, 202))
        run_id = import_res.data["run"]["id"]

        import time

        deadline = time.time() + 60
        final_status = import_res.data["run"]["status"]
        while time.time() < deadline and final_status == "running":
            time.sleep(0.5)
            poll = self.client.get(f"{BASE}/runs/{run_id}/")
            self.assertEqual(poll.status_code, 200)
            final_status = poll.data["run"]["status"]
        self.assertEqual(final_status, "ok")

        from apps.dimensoes_processos.models import DimCliente, ProjecaoSla

        self.assertTrue(DimCliente.objects.filter(id_cliente=100).exists())
        sla = ProjecaoSla.objects.get(cliente_id=100, workflow_id=700)
        self.assertEqual(sla.volume, 42.0)

    def test_denied_without_megazord_permission(self):
        res = self._upload(user=self.denied)
        self.assertEqual(res.status_code, 403)
