"""Testes do importador Identificação dos Processos."""

from __future__ import annotations

import tempfile
from datetime import date
from pathlib import Path

from django.test import TestCase
from openpyxl import Workbook

from apps.dimensoes_processos.models import (
    DimCliente,
    DimNivelHierarquico,
    DimProduto,
    DimWorkflow,
    ProjecaoSla,
)
from apps.dimensoes_processos.services.importer import import_identificacao_processos


def _write_min_workbook(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Produto")
    ws.append(["id_produto", "tipo_produto"])
    ws.append([1, "Documentoscopia"])

    ws = wb.create_sheet("Clientes")
    ws.append(["id_cliente", "nome_cliente", "operations", "id_classificacao"])
    ws.append([100, "Cliente A", 1, 0])
    ws.append([200, "Cliente B", 1, 0])

    ws = wb.create_sheet("Nivel_Hierarquico")
    ws.append(["id_nh", "nivel hierarquico", "ind_considerar"])
    ws.append([10, "NH Teste", 1])

    ws = wb.create_sheet("Workflow")
    ws.append(["id_workflow", "workflow", "ind_considerar", "id_produto", "tipo_atendimento"])
    ws.append([700, "WF Teste", 1, 1, "Manual"])
    ws.append([800, "WF Novo", 1, 1, "Manual"])

    ws = wb.create_sheet("Projeção_SLA")
    ws.append(
        [
            "id_cliente",
            "data_início",
            "data_fim",
            "id_workflow",
            "id_nivel_hierarquico",
            "dias_semana",
            "hora_início",
            "hora_fim",
            "duração_atendimento",
            "sla_segundos",
            "flag_ajuste_sla",
            "sla_ajuste",
            "volume",
        ]
    )
    ws.append([100, date(2026, 1, 1), None, 700, 10, "{0..4}", None, None, None, 3600, None, None, 100.0])
    ws.append([200, date(2026, 1, 1), None, 800, 10, "{0..4}", None, None, None, 3600, None, None, 50.0])

    wb.save(path)


class IdentificacaoProcessosImporterTests(TestCase):
    def setUp(self):
        DimProduto.objects.create(id_produto=1, tipo_produto="Documentoscopia")
        DimCliente.objects.create(id_cliente=100, nome="Cliente A", operations=True)
        DimNivelHierarquico.objects.create(id_nh=10, nome="NH Teste")
        DimWorkflow.objects.create(id_workflow=700, nome="WF Teste", produto_id=1)
        ProjecaoSla.objects.create(
            cliente_id=100,
            workflow_id=700,
            nivel_hierarquico_id=10,
            data_inicio=date(2026, 1, 1),
            dias_semana="{0..4}",
            volume=10.0,
        )
        fd, name = tempfile.mkstemp(suffix=".xlsx")
        self.workbook_path = Path(name)
        import os

        os.close(fd)
        _write_min_workbook(self.workbook_path)

    def tearDown(self):
        if getattr(self, "workbook_path", None) and self.workbook_path.exists():
            try:
                self.workbook_path.unlink()
            except PermissionError:
                pass

    def test_sync_updates_volume_and_creates_new_entities(self):
        summary = import_identificacao_processos(self.workbook_path, mode="sync")

        self.assertEqual(DimCliente.objects.filter(id_cliente=200).count(), 1)
        self.assertEqual(DimWorkflow.objects.filter(id_workflow=800).count(), 1)
        sla = ProjecaoSla.objects.get(cliente_id=100, workflow_id=700, data_inicio=date(2026, 1, 1))
        self.assertEqual(sla.volume, 100.0)
        self.assertEqual(ProjecaoSla.objects.filter(cliente_id=200, workflow_id=800).count(), 1)
        self.assertEqual(summary["entities"]["projecao_sla"]["updated"], 1)
        self.assertEqual(summary["entities"]["projecao_sla"]["created"], 1)

    def test_sync_does_not_remove_existing_absent_from_file(self):
        DimCliente.objects.create(id_cliente=999, nome="Só no portal")
        import_identificacao_processos(self.workbook_path, mode="sync")
        self.assertTrue(DimCliente.objects.filter(id_cliente=999).exists())

    def test_replace_clears_and_reloads(self):
        DimCliente.objects.create(id_cliente=999, nome="Só no portal")
        import_identificacao_processos(self.workbook_path, mode="replace")
        self.assertFalse(DimCliente.objects.filter(id_cliente=999).exists())
        self.assertEqual(DimCliente.objects.count(), 2)

    def test_cadastros_scope_does_not_touch_sla(self):
        summary = import_identificacao_processos(self.workbook_path, mode="sync", import_scope="cadastros")
        sla = ProjecaoSla.objects.get(cliente_id=100, workflow_id=700, data_inicio=date(2026, 1, 1))
        self.assertEqual(sla.volume, 10.0)
        self.assertFalse(ProjecaoSla.objects.filter(cliente_id=200).exists())
        self.assertNotIn("projecao_sla", summary.get("entities", {}))

    def test_rotates_sla_when_new_data_inicio(self):
        import os

        wb = Workbook()
        wb.remove(wb.active)
        for sheet, headers, row in (
            ("Produto", ["id_produto", "tipo_produto"], [1, "Documentoscopia"]),
            ("Clientes", ["id_cliente", "nome_cliente", "operations", "id_classificacao"], [100, "Cliente A", 1, 0]),
            ("Nivel_Hierarquico", ["id_nh", "nivel hierarquico", "ind_considerar"], [10, "NH", 1]),
            ("Workflow", ["id_workflow", "workflow", "ind_considerar", "id_produto", "tipo_atendimento"], [700, "WF", 1, 1, "Manual"]),
        ):
            ws = wb.create_sheet(sheet)
            ws.append(headers)
            ws.append(row)
        ws = wb.create_sheet("Projeção_SLA")
        ws.append(
            [
                "id_cliente", "data_início", "data_fim", "id_workflow", "id_nivel_hierarquico",
                "dias_semana", "hora_início", "hora_fim", "duração_atendimento", "sla_segundos",
                "flag_ajuste_sla", "sla_ajuste", "volume",
            ]
        )
        ws.append([100, date(2026, 6, 1), None, 700, 10, "{0..4}", None, None, None, 3600, None, None, 200.0])
        fd, name = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        rotate_path = Path(name)
        wb.save(rotate_path)
        summary = import_identificacao_processos(rotate_path, mode="sync", import_scope="projecao_sla")
        self.assertEqual(summary.get("sla_rotated"), 1)
        self.assertEqual(ProjecaoSla.objects.filter(data_fim__isnull=True).count(), 1)
        vigente = ProjecaoSla.objects.get(data_fim__isnull=True)
        self.assertEqual(vigente.data_inicio, date(2026, 6, 1))
        self.assertEqual(vigente.volume, 200.0)
        try:
            rotate_path.unlink()
        except PermissionError:
            pass

