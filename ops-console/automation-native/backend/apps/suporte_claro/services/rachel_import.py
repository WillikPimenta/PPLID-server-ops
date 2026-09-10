# -*- coding: utf-8 -*-
"""Importação idempotente dos INC da planilha DEMANDAS FY27 da Rachel."""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime, time
from io import BytesIO

import openpyxl
from django.db import transaction
from django.utils import timezone

from apps.suporte_claro.models import (
    SuporteClaroHistorico,
    SuporteClaroImportRef,
    SuporteClaroRegistro,
)
from apps.suporte_claro.services.chamados_externos import save_chamados_externos
from apps.suporte_claro.services.protocolos import save_protocolos

RACHEL_MARKERS = ("RACHEL", "RAQUEL")
RACHEL_USERNAME = "c93189a"
RACHEL_FULL_NAME = "Rachel Ramos De Lima Pereira"
IMPORT_PREFIX = "FY27-DEMANDAS-"

ESTADO_TO_STATUS = {
    "RESOLVIDA": SuporteClaroRegistro.STATUS_CONCLUIDO,
    "EM ANDAMENTO": SuporteClaroRegistro.STATUS_EM_ATENDIMENTO,
    "PAUSADA": SuporteClaroRegistro.STATUS_EM_ATENDIMENTO,
    "CANCELADA": SuporteClaroRegistro.STATUS_CONCLUIDO,
}


@dataclass
class RachelImportRow:
    row: int
    sheet: str
    numero: str
    titulo: str
    cliente: str
    desc: str
    estado: str
    received_at: datetime | None
    resolved_at: datetime | None
    tipo_incidente: str
    tipo_label: str
    action: str = "create"
    message: str = "Pronto para importar."
    registro_id: int | None = None


@dataclass
class RachelImportPreview:
    total: int
    valid: int
    skipped: int
    errors: int
    rows: list[RachelImportRow]
    found_sheets: list[str]
    missing_sheets: list[str]
    ignored_sheets: list[str]


@dataclass
class RachelImportApplyResult(RachelImportPreview):
    created: int = 0


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _parse_date(raw) -> datetime | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw
    text = _norm(raw)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _as_aware(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if timezone.is_naive(value):
        return timezone.make_aware(value, timezone.get_current_timezone())
    return value


def infer_incident_type(titulo: str, descricao: str) -> str:
    blob = unicodedata.normalize("NFKD", f"{titulo} {descricao}".casefold())
    blob = "".join(char for char in blob if not unicodedata.combining(char))
    if "lentid" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_LENTIDAO
    if "queda" in blob or "indispon" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_QUEDA
    if "trav" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_TRAVAMENTO
    if "erro" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_ERRO
    return SuporteClaroRegistro.TIPO_INCIDENTE_OUTRO


def _tipo_label(tipo: str) -> str:
    return dict(SuporteClaroRegistro.TIPO_INCIDENTE_CHOICES).get(tipo, "Outro")


HEADER_ALIASES = {
    "quem": {"quem tratou", "quem", "tratado por", "responsavel"},
    "data": {"data da analise", "data analise", "data de recebimento", "data"},
    "num": {
        "n da demanda",
        "no da demanda",
        "numero da demanda",
        "numero",
        "demanda",
        "servicenow",
    },
    "titulo": {"titulo", "assunto"},
    "cliente": {"cliente"},
    "desc": {"descricao", "detalhes", "irregularidade"},
    "estado": {"estado", "situacao"},
    "data_res": {"data resolvida", "data de resolucao", "data resolucao", "resolvido em"},
}


def _header_key(value) -> str:
    text = unicodedata.normalize("NFKD", _norm(value).casefold())
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", " ", text).strip()
    return text


def _detect_sheet_layout(worksheet) -> tuple[int, dict[str, int]] | None:
    """Encontra os cabeçalhos sem depender do nome/ordem da aba."""
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=10, values_only=True), start=1
    ):
        normalized = [_header_key(value) for value in row]
        columns: dict[str, int] = {}
        for field, aliases in HEADER_ALIASES.items():
            for index, header in enumerate(normalized):
                if header in aliases:
                    columns[field] = index
                    break
        if {"quem", "num", "titulo"}.issubset(columns):
            return row_number, columns
    return None


def _cell(row, index: int):
    return None if index < 0 or index >= len(row) else row[index]


def parse_rachel_incidents(
    file_bytes: bytes,
) -> tuple[list[RachelImportRow], list[str], list[str]]:
    try:
        workbook = openpyxl.load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Arquivo inválido. Selecione a planilha DEMANDAS FY27 em formato .xlsx.") from exc

    layouts = {
        name: layout
        for name in workbook.sheetnames
        if (layout := _detect_sheet_layout(workbook[name])) is not None
    }
    found = list(layouts)
    ignored = [name for name in workbook.sheetnames if name not in layouts]
    if not found:
        workbook.close()
        raise ValueError(
            "Nenhuma aba compatível foi encontrada. Verifique os cabeçalhos "
            "Quem tratou, Nº da demanda e Título."
        )

    output: list[RachelImportRow] = []
    seen: set[str] = set()
    try:
        for sheet_name, (header_row, columns) in layouts.items():
            for row_number, row in enumerate(
                workbook[sheet_name].iter_rows(min_row=header_row + 1, values_only=True),
                start=header_row + 1,
            ):
                numero = _norm(_cell(row, columns["num"])).upper()
                quem = _norm(_cell(row, columns["quem"])).upper()
                if not numero or "INC" not in numero or not any(x in quem for x in RACHEL_MARKERS):
                    continue
                if numero in seen:
                    continue
                seen.add(numero)
                titulo = _norm(_cell(row, columns["titulo"])) or numero
                descricao = _norm(_cell(row, columns.get("desc", -1)))
                tipo = infer_incident_type(titulo, descricao)
                output.append(RachelImportRow(
                    row=row_number,
                    sheet=sheet_name,
                    numero=numero,
                    titulo=titulo,
                    cliente=_norm(_cell(row, columns.get("cliente", -1))),
                    desc=descricao,
                    estado=_norm(_cell(row, columns.get("estado", -1))),
                    received_at=_parse_date(_cell(row, columns.get("data", -1))),
                    resolved_at=_parse_date(_cell(row, columns.get("data_res", -1))),
                    tipo_incidente=tipo,
                    tipo_label=_tipo_label(tipo),
                ))
    finally:
        workbook.close()
    return output, found, ignored


def _mark_existing(rows: list[RachelImportRow]) -> None:
    for item in rows:
        import_ref = SuporteClaroImportRef.objects.filter(
            linha_id=f"{IMPORT_PREFIX}{item.numero}"
        ).first()
        if import_ref:
            item.action = "skip"
            item.registro_id = import_ref.registro_id
            item.message = f"Já importado na demanda #{import_ref.registro_id}."
            continue
        existing = SuporteClaroRegistro.objects.filter(
            chamados_externos__sistema=SuporteClaroRegistro.CHAMADO_SERVICE,
            chamados_externos__codigo__iexact=item.numero,
        ).distinct().first()
        if existing:
            item.action = "skip"
            item.registro_id = existing.id
            item.message = f"ServiceNow já vinculado à demanda #{existing.id}."


def preview_rachel_import(file_bytes: bytes) -> RachelImportPreview:
    rows, found, ignored = parse_rachel_incidents(file_bytes)
    _mark_existing(rows)
    return RachelImportPreview(
        total=len(rows),
        valid=sum(row.action == "create" for row in rows),
        skipped=sum(row.action == "skip" for row in rows),
        errors=sum(row.action == "error" for row in rows),
        rows=rows,
        found_sheets=found,
        missing_sheets=[],
        ignored_sheets=ignored,
    )


def _map_status(estado: str) -> str:
    return ESTADO_TO_STATUS.get(_norm(estado).upper(), SuporteClaroRegistro.STATUS_EM_ATENDIMENTO)


@transaction.atomic
def apply_rachel_import(file_bytes: bytes, *, imported_by, record_owner) -> RachelImportApplyResult:
    preview = preview_rachel_import(file_bytes)
    created = 0
    for item in preview.rows:
        if item.action != "create":
            continue
        received = item.received_at or datetime.combine(timezone.localdate(), time(9, 0))
        status_value = _map_status(item.estado)
        irregularidade = item.desc or item.titulo
        if item.cliente:
            irregularidade = f"[{item.cliente}] {irregularidade}"
        avaliacao = ""
        retorno_at = None
        if status_value == SuporteClaroRegistro.STATUS_CONCLUIDO:
            avaliacao = (
                f"Importado da planilha DEMANDAS FY27 ({item.sheet}). "
                f"Estado de origem: {item.estado or 'RESOLVIDA'}."
            )
            retorno_at = _as_aware(item.resolved_at or item.received_at or received)
        registro = SuporteClaroRegistro.objects.create(
            titulo=item.titulo[:255],
            protocolo="",
            irregularidade=irregularidade,
            avaliacao=avaliacao,
            received_at=_as_aware(received),
            retorno_at=retorno_at,
            sent_by=item.cliente[:255],
            origem=SuporteClaroRegistro.ORIGEM_TEAMS,
            categoria=SuporteClaroRegistro.CATEGORIA_INCIDENTE,
            tipo_incidente=item.tipo_incidente,
            status=status_value,
            created_by=record_owner,
        )
        auto_code = f"INC-{registro.id}"
        registro.protocolo = auto_code
        registro.save(update_fields=["protocolo", "updated_at"])
        save_protocolos(registro, [(auto_code, "")], user=imported_by)
        save_chamados_externos(
            registro,
            [(SuporteClaroRegistro.CHAMADO_SERVICE, item.numero, "", RACHEL_FULL_NAME)],
            user=imported_by,
        )
        SuporteClaroImportRef.objects.create(
            linha_id=f"{IMPORT_PREFIX}{item.numero}", registro=registro, imported_by=imported_by
        )
        SuporteClaroHistorico.objects.create(
            registro=registro,
            user=imported_by,
            action=SuporteClaroHistorico.ACTION_IMPORT,
            field_name="import",
            old_value="",
            new_value=f"DEMANDAS FY27 {item.numero} ({item.sheet}, {item.estado or '-'})",
        )
        item.registro_id = registro.id
        item.message = f"Importado na demanda #{registro.id}."
        created += 1

    return RachelImportApplyResult(
        total=preview.total,
        valid=preview.valid,
        skipped=preview.skipped,
        errors=preview.errors,
        rows=preview.rows,
        found_sheets=preview.found_sheets,
        missing_sheets=preview.missing_sheets,
        ignored_sheets=preview.ignored_sheets,
        created=created,
    )


def serialize_rachel_preview(result: RachelImportPreview) -> dict:
    return {
        "total": result.total,
        "valid": result.valid,
        "skipped": result.skipped,
        "errors": result.errors,
        "found_sheets": result.found_sheets,
        "missing_sheets": result.missing_sheets,
        "ignored_sheets": result.ignored_sheets,
        "rows": [
            {
                "row": row.row,
                "sheet": row.sheet,
                "numero": row.numero,
                "titulo": row.titulo,
                "tipo_incidente": row.tipo_incidente,
                "tipo_label": row.tipo_label,
                "estado": row.estado,
                "action": row.action,
                "message": row.message,
                "registro_id": row.registro_id,
            }
            for row in result.rows
        ],
    }


def serialize_rachel_apply(result: RachelImportApplyResult) -> dict:
    return {**serialize_rachel_preview(result), "created": result.created}
