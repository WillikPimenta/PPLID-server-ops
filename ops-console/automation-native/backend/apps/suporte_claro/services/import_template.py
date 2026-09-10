# -*- coding: utf-8 -*-
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

IMPORT_SHEET_NAME = "Demandas"
INSTRUCTIONS_SHEET_NAME = "Instruções"

IMPORT_COLUMNS = [
    "Linha_ID",
    "Titulo",
    "Protocolo",
    "Protocolos",
    "Canal",
    "Recebido em",
    "Quem enviou",
    "Irregularidade (cliente)",
    "Status",
    "Retorno do suporte",
    "Comentarios protocolo",
    "Observações internas",
    "Caminho anexos",
]

CANAL_OPTIONS = ["Teams", "E-mail", "Ligação"]
STATUS_OPTIONS = ["Não iniciado", "Em andamento", "Concluído"]

EXAMPLE_ROW = [
    "OFF-001",
    "Divergencia no extrato",
    "PROT-EXEMPLO-001",
    "PROT-EXEMPLO-001, PROT-EXEMPLO-002",
    "Teams",
    "08/07/2026 09:30",
    "Maria Silva",
    "Cliente reportou divergência no extrato.",
    "Concluído",
    "Cliente orientado e caso encerrado.",
    "Protocolo principal | Segundo protocolo vinculado",
    "Atendimento offline — servidor indisponível.",
    "C:\\SuporteClaro\\OFF-001\\",
]

INSTRUCTIONS_LINES = [
    "Planilha paliativa — Suporte Claro (cadastro offline)",
    "",
    "1. Preencha uma linha por demanda atendida enquanto o portal estiver indisponível.",
    "2. Linha_ID deve ser única (ex.: OFF-001). Permite reimportar sem duplicar registros.",
    "3. Titulo: nome da demanda no kanban. Protocolo (legado) ou Protocolos (lista separada por virgula).",
    "4. Comentarios protocolo: um comentario por protocolo, separados por | (pipe).",
    "5. Canal: Teams, E-mail ou Ligação.",
    "6. Recebido em: DD/MM/AAAA HH:MM (ex.: 08/07/2026 09:30).",
    "7. Status Concluído exige Retorno do suporte preenchido.",
    "8. Anexos: guarde prints na pasta indicada em Caminho anexos; após import, abra a demanda no kanban e anexe manualmente.",
    "9. Quando o portal voltar: Kanban → Importar planilha → revisar preview → confirmar.",
    "10. Formalização no Jira continua pelo modal Pendentes de formalização após o import.",
]


def build_import_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = IMPORT_SHEET_NAME

    header_font = Font(bold=True)
    for col_idx, header in enumerate(IMPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font

    for col_idx, value in enumerate(EXAMPLE_ROW, start=1):
        ws.cell(row=2, column=col_idx, value=value)

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 40
    ws.column_dimensions["I"].width = 16
    ws.column_dimensions["J"].width = 40
    ws.column_dimensions["K"].width = 36
    ws.column_dimensions["L"].width = 32
    ws.column_dimensions["M"].width = 28

    canal_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(CANAL_OPTIONS)}"',
        allow_blank=False,
    )
    canal_validation.add(f"E2:E500")
    ws.add_data_validation(canal_validation)

    status_validation = DataValidation(
        type="list",
        formula1=f'"{",".join(STATUS_OPTIONS)}"',
        allow_blank=False,
    )
    status_validation.add(f"I2:I500")
    ws.add_data_validation(status_validation)

    instr = wb.create_sheet(INSTRUCTIONS_SHEET_NAME)
    for row_idx, line in enumerate(INSTRUCTIONS_LINES, start=1):
        cell = instr.cell(row=row_idx, column=1, value=line)
        if row_idx == 1:
            cell.font = header_font
    instr.column_dimensions["A"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
