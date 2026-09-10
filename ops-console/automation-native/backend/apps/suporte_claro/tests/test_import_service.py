# -*- coding: utf-8 -*-
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.suporte_claro.models import SuporteClaroImportRef, SuporteClaroRegistro
from apps.suporte_claro.services.import_service import apply_import, preview_import
from apps.suporte_claro.services.import_template import IMPORT_COLUMNS, IMPORT_SHEET_NAME

User = get_user_model()


def _build_import_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = IMPORT_SHEET_NAME
    for col_idx, header in enumerate(IMPORT_COLUMNS, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, values in enumerate(rows, start=2):
        for col_idx, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class SuporteClaroImportServiceTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="agent001",
            password="pass12345",
            email="agent001@test.local",
        )

    def _valid_row(self, linha_id: str = "OFF-001", protocolo: str = "PROT-001"):
        return [
            linha_id,
            f"Demanda {protocolo}",
            protocolo,
            protocolo,
            "Teams",
            "08/07/2026 09:30",
            "Maria",
            "Irregularidade teste",
            "Concluído",
            "Retorno teste",
            "",
            "",
            "",
        ]

    def test_preview_valid_row(self):
        content = _build_import_xlsx([self._valid_row()])
        preview = preview_import(content)
        self.assertEqual(preview.total, 1)
        self.assertEqual(preview.valid, 1)
        self.assertEqual(preview.errors, 0)
        self.assertEqual(preview.rows[0].status, "ok")

    def test_preview_error_when_concluido_sem_retorno(self):
        row = self._valid_row()
        row[9] = ""
        content = _build_import_xlsx([row])
        preview = preview_import(content)
        self.assertEqual(preview.errors, 1)
        self.assertEqual(preview.rows[0].status, "error")

    def test_apply_creates_registro_and_skips_reimport(self):
        content = _build_import_xlsx([self._valid_row()])
        result = apply_import(content, self.user)
        self.assertEqual(result.created, 1)
        self.assertEqual(SuporteClaroRegistro.objects.filter(protocolo="PROT-001").count(), 1)
        self.assertEqual(SuporteClaroImportRef.objects.filter(linha_id="OFF-001").count(), 1)

        preview2 = preview_import(content)
        self.assertEqual(preview2.skipped, 1)
        self.assertEqual(preview2.valid, 0)

        result2 = apply_import(content, self.user)
        self.assertEqual(result2.created, 0)
        self.assertEqual(result2.skipped, 1)


class SuporteClaroImportApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.user = User.objects.create_user(
            username="agent001",
            password="pass12345",
            email="agent001@test.local",
        )
        self.client.force_authenticate(user=self.user)

    def test_import_template_download(self):
        response = self.client.get("/api/v1/suporte-claro/registros/import-template.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            response["Content-Type"],
        )
        self.assertTrue(len(response.content) > 1000)

    def test_import_preview_and_apply(self):
        content = _build_import_xlsx(
            [
                [
                    "OFF-API-1",
                    "Demanda PROT-API-1",
                    "PROT-API-1",
                    "PROT-API-1",
                    "E-mail",
                    "08/07/2026 10:00",
                    "João",
                    "Caso offline",
                    "Em andamento",
                    "",
                    "",
                    "",
                    "",
                ]
            ]
        )
        upload = SimpleUploadedFile(
            "import.xlsx",
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        preview = self.client.post(
            "/api/v1/suporte-claro/registros/import/preview/",
            {"file": upload},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["valid"], 1)

        upload2 = SimpleUploadedFile(
            "import.xlsx",
            content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        applied = self.client.post(
            "/api/v1/suporte-claro/registros/import/apply/",
            {"file": upload2},
            format="multipart",
        )
        self.assertEqual(applied.status_code, 201)
        self.assertEqual(applied.json()["created"], 1)
        self.assertTrue(
            SuporteClaroRegistro.objects.filter(protocolo="PROT-API-1").exists()
        )
