# -*- coding: utf-8 -*-
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.suporte_claro.models import SuporteClaroImportRef, SuporteClaroRegistro
from apps.suporte_claro.services.rachel_import import (
    apply_rachel_import,
    infer_incident_type,
    preview_rachel_import,
)

User = get_user_model()


def _build_fy27_xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "MAIO"
    sheet.append(["Quem", "Data", "Número", "Título", "Cliente", "Descrição", "Obs", "Estado", "Resolução"])
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_dynamic_fy27_xlsx() -> bytes:
    workbook = Workbook()
    dados = workbook.active
    dados.title = "DADOS"
    dados.append(["ESTADO", "LISTA - TIPO"])
    dados.append(["RESOLVIDA", "INCIDENTE"])
    for sheet_name, numero in (("SETEMBRO", "INC901"), ("OUTUBRO FY27", "INC902")):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["QUEM TRATOU?", "DATA DA ANÁLISE", "Nº DA DEMANDA", "TÍTULO", "CLIENTE", "DESCRIÇÃO", "ESTADO", "DATA RESOLVIDA"])
        sheet.append(["RACHEL PEREIRA", "01/09/2026", numero, "Erro no Confer", "CONFER", "Caso com erro", "RESOLVIDA", "02/09/2026"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _row(numero="INC001", titulo="Lentidão no portal", quem="Rachel Ramos"):
    return [quem, "05/05/2026", numero, titulo, "Cliente A", "Falha observada", "", "RESOLVIDA", "06/05/2026"]


class RachelImportServiceTests(TestCase):
    def setUp(self):
        self.importer = User.objects.create_user(username="importer", email="importer@test.local")
        self.rachel = User.objects.create_user(username="c93189a", email="rachel@test.local")

    def test_preview_filters_rachel_inc_and_infers_types(self):
        content = _build_fy27_xlsx([
            _row("INC001", "Lentidão no portal"),
            _row("INC002", "Sistema indisponível"),
            _row("INC003", "Tela travando"),
            _row("INC004", "Erro na consulta"),
            _row("INC005", "Comportamento inesperado"),
            _row("INC999", "Não deve entrar", quem="Outro Analista"),
            _row("REQ001", "Não é incidente"),
        ])
        preview = preview_rachel_import(content)
        self.assertEqual(preview.total, 5)
        self.assertEqual(preview.valid, 5)
        self.assertEqual(
            [row.tipo_incidente for row in preview.rows],
            ["lentidao", "queda", "travamento", "erro", "outro"],
        )

    def test_reads_any_compatible_sheet_and_ignores_auxiliary_tabs(self):
        preview = preview_rachel_import(_build_dynamic_fy27_xlsx())
        self.assertEqual(preview.total, 2)
        self.assertEqual(preview.found_sheets, ["SETEMBRO", "OUTUBRO FY27"])
        self.assertEqual(preview.ignored_sheets, ["DADOS"])

    def test_infer_unknown_as_outro_never_sem_erro(self):
        self.assertEqual(infer_incident_type("Falha genérica", "Sem pista"), "outro")

    def test_apply_is_idempotent_and_preserves_owner_and_audit_user(self):
        content = _build_fy27_xlsx([_row()])
        result = apply_rachel_import(
            content, imported_by=self.importer, record_owner=self.rachel
        )
        self.assertEqual(result.created, 1)
        registro = SuporteClaroRegistro.objects.get()
        self.assertEqual(registro.created_by, self.rachel)
        self.assertEqual(registro.tipo_incidente, "lentidao")
        self.assertTrue(
            registro.chamados_externos.filter(sistema="service", codigo="INC001").exists()
        )
        ref = SuporteClaroImportRef.objects.get(linha_id="FY27-DEMANDAS-INC001")
        self.assertEqual(ref.imported_by, self.importer)

        second = apply_rachel_import(
            content, imported_by=self.importer, record_owner=self.rachel
        )
        self.assertEqual(second.created, 0)
        self.assertEqual(second.skipped, 1)
        self.assertEqual(SuporteClaroRegistro.objects.count(), 1)


class RachelImportApiTests(TestCase):
    def setUp(self):
        self.client = APIClient()
        self.importer = User.objects.create_user(username="importer", email="importer@test.local")
        self.rachel = User.objects.create_user(username="c93189a", email="rachel@test.local")
        self.client.force_authenticate(self.importer)
        self.content = _build_fy27_xlsx([_row()])

    def _upload(self):
        return SimpleUploadedFile(
            "DEMANDAS FY27.xlsx",
            self.content,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_preview_and_apply_endpoints(self):
        preview = self.client.post(
            "/api/v1/suporte-claro/registros/import-rachel/preview/",
            {"file": self._upload()},
            format="multipart",
        )
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.json()["valid"], 1)
        self.assertEqual(preview.json()["rows"][0]["tipo_label"], "Lentidão")

        applied = self.client.post(
            "/api/v1/suporte-claro/registros/import-rachel/apply/",
            {"file": self._upload()},
            format="multipart",
        )
        self.assertEqual(applied.status_code, 201)
        self.assertEqual(applied.json()["created"], 1)

    def test_apply_requires_rachel_user(self):
        self.rachel.delete()
        response = self.client.post(
            "/api/v1/suporte-claro/registros/import-rachel/apply/",
            {"file": self._upload()},
            format="multipart",
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("Rachel", response.json()["detail"])
