# -*- coding: utf-8 -*-
"""Importa INC (ServiceNow) tratados por Rachel da planilha DEMANDAS FY27 → Controle."""
from __future__ import annotations

import re
from datetime import datetime, time
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone

from apps.suporte_claro.models import (
    SuporteClaroHistorico,
    SuporteClaroImportRef,
    SuporteClaroRegistro,
)
from apps.suporte_claro.services.chamados_externos import save_chamados_externos
from apps.suporte_claro.services.protocolos import save_protocolos

MONTH_SHEETS = ("ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO")
RACHEL_MARKERS = ("RACHEL", "RAQUEL")
IMPORT_PREFIX = "FY27-DEMANDAS-"

ESTADO_TO_STATUS = {
    "RESOLVIDA": SuporteClaroRegistro.STATUS_CONCLUIDO,
    "EM ANDAMENTO": SuporteClaroRegistro.STATUS_EM_ATENDIMENTO,
    "PAUSADA": SuporteClaroRegistro.STATUS_EM_ATENDIMENTO,
    "CANCELADA": SuporteClaroRegistro.STATUS_CONCLUIDO,
}


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def _is_rachel(quem: str) -> bool:
    upper = quem.upper()
    return any(marker in upper for marker in RACHEL_MARKERS)


def _is_inc(numero: str) -> bool:
    return "INC" in numero.upper()


def _parse_date(raw) -> datetime | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, datetime):
        return raw
    text = _norm(raw)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _as_aware(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    if timezone.is_naive(dt):
        return timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def _map_status(estado: str) -> str:
    key = _norm(estado).upper()
    return ESTADO_TO_STATUS.get(key, SuporteClaroRegistro.STATUS_EM_ATENDIMENTO)


def _infer_tipo(titulo: str, desc: str) -> str:
    blob = f"{titulo} {desc}".lower()
    if "lentid" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_LENTIDAO
    if "queda" in blob or "indisponib" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_QUEDA
    if "trav" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_TRAVAMENTO
    if "erro" in blob:
        return SuporteClaroRegistro.TIPO_INCIDENTE_ERRO
    return SuporteClaroRegistro.TIPO_INCIDENTE_OUTRO


def _sheet_columns(sheet_name: str) -> dict[str, int]:
    if sheet_name == "ABRIL":
        return {
            "quem": 2,
            "data": 3,
            "num": 4,
            "titulo": 5,
            "cliente": 6,
            "desc": 7,
            "estado": 9,
            "data_res": -1,
        }
    return {
        "quem": 0,
        "data": 1,
        "num": 2,
        "titulo": 3,
        "cliente": 4,
        "desc": 5,
        "estado": 7,
        "data_res": 8,
    }


def _cell(row, idx: int):
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def iter_rachel_incs(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows_out: list[dict] = []
    seen: set[str] = set()
    try:
        for sheet_name in MONTH_SHEETS:
            if sheet_name not in wb.sheetnames:
                continue
            cols = _sheet_columns(sheet_name)
            for row in wb[sheet_name].iter_rows(min_row=2, values_only=True):
                if not row:
                    continue
                numero = _norm(_cell(row, cols["num"]))
                quem = _norm(_cell(row, cols["quem"]))
                if not numero or not _is_inc(numero) or not _is_rachel(quem):
                    continue
                key = numero.upper()
                if key in seen:
                    continue
                seen.add(key)
                rows_out.append(
                    {
                        "sheet": sheet_name,
                        "numero": key,
                        "quem": quem,
                        "titulo": _norm(_cell(row, cols["titulo"])) or key,
                        "cliente": _norm(_cell(row, cols["cliente"])),
                        "desc": _norm(_cell(row, cols["desc"])),
                        "estado": _norm(_cell(row, cols["estado"])),
                        "data": _parse_date(_cell(row, cols["data"])),
                        "data_res": _parse_date(_cell(row, cols["data_res"])),
                    }
                )
    finally:
        wb.close()
    return rows_out


class Command(BaseCommand):
    help = (
        "Importa nº DA DEMANDA com INC tratados por Rachel/Raquel "
        "da planilha DEMANDAS FY27 como incidentes no Controle."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "xlsx",
            nargs="?",
            default=str(Path.home() / "Downloads" / "DEMANDAS FY27.xlsx"),
            help="Caminho da planilha DEMANDAS FY27.xlsx",
        )
        parser.add_argument(
            "--username",
            default="c93189a",
            help="Usuário created_by / imported_by (default: c93189a Rachel Planejamento)",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Só lista o que seria importado, sem gravar.",
        )

    def handle(self, *args, **options):
        path = Path(options["xlsx"]).expanduser()
        if not path.is_file():
            raise CommandError(f"Arquivo não encontrado: {path}")

        User = get_user_model()
        user = User.objects.filter(username=options["username"]).first()
        if user is None:
            raise CommandError(f"Usuário não encontrado: {options['username']}")

        items = iter_rachel_incs(path)
        if not items:
            self.stdout.write(self.style.WARNING("Nenhum INC da Rachel encontrado."))
            return

        self.stdout.write(f"Encontrados {len(items)} INC (Rachel) em {path.name}")
        if options["dry_run"]:
            for item in items:
                self.stdout.write(
                    f"  {item['numero']} | {item['estado'] or '-'} | {item['titulo'][:60]}"
                )
            return

        created = skipped = 0
        with transaction.atomic():
            for item in items:
                linha_id = f"{IMPORT_PREFIX}{item['numero']}"
                existing = (
                    SuporteClaroImportRef.objects.filter(linha_id=linha_id)
                    .select_related("registro")
                    .first()
                )
                if existing:
                    skipped += 1
                    self.stdout.write(
                        f"skip {item['numero']} -> demanda #{existing.registro_id}"
                    )
                    continue

                # Também evita duplicar se já houver ServiceNow com o mesmo código.
                already = (
                    SuporteClaroRegistro.objects.filter(
                        chamados_externos__sistema=SuporteClaroRegistro.CHAMADO_SERVICE,
                        chamados_externos__codigo__iexact=item["numero"],
                    )
                    .distinct()
                    .first()
                )
                if already:
                    SuporteClaroImportRef.objects.get_or_create(
                        linha_id=linha_id,
                        defaults={"registro": already, "imported_by": user},
                    )
                    skipped += 1
                    self.stdout.write(
                        f"skip {item['numero']} (já existe #{already.id})"
                    )
                    continue

                received = item["data"] or datetime.combine(
                    timezone.localdate(), time(9, 0)
                )
                received_at = _as_aware(received)
                status = _map_status(item["estado"])
                desc = item["desc"] or item["titulo"]
                cliente = item["cliente"]
                irregularidade = desc
                if cliente:
                    irregularidade = f"[{cliente}] {desc}".strip()

                avaliacao = ""
                retorno_at = None
                if status == SuporteClaroRegistro.STATUS_CONCLUIDO:
                    avaliacao = (
                        f"Importado da planilha DEMANDAS FY27 ({item['sheet']}). "
                        f"Estado origem: {item['estado'] or 'RESOLVIDA'}."
                    )
                    retorno_at = _as_aware(item["data_res"] or item["data"] or received)

                registro = SuporteClaroRegistro.objects.create(
                    titulo=item["titulo"][:255],
                    protocolo="",
                    irregularidade=irregularidade,
                    avaliacao=avaliacao,
                    received_at=received_at,
                    retorno_at=retorno_at,
                    sent_by=cliente[:255],
                    origem=SuporteClaroRegistro.ORIGEM_TEAMS,
                    categoria=SuporteClaroRegistro.CATEGORIA_INCIDENTE,
                    tipo_incidente=_infer_tipo(item["titulo"], item["desc"]),
                    status=status,
                    created_by=user,
                )
                auto_code = f"INC-{registro.id}"
                registro.protocolo = auto_code
                registro.save(update_fields=["protocolo", "updated_at"])
                save_protocolos(registro, [(auto_code, "")], user=user)
                save_chamados_externos(
                    registro,
                    [
                        (
                            SuporteClaroRegistro.CHAMADO_SERVICE,
                            item["numero"],
                            "",
                            "Rachel Ramos De Lima Pereira",
                        )
                    ],
                    user=user,
                )
                SuporteClaroImportRef.objects.create(
                    linha_id=linha_id,
                    registro=registro,
                    imported_by=user,
                )
                SuporteClaroHistorico.objects.create(
                    registro=registro,
                    user=user,
                    action=SuporteClaroHistorico.ACTION_IMPORT,
                    field_name="import",
                    old_value="",
                    new_value=(
                        f"DEMANDAS FY27 {item['numero']} "
                        f"({item['sheet']}, {item['estado'] or '-'})"
                    ),
                )
                created += 1
                self.stdout.write(
                    self.style.SUCCESS(
                        f"ok {item['numero']} -> #{registro.id} "
                        f"status={status} tratado_por={item['quem']}"
                    )
                )

        self.stdout.write(
            self.style.SUCCESS(f"Concluído: created={created} skipped={skipped}")
        )
