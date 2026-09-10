# -*- coding: utf-8 -*-

"""Exportação XLSX do relatório (Etapa 2)."""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from report_falhas.io.data_loader import norm_matricula, norm_protocolo, safe_str
from report_falhas.config_report import (
    COL_DATA_ANALISE,
    COL_DATA_AUDITORIA,
    COL_MATRICULA,
    COL_PROTOCOLO,
)


def strip_html(s: str) -> str:
    """Remove tags HTML simples para exportação em Excel/Texto."""
    if s is None:
        return ""
    try:
        import re

        txt = str(s)
        txt = txt.replace("&nbsp;", " ")
        txt = re.sub(r"<[^>]+>", "", txt)
        txt = re.sub(r"\s+", " ", txt).strip()
        return txt
    except Exception:
        return str(s)


def _snapshot_path(out_dir: Path) -> Path:
    return Path(out_dir) / "_snapshot_base_protocolos.pkl"


def load_protocol_snapshot(out_dir: Path) -> dict:
    p = _snapshot_path(out_dir)
    if not p.exists():
        return {}
    try:
        return pd.read_pickle(p)
    except Exception:
        return {}


def save_protocol_snapshot(out_dir: Path, payload: dict) -> None:
    p = _snapshot_path(out_dir)
    try:
        pd.to_pickle(payload, p)
    except Exception:
        pass


def compute_protocol_snapshot(df_base: pd.DataFrame) -> dict:
    """Gera snapshot simples de protocolos (normalizados) e contagens."""
    if df_base is None or df_base.empty or (COL_PROTOCOLO not in df_base.columns):
        return {
            "generated_at": datetime.now(),
            "row_count": int(len(df_base) if df_base is not None else 0),
            "protocol_counts": {},
            "unique_protocols": 0,
        }
    s = df_base[COL_PROTOCOLO].apply(norm_protocolo).apply(safe_str)
    s = s[s != ""]
    vc = s.value_counts()
    return {
        "generated_at": datetime.now(),
        "row_count": int(len(df_base)),
        "protocol_counts": {str(k): int(v) for k, v in vc.items()},
        "unique_protocols": int(vc.shape[0]),
    }


def diff_protocol_snapshots(prev: dict, cur: dict) -> dict:
    prev_counts = (prev or {}).get("protocol_counts") or {}
    cur_counts = (cur or {}).get("protocol_counts") or {}
    prev_set = set(prev_counts.keys())
    cur_set = set(cur_counts.keys())

    missing = sorted(prev_set - cur_set)
    added = sorted(cur_set - prev_set)

    changed = []
    for p in (prev_set & cur_set):
        dv = int(cur_counts.get(p, 0)) - int(prev_counts.get(p, 0))
        if dv != 0:
            changed.append((p, int(prev_counts.get(p, 0)), int(cur_counts.get(p, 0)), dv))
    changed_sorted = sorted(changed, key=lambda x: abs(x[3]), reverse=True)

    return {
        "prev_row_count": int((prev or {}).get("row_count", 0)),
        "cur_row_count": int((cur or {}).get("row_count", 0)),
        "prev_unique_protocols": int((prev or {}).get("unique_protocols", 0)),
        "cur_unique_protocols": int((cur or {}).get("unique_protocols", 0)),
        "missing_protocols": missing,
        "added_protocols": added,
        "changed_counts": changed_sorted,
    }


def export_protocol_diff_xlsx(proto_diff: dict, destino: Path) -> Path:
    """Exporta o diff de protocolos (snapshot Base) para um XLSX separado."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    proto_diff = proto_diff or {}

    resumo = [
        ("Linhas Base (Anterior)", int(proto_diff.get("prev_row_count", 0) or 0)),
        ("Linhas Base (Atual)", int(proto_diff.get("cur_row_count", 0) or 0)),
        (
            "Protocolos únicos (Anterior)",
            int(proto_diff.get("prev_unique_protocols", 0) or 0),
        ),
        ("Protocolos únicos (Atual)", int(proto_diff.get("cur_unique_protocols", 0) or 0)),
    ]

    missing = [safe_str(x) for x in (proto_diff.get("missing_protocols") or []) if safe_str(x)]
    added = [safe_str(x) for x in (proto_diff.get("added_protocols") or []) if safe_str(x)]
    changed = proto_diff.get("changed_counts") or []

    try:
        df_resumo = pd.DataFrame(resumo, columns=["Indicador", "Valor"])
        df_missing = pd.DataFrame({"Protocolo": missing})
        df_added = pd.DataFrame({"Protocolo": added})
        df_changed = (
            pd.DataFrame(changed, columns=["Protocolo", "Anterior", "Atual", "Delta"])
            if changed
            else pd.DataFrame(columns=["Protocolo", "Anterior", "Atual", "Delta"])
        )

        with pd.ExcelWriter(destino, engine="openpyxl", mode="w") as writer:
            df_resumo.to_excel(writer, sheet_name="Resumo", index=False)
            df_missing.to_excel(writer, sheet_name="Missing", index=False)
            df_added.to_excel(writer, sheet_name="Added", index=False)
            df_changed.to_excel(writer, sheet_name="Changed", index=False)
    except Exception:
        # Falha na exportação não pode quebrar o report
        pass

    return destino


def export_reincidencia_xlsx(
    insights,
    df_reinc_full: pd.DataFrame,
    df_cur_reinc: pd.DataFrame,
    tempo_map_etapa: dict,
    tempo_map_casa: dict,
    nome_map: dict,
    atividade_atual_hc_map: dict,
    dificuldade_counts_map: dict,
    destino: Path,
    turno_map: dict | None = None,
    team_categoria_map: dict | None = None,
    etapa_map: dict | None = None,
    dificuldade_map: dict | None = None,
    spill_df: pd.DataFrame | None = None,
    df_reinc_full_oficial: pd.DataFrame | None = None,
    df_cur_reinc_oficial: pd.DataFrame | None = None,
    kpis_total: dict | None = None,
    kpis_oficial: dict | None = None,
    rank_tipo_uf_total: dict | None = None,
    rank_tipo_uf_oficial: dict | None = None,
    ult3m_total_df: pd.DataFrame | None = None,
    ult3m_oficial_df: pd.DataFrame | None = None,
    adicionar_col_nome_agente: bool = True,
    dual_metric_mode: bool = True,
):
    """Exporta um XLSX com abas TOTAL e OFICIAL (ou só OFICIAL a partir de jul/2026)."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    etapa_map = etapa_map or {}
    dificuldade_map = dificuldade_map or {}
    team_categoria_map = team_categoria_map or {}
    from report_falhas.team_category import CATEGORIA_NAO_CLASSIFICADO

    fill_red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fill_green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    def _apply_reinc_color(ws_):
        try:
            headers = {cell.value: cell.column for cell in ws_[1]}
            if "Reincidente" in headers:
                col_idx = headers["Reincidente"]
                col_letter = get_column_letter(col_idx)
                for row in range(2, ws_.max_row + 1):
                    cell = ws_[f"{col_letter}{row}"]
                    val = safe_str(cell.value).lower()
                    if val == "sim":
                        cell.fill = fill_red
                    elif val in ("não", "nao"):
                        cell.fill = fill_green
        except Exception:
            pass

    def _enrich(df_in: pd.DataFrame, etapa_map_in: dict):
        df_out = df_in.copy()
        if df_out is None or df_out.empty:
            return df_out
        if adicionar_col_nome_agente and COL_MATRICULA in df_out.columns:
            df_out["Matrícula"] = df_out[COL_MATRICULA]
            df_out["Team/Category"] = df_out[COL_MATRICULA].map(
                lambda m: team_categoria_map.get(norm_matricula(m), CATEGORIA_NAO_CLASSIFICADO)
            )
            df_out["Turno"] = df_out[COL_MATRICULA].map(lambda m: (turno_map or {}).get(norm_matricula(m), ""))
            df_out["Etapa (Mês Atual)"] = df_out[COL_MATRICULA].map(
                lambda m: etapa_map_in.get(norm_matricula(m), "")
            )
        if COL_MATRICULA in df_out.columns:
            df_out["Atividade atual (HC)"] = df_out[COL_MATRICULA].map(
                lambda m: atividade_atual_hc_map.get(norm_matricula(m), "")
            )
            df_out["Nível de Dificuldade (contagem)"] = df_out[COL_MATRICULA].map(
                lambda m: dificuldade_counts_map.get(norm_matricula(m), "")
            )
            df_out["Tempo desde última Alteração de Atividade (HC)"] = df_out[COL_MATRICULA].map(
                lambda m: tempo_map_etapa.get(norm_matricula(m), "—")
            )
            df_out["Tempo de Casa"] = df_out[COL_MATRICULA].map(
                lambda m: tempo_map_casa.get(norm_matricula(m), "—")
            )
        return df_out

    df_total = _enrich(df_reinc_full if df_reinc_full is not None else pd.DataFrame(), etapa_map)

    with pd.ExcelWriter(destino, engine="openpyxl", mode="w") as writer:
        if dual_metric_mode:
            df_total.to_excel(writer, sheet_name="Reincidência (Total)", index=False)
            _apply_reinc_color(writer.sheets["Reincidência (Total)"])

        try:
            if dual_metric_mode and ult3m_total_df is not None and not ult3m_total_df.empty:
                ult3m_total_df.to_excel(writer, sheet_name="Falhas últimos 3 meses (Total)", index=False)
            if ult3m_oficial_df is not None and not ult3m_oficial_df.empty:
                ult3m_oficial_df.to_excel(writer, sheet_name="Falhas últimos 3 meses (Oficial)", index=False)
        except Exception:
            pass

        try:
            def _write_rank(sheet_name: str, rank: dict):
                if not rank or (not rank.get("impacto") and not rank.get("crescimento")):
                    return
                impact = pd.DataFrame(rank.get("impacto") or [])
                growth = pd.DataFrame(rank.get("crescimento") or [])
                startrow = 0
                if not impact.empty:
                    impact = impact.rename(columns={"label": "Tipo | UF", "cur": "Atual", "share": "%"})
                    impact["%"] = impact["%"].apply(
                        lambda v: round(float(v), 1)
                        if v is not None and not (hasattr(pd, "isna") and pd.isna(v))
                        else ""
                    )
                    impact = impact[["Tipo | UF", "Atual", "%"]]
                    impact.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow)
                    startrow += len(impact) + 3
                if not growth.empty:
                    growth = growth.rename(
                        columns={
                            "label": "Tipo | UF",
                            "prev": "Anterior",
                            "cur": "Atual",
                            "delta": "Δ",
                            "delta_pct": "Δ%",
                        }
                    )
                    if "Δ%" in growth.columns:
                        growth["Δ%"] = growth["Δ%"].apply(
                            lambda v: ""
                            if v is None or (hasattr(pd, "isna") and pd.isna(v))
                            else round(float(v), 0)
                        )
                    growth = growth[["Tipo | UF", "Anterior", "Atual", "Δ", "Δ%"]]
                    growth.to_excel(writer, sheet_name=sheet_name, index=False, startrow=startrow)

            if dual_metric_mode:
                _write_rank("Treinamento FF Tipo|UF (Total)", rank_tipo_uf_total or {})
            _write_rank("Treinamento FF Tipo|UF (Oficial)", rank_tipo_uf_oficial or {})
        except Exception:
            pass

        if dual_metric_mode and kpis_total is not None and isinstance(kpis_total, dict):
            pd.DataFrame(
                {
                    "Indicador": [
                        "Período (MTD)",
                        "Falhas (MTD) - Data de Análise",
                        "Falhas mês anterior (mesmos dias de calendário)",
                        "Variação (%)",
                        "Variação (delta)",
                        "Agentes reincidentes",
                        "Agentes novos no mês",
                        "Top 1 cenário (qtd)",
                        "Top 1 cenário (nome)",
                        "Período equivalente (mês anterior)",
                    ],
                    "Valor": [
                        kpis_total.get("periodo_mtd", ""),
                        kpis_total.get("fail_mtd_atual", 0),
                        kpis_total.get("fail_prev_equal", 0),
                        kpis_total.get("variacao_perc", ""),
                        kpis_total.get("variacao_delta", ""),
                        kpis_total.get("reinc_total", 0),
                        kpis_total.get("novos_mes_count", 0),
                        kpis_total.get("top1_cenario_qtd", 0),
                        kpis_total.get("top1_cenario_nome", ""),
                        kpis_total.get("periodo_equal_prev", ""),
                    ],
                }
            ).to_excel(writer, sheet_name="Consolidado (Total)", index=False)

        if df_reinc_full_oficial is not None and not df_reinc_full_oficial.empty:
            df_of = _enrich(df_reinc_full_oficial, etapa_map)
            df_of.to_excel(writer, sheet_name="Reincidência (Oficial)", index=False)
            _apply_reinc_color(writer.sheets["Reincidência (Oficial)"])

        if kpis_oficial is not None and isinstance(kpis_oficial, dict):
            pd.DataFrame(
                {
                    "Indicador": [
                        "Período (MTD)",
                        "Falhas (MTD) - Data de Análise",
                        "Falhas mês anterior (mesmos dias de calendário)",
                        "Variação (%)",
                        "Variação (delta)",
                        "Agentes reincidentes",
                        "Agentes novos no mês",
                        "Top 1 cenário (qtd)",
                        "Top 1 cenário (nome)",
                        "Período equivalente (mês anterior)",
                    ],
                    "Valor": [
                        kpis_oficial.get("periodo_mtd", ""),
                        kpis_oficial.get("fail_mtd_atual", 0),
                        kpis_oficial.get("fail_prev_equal", 0),
                        kpis_oficial.get("variacao_perc", ""),
                        kpis_oficial.get("variacao_delta", ""),
                        kpis_oficial.get("reinc_total", 0),
                        kpis_oficial.get("novos_mes_count", 0),
                        kpis_oficial.get("top1_cenario_qtd", 0),
                        kpis_oficial.get("top1_cenario_nome", ""),
                        kpis_oficial.get("periodo_equal_prev", ""),
                    ],
                }
            ).to_excel(writer, sheet_name="Consolidado (Oficial)", index=False)

        try:
            insights_list = [] if insights is None else ([insights] if isinstance(insights, str) else list(insights))
            insights_list = [strip_html(x) for x in insights_list]
            pd.DataFrame({"Insights": insights_list}).to_excel(writer, sheet_name="Insights", index=False)
        except Exception:
            pass

    try:
        if spill_df is not None and not isinstance(spill_df, dict) and not spill_df.empty:
            cols = []
            if COL_PROTOCOLO in spill_df.columns:
                cols.append(COL_PROTOCOLO)
            if COL_DATA_AUDITORIA in spill_df.columns:
                cols.append(COL_DATA_AUDITORIA)
            if COL_DATA_ANALISE in spill_df.columns:
                cols.append(COL_DATA_ANALISE)
            spill_to_save = spill_df[cols].copy() if cols else spill_df.copy()
            with pd.ExcelWriter(destino, engine="openpyxl", mode="a") as writer2:
                spill_to_save.to_excel(writer2, sheet_name="Fora de Fase", index=False)
    except Exception:
        pass
