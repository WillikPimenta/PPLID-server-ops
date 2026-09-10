from pathlib import Path
import pandas as pd
import csv
import time
import shutil
from typing import Iterable, Optional, Set, Union

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
except Exception:
    load_workbook = Workbook = None  # write_xlsx_from_df will check

DEFAULT_ENCODINGS = ("utf-8", "latin1", "cp1252")


def detect_delimiter(path: Union[str, Path], encodings: Iterable[str] = DEFAULT_ENCODINGS, sample_size: int = 4096) -> str:
    """Tenta detectar delimitador do CSV lendo um sample com vários encodings."""
    p = Path(path)
    for enc in encodings:
        try:
            with p.open("r", encoding=enc, errors="replace") as f:
                sample = f.read(sample_size)
                if not sample:
                    continue
                dialect = csv.Sniffer().sniff(sample)
                return dialect.delimiter
        except Exception:
            continue
    # fallback comum
    return ";"


def read_csv(path: Union[str, Path], encodings: Iterable[str] = DEFAULT_ENCODINGS, detect_sep: bool = True, **pd_kwargs) -> pd.DataFrame:
    """
    Robust read_csv: tenta múltiplos encodings e detecta delimitador.
    Retorna DataFrame (vazio se falhar).
    pd_kwargs são passados para pandas.read_csv (ex: on_bad_lines="skip").
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(p)
    pd_kwargs.setdefault("engine", "python")
    pd_kwargs.setdefault("on_bad_lines", "skip")
    if detect_sep:
        sep = detect_delimiter(p, encodings=encodings)
        pd_kwargs.setdefault("sep", sep)
    for enc in encodings:
        try:
            return pd.read_csv(p, encoding=enc, **pd_kwargs)
        except UnicodeDecodeError:
            continue
        except Exception:
            # parsing issues -> try next encoding
            continue
    # last resort: read with replace errors
    try:
        with p.open("r", encoding="utf-8", errors="replace") as f:
            return pd.read_csv(f, **pd_kwargs)
    except Exception:
        # return empty df if all fails
        return pd.DataFrame()


def write_csv(path: Union[str, Path], df: pd.DataFrame, index: bool = False):
    """Salva CSV garantindo pasta destino."""
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(p, index=index)


def write_xlsx_from_df(dest_path: Union[str, Path], df: pd.DataFrame, model_path: Optional[Union[str, Path]] = None,
                       sheet_name: str = None, table_name: Optional[str] = None):
    """
    Grava DataFrame em arquivo Excel/XLSX/XLSM.
    - Se model_path existir, ele é copiado para dest_path antes de escrever.
    - Usa openpyxl para preservar formato do modelo.
    """
    dest = Path(dest_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if model_path and Path(model_path).exists():
        shutil.copy2(model_path, dest)
    else:
        if Workbook is None:
            raise RuntimeError("openpyxl não disponível")
        Workbook().save(dest)
    wb = load_workbook(dest)
    ws = wb.active if sheet_name is None else (wb[sheet_name] if sheet_name in wb.sheetnames else wb.active)
    # clear existing cells
    try:
        if ws.max_row:
            ws.delete_rows(1, ws.max_row)
    except Exception:
        pass
    # write header + rows
    cols = list(df.columns)
    for c_idx, col in enumerate(cols, start=1):
        ws.cell(row=1, column=c_idx, value=col)
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, val in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)
    # optional table
    if table_name:
        try:
            ref = f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"
            tab = Table(displayName=table_name, ref=ref)
            style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            tab.tableStyleInfo = style
            if hasattr(ws._tables, "add"):
                ws._tables.add(tab)
            else:
                ws._tables.append(tab)
        except Exception:
            pass
    wb.save(dest)
    return dest


def list_csvs(directory: Union[str, Path]):
    """Retorna lista de Path para CSVs na pasta."""
    d = Path(directory)
    if not d.exists():
        return []
    return list(d.glob("*.csv"))


def newest_csv(directory: Union[str, Path], prefix: Optional[str] = None, suffix: Optional[str] = None) -> Optional[Path]:
    """Retorna o CSV mais recente que combine prefix/suffix (ou None)."""
    files = list_csvs(directory)
    if prefix:
        files = [f for f in files if f.name.startswith(prefix)]
    if suffix:
        files = [f for f in files if f.name.endswith(suffix)]
    if not files:
        return None
    return max(files, key=lambda p: p.stat().st_ctime)


def wait_for_new_file(before: Optional[Iterable[str]], directory: Union[str, Path], prefix: str, suffix: str, timeout: int = 60) -> Optional[Path]:
    """
    Observa 'directory' até que apareça um arquivo novo que comece com prefix e termine com suffix.
    'before' é um iterable de nomes (strings) existentes antes do download.
    Retorna Path do arquivo encontrado ou None.
    """
    start = time.time()
    d = Path(directory)
    before_set: Set[str] = set(before or [])
    while time.time() - start < timeout:
        if not d.exists():
            time.sleep(0.5)
            continue
        current = {p.name for p in d.iterdir()}
        added = current - before_set
        candidates = [d / name for name in added if name.startswith(prefix) and name.endswith(suffix)]
        if candidates:
            return max(candidates, key=lambda p: p.stat().st_ctime)
        time.sleep(0.5)
    return None