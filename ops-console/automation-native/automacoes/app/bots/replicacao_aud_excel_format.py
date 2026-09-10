"""Formatação visual do relatório Excel consolidado de replicação de auditoria."""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Sequence

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import (
    COLUNA_CONFIG_AUTOMATICOS,
    COLUNA_CONFIG_CATEGORIA,
    COLUNA_CONFIG_CLIENTE,
    COLUNA_CONFIG_MANUAIS,
    COLUNA_CONFIG_SEGMENTO,
    COLUNA_CONFIG_TOTAL,
)

ABA_VISAO_GERAL = "Visão Geral"
ABA_PLANO = "Plano"
ABA_RESUMO = "Resumo"
ABA_RESUMO_BRFLOW = "Resumo BRFlow"
ABA_RESUMO_CASE = "Resumo Case"
ABA_DASHBOARD = "Dashboard"

COR_CABECALHO = "1F4E79"
COR_CABECALHO_TEXTO = "FFFFFF"
COR_TOTAL = "E7EEF7"
COR_SECAO = "D9E2F3"
COR_SALVO_OK = "C6EFCE"
COR_ERRO = "FFC7CE"
COR_PENDENTE = "FFEB9C"
COR_NEUTRO = "F2F2F2"

COLUNAS_RESUMO_ORDEM: List[str] = [
    "Workflow",
    "Workflow BRFlow",
    "Fila",
    "Subpasta",
    COLUNA_CONFIG_CLIENTE,
    COLUNA_CONFIG_SEGMENTO,
    COLUNA_CONFIG_CATEGORIA,
    "Workflow D1",
    "Amostra Diaria",
    "Amostra Solicitada",
    "Amostra Redistribuida",
    "Amostra Efetiva",
    "Protocolos Salvos",
    "Pct Atingido",
    "Status",
    "Status BRFlow",
    "Data Hora Upload BRFlow",
    "Data Hora Ultima Atualizacao BRFlow",
    "Disponivel D1",
    "Horas Utilizadas",
    "Faixa Horaria",
    COLUNA_CONFIG_AUTOMATICOS,
    COLUNA_CONFIG_MANUAIS,
    COLUNA_CONFIG_TOTAL,
    "Arquivo CSV",
    "Observacao",
    "Amostra Total Config",
    "Amostra Conf Prod Config",
    "RunId",
    "Parquet Referencia",
    "Data Referencia D1",
    "Data Execucao",
    "Seed Amostra",
    "Pasta Volumetria",
]

COLUNAS_RESUMO_D1_ORDEM: List[str] = [
    "Workflow",
    "Workflow D1",
    "Workflow BRFlow",
    "Canal Destino",
    COLUNA_CONFIG_CLIENTE,
    COLUNA_CONFIG_SEGMENTO,
    COLUNA_CONFIG_CATEGORIA,
    "Amostra Diaria",
    "Amostra Solicitada",
    "Amostra Efetiva",
    "Protocolos Salvos",
    "Pct Atingido",
    "Status",
    "Status BRFlow",
    "Data Hora Upload BRFlow",
    "Disponivel D1",
    "Faixa Horaria",
    "Observacao",
]

COLUNAS_METADADOS_OCULTAR = frozenset(
    {
        "RunId",
        "Parquet Referencia",
        "Data Referencia D1",
        "Data Execucao",
        "Seed Amostra",
        "Pasta Volumetria",
    }
)

COLUNAS_METADADOS_OCULTAR_D1 = COLUNAS_METADADOS_OCULTAR | frozenset(
    {
        "Fila",
        "Subpasta",
        "Amostra Redistribuida",
        "Data Hora Ultima Atualizacao BRFlow",
        "Horas Utilizadas",
        COLUNA_CONFIG_AUTOMATICOS,
        COLUNA_CONFIG_MANUAIS,
        COLUNA_CONFIG_TOTAL,
        "Arquivo CSV",
        "Arquivo CSV Relativo",
        "Amostra Total Config",
        "Amostra Conf Prod Config",
        "Auditores Ativos",
        "RedistribuicaoTier",
        "Fonte Volumetria",
        "Excluidos Historico",
    }
)

COLUNAS_NUMERICAS_RESUMO = frozenset(
    {
        "Amostra Diaria",
        "Amostra Solicitada",
        "Amostra Redistribuida",
        "Amostra Efetiva",
        "Protocolos Salvos",
        "Excluidos Historico",
        "Disponivel D1",
        "Horas Utilizadas",
        COLUNA_CONFIG_AUTOMATICOS,
        COLUNA_CONFIG_MANUAIS,
        COLUNA_CONFIG_TOTAL,
        "Amostra Total Config",
        "Amostra Conf Prod Config",
    }
)

STATUS_BRFLOW_CORES = {
    "SALVO_OK": COR_SALVO_OK,
    "UPLOAD_OK": COR_SALVO_OK,
    "ERRO": COR_ERRO,
    "PENDENTE": COR_PENDENTE,
    "INATIVO": COR_NEUTRO,
    "PULADO": COR_NEUTRO,
}

STATUS_AMOSTRA_CORES = {
    "OK": COR_SALVO_OK,
    "PARCIAL": COR_PENDENTE,
    "SEM_REGISTRO_D1": COR_NEUTRO,
    "VAZIO": COR_NEUTRO,
}


def _fill(hex_cor: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_cor)


def _font(bold: bool = False, color: str = "000000", size: int = 11) -> Font:
    return Font(bold=bold, color=color, size=size)


def _thin_border() -> Border:
    side = Side(style="thin", color="B4B4B4")
    return Border(left=side, right=side, top=side, bottom=side)


def ordenar_colunas_resumo(df: pd.DataFrame, *, d1: bool = False) -> pd.DataFrame:
    """Reordena colunas do resumo para leitura agrupada."""
    if df.empty:
        return df
    ordem = COLUNAS_RESUMO_D1_ORDEM if d1 else COLUNAS_RESUMO_ORDEM
    existentes = [c for c in ordem if c in df.columns]
    extras = [c for c in df.columns if c not in existentes]
    return df[existentes + extras]


def montar_dataframe_dashboard_secoes(metricas: Dict[str, Any]) -> pd.DataFrame:
    """Dashboard com colunas Secao | Metrica | Valor."""
    secoes: List[tuple[str, List[str]]] = [
        (
            "Execucao",
            [
                "RunId",
                "DataReferenciaD1",
                "ParquetReferencia",
                "PastaVolumetria",
                "PastaProtocolos",
                "PastaResumo",
                "RelatorioExcel",
                "DataInicioExecucao",
                "DataFimExecucao",
            ],
        ),
        (
            "Amostragem",
            [
                "TotalWorkflowsConfig",
                "WorkflowsComD1",
                "WorkflowsSemD1",
                "AmostraSolicitadaTotal",
                "AmostraEfetivaTotal",
                "ProtocolosSalvosTotal",
                "PctAtingidoGlobal",
                "PoolRedistribuido",
                "ExcluidosHistoricoTotal",
                "WorkflowsParciais",
                "WorkflowsComUpload",
                "TempoEstimadoMin",
            ],
        ),
        (
            "BRFlow",
            [
                "WorkflowsSalvosBRFlow",
                "WorkflowsErroBRFlow",
                "WorkflowsPuladosBRFlow",
                "WorkflowsInativosBRFlow",
                "WorkflowsPendentesBRFlow",
            ],
        ),
        (
            "Capacidade",
            [
                "AuditoresAtivos",
                "MetaProdu",
                "CapacidadeProdutiva",
                "SomaAmostraDiaria",
                "SomaAmostraAjustada",
                "FatorCapacidade",
                "AderenciaCapacidadePct",
            ],
        ),
    ]
    linhas: List[dict] = []
    for secao, chaves in secoes:
        linhas.append({"Secao": secao, "Metrica": "", "Valor": ""})
        for chave in chaves:
            if chave in metricas:
                linhas.append({"Secao": "", "Metrica": chave, "Valor": metricas[chave]})
    return pd.DataFrame(linhas)


def montar_dataframe_dashboard_secoes_d1(metricas: Dict[str, Any]) -> pd.DataFrame:
    """Dashboard D-1 enxuto: capacidade e redistribuição (sem duplicar Visão Geral)."""
    secoes: List[tuple[str, List[str]]] = [
        (
            "Capacidade",
            [
                "AuditoresAtivos",
                "AuditoresAtivosCase",
                "MetaProdu",
                "CapacidadeProdutiva",
                "SomaAmostraDiaria",
                "FatorCapacidade",
                "AderenciaCapacidadePct",
            ],
        ),
        (
            "Redistribuicao",
            [
                "PoolRedistribuido",
                "ExcluidosHistoricoTotal",
                "VolumeRedistribuicaoNaoAlocado",
            ],
        ),
    ]
    linhas: List[dict] = []
    for secao, chaves in secoes:
        linhas.append({"Secao": secao, "Metrica": "", "Valor": ""})
        for chave in chaves:
            if chave in metricas:
                linhas.append({"Secao": "", "Metrica": chave, "Valor": metricas[chave]})
    return pd.DataFrame(linhas)


def _kpi_distribuicao_canal(df_resumo: pd.DataFrame) -> Dict[str, int]:
    if df_resumo.empty or "Canal Destino" not in df_resumo.columns:
        return {"workflows_brflow": 0, "workflows_case": 0, "protocolos_brflow": 0, "protocolos_case": 0}
    wf = df_resumo[df_resumo["Workflow"].astype(str) != "TOTAL"]
    mask_br = wf["Canal Destino"].astype(str) == "BRFlow"
    mask_case = wf["Canal Destino"].astype(str) == "Case Manager"
    prot_col = "Protocolos Salvos" if "Protocolos Salvos" in wf.columns else None
    return {
        "workflows_brflow": int(mask_br.sum()),
        "workflows_case": int(mask_case.sum()),
        "protocolos_brflow": int(wf.loc[mask_br, prot_col].sum()) if prot_col else 0,
        "protocolos_case": int(wf.loc[mask_case, prot_col].sum()) if prot_col else 0,
    }


def _metricas_dict_de_dashboard(df_dashboard: pd.DataFrame) -> Dict[str, Any]:
    if df_dashboard.empty:
        return {}
    if "Secao" in df_dashboard.columns:
        mask = df_dashboard["Metrica"].astype(str).str.len() > 0
        sub = df_dashboard.loc[mask, ["Metrica", "Valor"]]
        return dict(zip(sub["Metrica"], sub["Valor"]))
    if "Metrica" in df_dashboard.columns:
        return dict(zip(df_dashboard["Metrica"], df_dashboard["Valor"]))
    return {}


def _kpi_resumo(df_resumo: pd.DataFrame) -> Dict[str, Any]:
    if df_resumo.empty:
        return {}
    wf = df_resumo[df_resumo["Workflow"].astype(str) != "TOTAL"]
    total_efetiva = int(wf["Amostra Efetiva"].sum()) if "Amostra Efetiva" in wf.columns else 0
    total_salvos = int(wf["Protocolos Salvos"].sum()) if "Protocolos Salvos" in wf.columns else 0
    pct = round(100 * total_salvos / max(1, total_efetiva), 1)
    return {
        "protocolos_salvos": total_salvos,
        "amostra_efetiva": total_efetiva,
        "pct_atingido": pct,
        "ok": int((wf["Status"] == "OK").sum()) if "Status" in wf.columns else 0,
        "parcial": int((wf["Status"] == "PARCIAL").sum()) if "Status" in wf.columns else 0,
        "sem_d1": int((wf["Status"] == "SEM_REGISTRO_D1").sum()) if "Status" in wf.columns else 0,
    }


def escrever_aba_visao_geral(
    ws,
    plano,
    df_resumo: pd.DataFrame,
    estado: Optional[dict],
    metricas: Dict[str, Any],
) -> None:
    """Monta capa executiva com KPIs e tabela compacta de workflows."""
    ws.merge_cells("A1:F1")
    titulo = ws["A1"]
    titulo.value = "Replicação de Auditoria — Relatório Consolidado"
    titulo.font = _font(bold=True, color=COR_CABECALHO_TEXTO, size=14)
    titulo.fill = _fill(COR_CABECALHO)
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    kpi = _kpi_resumo(df_resumo)
    estado = estado or {}
    brflow = {
        "salvos": metricas.get("WorkflowsSalvosBRFlow", 0),
        "erro": metricas.get("WorkflowsErroBRFlow", 0),
        "pendente": metricas.get("WorkflowsPendentesBRFlow", 0),
        "pulado": metricas.get("WorkflowsPuladosBRFlow", 0),
        "inativo": metricas.get("WorkflowsInativosBRFlow", 0),
    }

    blocos: List[tuple[str, List[tuple[str, Any]]]] = [
        (
            "Identificação da execução",
            [
                ("RunId", plano.run_id),
                ("Data Referência D-1", metricas.get("DataReferenciaD1", plano.data_referencia_d1_fmt)),
                ("Data Execução", plano.data_execucao_fmt or metricas.get("DataFimExecucao", "")),
                ("Parquet", metricas.get("ParquetReferencia", plano.parquet_referencia)),
                ("Pasta Volumetria", metricas.get("PastaVolumetria", plano.pasta_volumetria)),
                ("Seed Amostra", getattr(plano, "seed_amostra", "")),
            ],
        ),
        (
            "Resultado da amostragem",
            [
                ("Protocolos salvos", kpi.get("protocolos_salvos", 0)),
                ("Amostra efetiva", kpi.get("amostra_efetiva", 0)),
                ("% atingido global", f"{kpi.get('pct_atingido', 0)}%"),
                ("Workflows OK", kpi.get("ok", 0)),
                ("Workflows parciais", kpi.get("parcial", 0)),
                ("Workflows sem D-1", kpi.get("sem_d1", 0)),
            ],
        ),
        (
            "Resultado BRFlow",
            [
                ("Workflows salvos no BRFlow", brflow["salvos"]),
                ("Workflows com erro", brflow["erro"]),
                ("Workflows pendentes", brflow["pendente"]),
                ("Workflows pulados", brflow["pulado"]),
                ("Workflows inativos", brflow["inativo"]),
                ("Início execução Selenium", metricas.get("DataInicioExecucao", estado.get("iniciado_em", ""))),
                ("Fim / atualização relatório", metricas.get("DataFimExecucao", "")),
            ],
        ),
    ]

    canal = _kpi_distribuicao_canal(df_resumo)
    if canal["workflows_brflow"] + canal["workflows_case"] > 0:
        blocos.append(
            (
                "Distribuição por canal",
                [
                    ("Workflows BRFlow", canal["workflows_brflow"]),
                    ("Workflows Case Manager", canal["workflows_case"]),
                    ("Protocolos BRFlow", canal["protocolos_brflow"]),
                    ("Protocolos Case Manager", canal["protocolos_case"]),
                ],
            )
        )

    linha = 3
    for titulo_bloco, pares in blocos:
        ws.cell(row=linha, column=1, value=titulo_bloco).font = _font(bold=True, size=12)
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        linha += 1
        for rotulo, valor in pares:
            ws.cell(row=linha, column=1, value=rotulo).font = _font(bold=True)
            ws.cell(row=linha, column=2, value=valor)
            linha += 1
        linha += 1

    ws.cell(row=linha, column=1, value="Status por workflow").font = _font(bold=True, size=12)
    linha += 1
    cabecalhos = [
        "Workflow",
        "Canal Destino",
        "Workflow BRFlow",
        "Status amostra",
        "Status BRFlow",
        "Data upload BRFlow",
        "% atingido",
    ]
    for col, nome in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=linha, column=col, value=nome)
        cell.font = _font(bold=True, color=COR_CABECALHO_TEXTO)
        cell.fill = _fill(COR_CABECALHO)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    linha += 1

    wf_rows = df_resumo[df_resumo["Workflow"].astype(str) != "TOTAL"].head(50)
    for _, row in wf_rows.iterrows():
        valores = [
            row.get("Workflow", ""),
            row.get("Canal Destino", ""),
            row.get("Workflow BRFlow", ""),
            row.get("Status", ""),
            row.get("Status BRFlow", ""),
            row.get("Data Hora Upload BRFlow", ""),
            row.get("Pct Atingido", ""),
        ]
        for col, val in enumerate(valores, start=1):
            cell = ws.cell(row=linha, column=col, value=val)
            cell.border = _thin_border()
            if col == 4:
                _aplicar_cor_status(cell, str(val), STATUS_BRFLOW_CORES)
            elif col == 3:
                _aplicar_cor_status(cell, str(val), STATUS_AMOSTRA_CORES)
        linha += 1

    _ajustar_larguras(ws, max_col=6, cap=40)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 36


def _aplicar_cor_status(cell, valor: str, mapa: dict) -> None:
    cor = mapa.get(valor.strip().upper() if valor else "", None)
    if cor:
        cell.fill = _fill(cor)


def _ajustar_larguras(ws, max_col: int, cap: int = 45) -> None:
    for col_idx in range(1, max_col + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), cap)


def _formatar_cabecalho_planilha(ws, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _font(bold=True, color=COR_CABECALHO_TEXTO)
        cell.fill = _fill(COR_CABECALHO)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _thin_border()
    ws.row_dimensions[1].height = 28


def _indice_coluna(ws, nome: str) -> Optional[int]:
    for col in range(1, ws.max_column + 1):
        if str(ws.cell(row=1, column=col).value or "") == nome:
            return col
    return None


def _formatar_aba_resumo(ws, *, d1: bool = False) -> None:
    if ws.max_row < 1:
        return
    num_cols = ws.max_column
    _formatar_cabecalho_planilha(ws, num_cols)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    col_workflow = _indice_coluna(ws, "Workflow")
    col_status = _indice_coluna(ws, "Status")
    col_status_br = _indice_coluna(ws, "Status BRFlow")
    col_pct = _indice_coluna(ws, "Pct Atingido")
    ocultar = COLUNAS_METADADOS_OCULTAR_D1 if d1 else COLUNAS_METADADOS_OCULTAR

    for row_idx in range(2, ws.max_row + 1):
        wf_val = ""
        if col_workflow:
            wf_val = str(ws.cell(row=row_idx, column=col_workflow).value or "")
        if wf_val == "TOTAL":
            for col in range(1, num_cols + 1):
                cell = ws.cell(row=row_idx, column=col)
                cell.font = _font(bold=True)
                cell.fill = _fill(COR_TOTAL)
            continue
        if col_status:
            val = str(ws.cell(row=row_idx, column=col_status).value or "")
            _aplicar_cor_status(ws.cell(row=row_idx, column=col_status), val, STATUS_AMOSTRA_CORES)
        if col_status_br:
            val = str(ws.cell(row=row_idx, column=col_status_br).value or "")
            _aplicar_cor_status(ws.cell(row=row_idx, column=col_status_br), val, STATUS_BRFLOW_CORES)
        if col_pct:
            cell = ws.cell(row=row_idx, column=col_pct)
            if isinstance(cell.value, (int, float)):
                cell.number_format = '0.0"%"'
        for col in range(1, num_cols + 1):
            nome = str(ws.cell(row=1, column=col).value or "")
            if nome in COLUNAS_NUMERICAS_RESUMO:
                cell = ws.cell(row=row_idx, column=col)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0"

    for col in range(1, num_cols + 1):
        nome = str(ws.cell(row=1, column=col).value or "")
        letter = get_column_letter(col)
        if nome in ocultar:
            ws.column_dimensions[letter].hidden = True

    _ajustar_larguras(ws, num_cols)


def _formatar_aba_plano(ws) -> None:
    if ws.max_row < 1:
        return
    num_cols = ws.max_column
    _formatar_cabecalho_planilha(ws, num_cols)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _ajustar_larguras(ws, num_cols)


def _formatar_aba_dashboard(ws) -> None:
    if ws.max_row < 1:
        return
    num_cols = min(ws.max_column, 3)
    _formatar_cabecalho_planilha(ws, num_cols)

    for row_idx in range(2, ws.max_row + 1):
        secao = str(ws.cell(row=row_idx, column=1).value or "").strip()
        metrica = str(ws.cell(row=row_idx, column=2).value or "").strip()
        if secao and not metrica:
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=3)
            cell = ws.cell(row=row_idx, column=1)
            cell.font = _font(bold=True, color=COR_CABECALHO)
            cell.fill = _fill(COR_SECAO)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            continue
        valor = ws.cell(row=row_idx, column=3).value
        if isinstance(valor, (int, float)):
            ws.cell(row=row_idx, column=3).alignment = Alignment(horizontal="right")
            if isinstance(valor, float) and not valor.is_integer():
                ws.cell(row=row_idx, column=3).number_format = "#,##0.00"
            else:
                ws.cell(row=row_idx, column=3).number_format = "#,##0"

    _ajustar_larguras(ws, num_cols, cap=50)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 32


def aplicar_formato_relatorio_excel(writer, *, d1: bool = False) -> None:
    """Aplica estilos às abas já escritas no ExcelWriter."""
    wb = writer.book
    if ABA_RESUMO in wb.sheetnames:
        _formatar_aba_resumo(wb[ABA_RESUMO], d1=d1)
    if ABA_PLANO in wb.sheetnames:
        _formatar_aba_plano(wb[ABA_PLANO])
    if ABA_DASHBOARD in wb.sheetnames:
        _formatar_aba_dashboard(wb[ABA_DASHBOARD])
    if ABA_VISAO_GERAL in wb.sheetnames:
        wb.active = wb.sheetnames.index(ABA_VISAO_GERAL)


def exportar_workbook_formatado(
    writer,
    plano,
    df_plano: pd.DataFrame,
    df_resumo: pd.DataFrame,
    df_dashboard: pd.DataFrame,
    estado: Optional[dict],
    *,
    modo_d1: bool = False,
    somente_plano_resumo: bool = False,
) -> None:
    """Escreve todas as abas no writer e aplica formatação."""
    df_resumo_ord = ordenar_colunas_resumo(df_resumo, d1=modo_d1)

    df_plano.to_excel(writer, sheet_name=ABA_PLANO, index=False)
    df_resumo_ord.to_excel(writer, sheet_name=ABA_RESUMO, index=False)

    if somente_plano_resumo:
        aplicar_formato_relatorio_excel(writer, d1=modo_d1)
        return

    metricas = _metricas_dict_de_dashboard(df_dashboard)

    if not modo_d1 and "Fila" in df_resumo_ord.columns:
        mask_case = df_resumo_ord["Fila"].astype(str).str.strip().isin(("3.1", "3,1"))
        df_brflow = df_resumo_ord[~mask_case]
        df_case = df_resumo_ord[mask_case]
        if not df_brflow.empty:
            df_brflow.to_excel(writer, sheet_name=ABA_RESUMO_BRFLOW, index=False)
        if not df_case.empty:
            df_case.to_excel(writer, sheet_name=ABA_RESUMO_CASE, index=False)
    df_dashboard.to_excel(writer, sheet_name=ABA_DASHBOARD, index=False)

    wb = writer.book
    ws_vg = wb.create_sheet(ABA_VISAO_GERAL, 0)
    escrever_aba_visao_geral(ws_vg, plano, df_resumo_ord, estado, metricas)

    aplicar_formato_relatorio_excel(writer, d1=modo_d1)
