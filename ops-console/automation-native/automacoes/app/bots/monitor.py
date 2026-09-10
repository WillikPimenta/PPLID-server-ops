"""
Compact Monitor de Eventos (usa csv_processing + selenium_helpers)
API: start(settings)/stop()/set_status_callback(fn)
"""

import sys
from pathlib import Path
import uuid


import os
import re
import time
import logging
import threading
import pandas as pd
from pathlib import Path as PathLib
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException, WebDriverException
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
import shutil
import csv
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from app.config.paths import (
    DEFAULT_DOWNLOAD,
    DEFAULT_INTERVAL,
    DEFAULT_MODELM,
    DEFAULT_SHAREPOINT,
    DEFAULT_TIMEOUT,
)
from app.config.selectors import brflow, okta
from app.core.credentials import get_credentials
from app.core.common import ErrorRecoveryHandler, safe_close_driver
from app.infrastructure.selenium_helpers import create_driver, click_element, send_keys_to_element, wait_for_element, take_error_screenshot, login_okta_resiliente
from app.infrastructure.csv_processing import read_csv as cp_read_csv, write_xlsx_from_df as cp_write_xlsx, wait_for_new_file as cp_wait_for_new_file


# Logger centralizado - usa o sistema configurado via configure_logging()
log = logging.getLogger("robots.monitor")

from app.core.bot_runtime import BotRuntime

_runtime = BotRuntime(mode="monitor")
set_status_callback = _runtime.set_status_callback
set_progress_callback = _runtime.set_progress_callback
parar_event = _runtime.parar_event
_set_status = _runtime.set_status
_set_progress = _runtime.set_progress
_reset_progress_state = _runtime.reset_progress_state
_begin_cycle = _runtime.begin_cycle


def _download_and_convert(driver, local_download, local_sharepoint, modelo, prefixo, extensao, timeout):
    _set_progress(30, "Monitor: acionando download CSV")
    # aciona CSV e espera arquivo
    if brflow and hasattr(brflow, "B_M_csv"):
        try:
            # helper espera (driver, selector_type, locator)
            click_element(driver, "id", brflow.B_M_csv)
            click_element(driver, "xpath", brflow.B_M_pesquisar)
        except TypeError:
            try:
                el = driver.find_element(By.ID, brflow.B_M_csv)
                driver.execute_script("arguments[0].click();", el)
            except Exception:
                pass
        except Exception:
            pass
    before = {p.name for p in Path(local_download).iterdir()} if Path(local_download).exists() else set()
    _set_progress(40, "Monitor: aguardando arquivo de download")
    novo = cp_wait_for_new_file(before, local_download, prefixo, extensao, timeout=timeout)
    if not novo:
        _set_progress(20, "Monitor: falha no download")
        return False, None
    caminho_csv = str(novo)
    destino = str(Path(local_sharepoint) / "Monitor de eventos.xlsx")
    _set_progress(60, "Monitor: convertendo CSV")
    df = cp_read_csv(caminho_csv, detect_sep=True)
    if df.empty:
        _set_progress(80, "Monitor: CSV vazio")
        return False, caminho_csv
    csv_para_xlsm_monitor_online(caminho_csv, destino, modelo, nome_tabela="TabelaMonitor")
    _set_progress(70, "Monitor: finalizando conversao")
    _set_progress(100, "Monitor: ciclo concluido")
    return True, caminho_csv

def _recreate_file_from_model(dest_path, model_path):
    """
    Remove dest_path se existir e copia model_path -> dest_path.
    Se model_path não existir, tenta encontrar o modelo no mesmo diretório do script.
    Se ainda não encontrar, cria um workbook vazio (.xlsx).
    """
    try:
        dest_dir = os.path.dirname(dest_path)
        if dest_dir and not os.path.isdir(dest_dir):
            os.makedirs(dest_dir, exist_ok=True)

        # tenta copiar o modelo original
        if model_path and os.path.exists(model_path):
            shutil.copy2(model_path, dest_path)
            return

        model_basename = os.path.basename(model_path) if model_path else None
        alt_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), model_basename) if model_basename else None
        if alt_path and os.path.exists(alt_path):
            shutil.copy2(alt_path, dest_path)
            return

        from openpyxl import Workbook
        wb = Workbook()
        wb.save(dest_path)
    except PermissionError:
        log.error(f"Permissão negada ao recriar {dest_path}")
        raise
    except Exception as e:
        log.error(f"Falha ao recriar {dest_path}")
        raise

def detectar_separador(caminho_csv, encodings=('latin1', 'utf-8')):
    """Tenta detectar delimitador lendo um sample maior; retorna (delimiter, encoding)."""
    for enc in encodings:
        try:
            with open(caminho_csv, 'r', encoding=enc, errors='ignore') as f:
                sample = f.read(4096)
                if not sample:
                    continue
                try:
                    dialect = csv.Sniffer().sniff(sample)
                    return dialect.delimiter, enc
                except Exception:
                    # fallback heurístico
                    if sample.count(';') > sample.count(','):
                        return ';', enc
                    return ',', enc
        except Exception:
            continue
    return ';', encodings[0]

def _normalizar_nome(col):
    import unicodedata
    return unicodedata.normalize('NFKD', str(col)).encode('ASCII', 'ignore').decode('ASCII').lower().strip()

def eh_matricula_val(valor):
    """
    Retorna True para usuário que começa com 'c', sequência de dígitos e termina com uma letra.
    Ex.: c12345a
    """
    v = re.sub(r'[^A-Za-z0-9]', '', str(valor)).lower().strip()
    return bool(re.match(r'^c\d+[a-z]$', v))

def _calc_status(linha, evento_col, data_col=None, horas_limite=6):
    try:
        ev = _normalizar_nome(str(linha.get(evento_col, "")))
    except Exception:
        ev = ""
    if ev == _normalizar_nome("logout"):
        return "Offline"
    if ev in (_normalizar_nome("Autenticação com sucesso"), _normalizar_nome("Autenticacao com sucesso")):
        if data_col and data_col in linha:
            try:
                ts = pd.to_datetime(linha[data_col], dayfirst=True, errors='coerce')
                if pd.isna(ts):
                    return "Offline"
                # usar pandas Timestamp para calcular diff com precisão
                diff_h = (pd.Timestamp.now() - ts).total_seconds() / 3600.0
                return "Online" if diff_h <= horas_limite else "Offline"
            except Exception:
                return "Offline"
        else:
            return "Offline"
    return "Offline"

def csv_para_xlsm_monitor_online(caminho_csv, caminho_xlsm, caminho_modelo_xlsm, nome_tabela="TabelaMonitor"):
    # usar csv_processing.read_csv (mais robusta) quando disponível
    try:
        sep, encoding = detectar_separador(caminho_csv)
    except Exception:
        sep, encoding = ',', 'utf-8'
    try:
        df = cp_read_csv(caminho_csv, detect_sep=True)
    except Exception:
        try:
            df = pd.read_csv(caminho_csv, sep=sep, encoding=encoding, engine='python', on_bad_lines='skip')
        except Exception as e:
            log.error(f"Falha ao ler CSV")
            df = pd.DataFrame()

    cols_norm = {col: _normalizar_nome(col) for col in df.columns}
    usuario_col = next((orig for orig, norm in cols_norm.items() if norm == 'usuario'), None)
    evento_col = next((orig for orig, norm in cols_norm.items() if norm == 'evento'), None)
    data_col = next((orig for orig, norm in cols_norm.items() if 'data' in norm), None)

    df_final = pd.DataFrame()
    if not df.empty and usuario_col and evento_col:
        # filtra por matrículas (ex: C99999A); se não achar, mantém todos
        df_matriculas = df[df[usuario_col].apply(eh_matricula_val)].copy()
        df_work = df_matriculas if not df_matriculas.empty else df.copy()

        # converte data (se houver) e garante ordenação cronológica asc por usuário
        if data_col:
            try:
                df_work.loc[:, data_col] = pd.to_datetime(df_work.loc[:, data_col], dayfirst=True, errors='coerce')
                df_work = df_work.sort_values(by=[usuario_col, data_col], ascending=[True, True])
            except Exception:
                logging.warning("[MONITOR CSV] Falha ao converter/ordenar data.")

        resultados = []
        # para cada usuário calcula o tempo a partir do PRIMEIRO login bem-sucedido
        success_keys = {_normalizar_nome("Autenticação com sucesso"), _normalizar_nome("Autenticacao com sucesso")}
        logout_norm = _normalizar_nome("logout")
        for usuario, grupo in df_work.groupby(usuario_col):
            # garante ordenação asc por data antes de operar (se data_col existir)
            if data_col and data_col in grupo.columns:
                try:
                    grupo = grupo.sort_values(by=data_col, ascending=True)
                except Exception:
                    pass

            # escolhe uma linha base para manter demais colunas (usa a última linha cronológica)
            base = grupo.iloc[-1].copy()

            # Se último evento for logout, considerar Offline imediatamente
            try:
                ultimo_ev_norm = _normalizar_nome(str(base.get(evento_col, "")))
            except Exception:
                ultimo_ev_norm = ""
            if ultimo_ev_norm == logout_norm:
                base['FirstLogin'] = None
                base['Status'] = "Offline"
                resultados.append(base)
                continue

            # procura o primeiro evento de sucesso (data mais antiga com evento sucesso)
            first_success_ts = None
            try:
                if data_col and evento_col in grupo.columns:
                    evs = grupo[evento_col].astype(str).apply(_normalizar_nome)
                    success_mask = evs.isin(success_keys)
                    if success_mask.any():
                        first_success_ts = pd.to_datetime(grupo.loc[success_mask, data_col].min(), dayfirst=True, errors='coerce')
                    else:
                        # fallback: usa a data mais antiga disponível no grupo
                        first_success_ts = pd.to_datetime(grupo[data_col].min(), dayfirst=True, errors='coerce')
            except Exception:
                first_success_ts = None

            # calcula status a partir do primeiro login (se disponível)
            status = None
            if first_success_ts is not None and not pd.isna(first_success_ts):
                try:
                    diff_h = (pd.Timestamp.now() - first_success_ts).total_seconds() / 3600.0
                    status = "Online" if diff_h <= 6 else "Offline"
                except Exception:
                    status = _calc_status(base, evento_col, data_col, horas_limite=6)
            else:
                # fallback para lógica anterior (último registro)
                status = _calc_status(base, evento_col, data_col, horas_limite=6)

            base['FirstLogin'] = first_success_ts
            base['Status'] = status
            resultados.append(base)

        if resultados:
            df_final = pd.DataFrame(resultados)

    # garante colunas mínimas se vazio
    if df_final.empty:
        headers = []
        if usuario_col:
            headers.append(usuario_col)
        if evento_col:
            headers.append(evento_col)
        if data_col:
            headers.append(data_col)
        headers.append('Status')
        df_final = pd.DataFrame(columns=headers)

    # adiciona id único por linha
    df_final['id'] = [str(uuid.uuid4()) for _ in range(len(df_final))]

    try:
        # Recria o arquivo destino a partir do modelo (.xlsx) e abre
        _recreate_file_from_model(caminho_xlsm, caminho_modelo_xlsm)
        wb = load_workbook(caminho_xlsm)
    except PermissionError:
        log.error(f"Permissão negada ao abrir workbook")
        return
    except Exception as e:
        log.error(f"Falha ao preparar workbook")
        return

    ws = wb.active
    try:
        if ws.max_row > 0:
            ws.delete_rows(1, ws.max_row)
    except Exception:
        pass

    try:
        if hasattr(ws, "_tables"):
            try:
                ws._tables.clear()
            except Exception:
                ws._tables = []
    except Exception:
        ws._tables = []

    # escreve dados (todos os usuários com último evento e Status)
    for c_idx, col_name in enumerate(df_final.columns, 1):
        ws.cell(row=1, column=c_idx, value=col_name)
    for r_idx, row in enumerate(df_final.values, 2):
        for c_idx, value in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=value)
    try:
        ref = f"A1:{ws.cell(row=ws.max_row, column=ws.max_column).coordinate}"
        tab = Table(displayName=nome_tabela, ref=ref)
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        tab.tableStyleInfo = style
        if hasattr(ws._tables, 'add'):
            ws._tables.add(tab)
        else:
            ws._tables.append(tab)
    except Exception as e:
        logging.warning(f"Não foi possível adicionar tabela")

    try:
        wb.save(caminho_xlsm)
    except PermissionError:
        log.error(f"Permissão negada ao salvar workbook")
    except Exception as e:
        log.error(f"Erro ao salvar workbook")

def baixar_monitor_loop(intervalo=DEFAULT_INTERVAL, tempo_max=None, settings=None):
    _reset_progress_state()
    _set_status("Monitor: iniciando loop")
    _set_progress(0, "Monitor: iniciando loop")
    inicio = time.time()
    ciclo = 1
    error_handler = ErrorRecoveryHandler(max_errors=5, retry_delay=2)
    s = settings or {}
    while not parar_event.is_set():
        try:
            _begin_cycle()
            if tempo_max and (time.time() - inicio) > tempo_max:
                _set_status("Tempo máximo atingido, reiniciando ciclo")
                inicio = time.time()
                ciclo += 1
            local_download = Path(s.get("local_download") or DEFAULT_DOWNLOAD)
            output_cfg = (os.getenv("MONITOR_OUTPUT_DIR") or os.getenv("ROBOT_OUTPUT_DIR") or "").strip()
            local_sharepoint = Path(s.get("local_sharepoint_usuario") or output_cfg or DEFAULT_SHAREPOINT)
            modelo = s.get("modelo_xlsm_monitor") or DEFAULT_MODELM
            prefixo = s.get("prefixo") or "relatorio_detalhado"
            extensao = s.get("extensao") or ".csv"
            matricula = s.get("matricula") or os.getenv("MONITOR_USER")
            senha = s.get("senha") or os.getenv("MONITOR_PASS")
            headless_raw = (os.getenv("MONITOR_HEADLESS") or os.getenv("ROBOT_HEADLESS") or os.getenv("HEADLESS") or "0").strip().lower()
            # start driver
            drv = create_driver(headless=headless_raw in ("1", "true", "yes", "on"), download_dir=str(local_download)) if create_driver else None
            wait = WebDriverWait(drv, 20)
            if drv is None or okta is None or brflow is None:
                _set_status("Dependências (selenium_helpers/path) não configuradas")
                break
            drv.get(okta.O_LINK)
            # login
            #verifique esse try except para o novo login okta
            try:
                user, pwd = matricula, senha
                if not user or not pwd:
                    creds = get_credentials()
                    user = user or creds[0]
                    pwd = pwd or creds[1]

                if not user or not pwd:
                    _set_status("Credenciais não fornecidas")
                    safe_close_driver(drv)
                    return
                login_okta_resiliente(drv, user, pwd, timeout=20)
            except TimeoutException as e:
                _set_status(f"Erro no login Okta: {e}")
                safe_close_driver(drv)

            # pesquisa brflow
            _set_progress(35, "BRFlow: pesquisando no Okta")
            send_keys_to_element(drv, "id", okta.O_pesquisar, "brflow")
            send_keys_to_element(drv, "id", okta.O_pesquisar, Keys.ENTER)
            # espera nova janela e navega para monitor
            WebDriver = getattr(__import__("selenium.webdriver.support.ui", fromlist=["WebDriverWait"]), "WebDriverWait")
            WebDriver(drv, 20).until(lambda d: len(d.window_handles) > 1)
            _set_progress(40, "BRFlow: abrindo nova janela")
            drv.switch_to.window(drv.window_handles[-1])
            
            # Aguarda a página carregar completamente
            try:
                wait.until(EC.element_to_be_clickable((By.XPATH, brflow.B_usuario)))
                _set_progress(45, "BRFlow: navegando para monitor")
            except TimeoutException:
                log.warning("Elemento B_usuario não encontrado, continuando mesmo assim")
            
            # Clica no menu Monitor com retries
            from app.infrastructure.brflow_monitor import navegar_monitor_e_baixar_csv
            from datetime import date as date_cls

            caminho_csv = navegar_monitor_e_baixar_csv(
                drv,
                local_download,
                data_ref=date_cls.today(),
                prefixo=prefixo,
                extensao=extensao,
                timeout=DEFAULT_TIMEOUT,
                set_progress=_set_progress,
            )
            ok = False
            if caminho_csv:
                destino = str(Path(local_sharepoint) / "Monitor de eventos.xlsx")
                csv_para_xlsm_monitor_online(str(caminho_csv), destino, modelo, nome_tabela="TabelaMonitor")
                ok = True
                try:
                    os.remove(caminho_csv)
                except Exception:
                    pass
            _set_progress(100, "Monitor: ciclo concluido" if ok else "Monitor: falha no download")
            _set_status(f"[Monitor] Ciclo {ciclo} finalizado ({'ok' if ok else 'erro'})")
            error_handler.reset()
            safe_close_driver(drv)
            # wait interval
            waited = 0.0
            while waited < intervalo and not parar_event.is_set():
                time.sleep(0.5); waited += 0.5
            ciclo += 1
        except (TimeoutException, WebDriverException) as e:
            take_error_screenshot(drv, str(type(e).__name__)[:30], "monitor")
            if not error_handler.handle_error(drv, "Selenium error, reiniciando ciclo", _set_status, log):
                break
            ciclo += 1
            continue
        except Exception as e:
            take_error_screenshot(drv, str(type(e).__name__)[:30], "monitor")
            if not error_handler.handle_error(drv, "Erro inesperado no loop do monitor", _set_status, log):
                break
            ciclo += 1
            continue
    _set_status("Monitor: loop encerrado")

# start/stop API
_thread = None
def start(settings=None, intervalo=DEFAULT_INTERVAL, tempo_max=None):
    global _thread
    parar_event.clear()
    if _thread and _thread.is_alive():
        return _thread
    _thread = threading.Thread(target=baixar_monitor_loop, kwargs={"intervalo": intervalo, "tempo_max": tempo_max, "settings": settings}, daemon=True)
    _thread.start()
    return _thread

def stop(timeout: int = 5):
    parar_event.set()
    global _thread
    if _thread:
        _thread.join(timeout=timeout)
    _thread = None

if __name__ == "__main__":
    start()
    try:
        while not parar_event.is_set():
            time.sleep(0.5)
    except KeyboardInterrupt:
        stop()